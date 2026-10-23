"""Questions and jobs on the person's own sheets, in English / Hinglish / Hindi. No AI.

understand(question, tables) -> a plan:
  totals, counts, averages, highest/lowest, "X-wise" summaries (also by month), top N, filtered lists,
  lookups from another sheet, invoice aging, combining sheets, highlighting rows, describing a sheet.
run(plan, path, out) -> the answer in words + a new sheet with the result + the Excel formula that does the same.
If the question isn't clear enough, it says exactly what's missing and offers the columns to pick. It never guesses wildly.
"""

import difflib
import re
from collections import defaultdict
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.ml.normalize import tokens
from app.office.sheetdata import Table, load_tables

SYNONYMS = {
    "amount": {"amount", "total", "sales", "sale", "revenue", "value", "bikri", "paisa", "paise", "rakam", "turnover", "net"},
    "salary": {"salary", "tankhwah", "tankha", "pay", "ctc", "wage"},
    "quantity": {"quantity", "qty", "units", "matra", "pieces", "nag"},
    "price": {"price", "rate", "daam", "cost", "mrp"},
    "date": {"date", "tarikh", "tareekh", "day", "din", "joining"},
    "name": {"name", "naam"},
    "department": {"department", "dept", "vibhag", "team"},
    "status": {"status", "sthiti", "state"},
    "city": {"city", "shehar", "location", "place", "jagah"},
    "region": {"region", "area", "zone", "kshetra"},
    "customer": {"customer", "client", "grahak", "party"},
    "vendor": {"vendor", "supplier"},
    "product": {"product", "item", "saman", "maal"},
    "rep": {"rep", "salesperson", "salesman", "executive", "employee"},
    "stock": {"stock", "inventory"},
    "invoice": {"invoice", "bill"},
    "tax": {"tax", "gst"},
    "discount": {"discount", "chhoot", "off"},
    "nights": {"nights", "night", "stay", "stays", "duration", "days stayed"},
    "rating": {"rating", "ratings", "rated", "stars", "star"},
}
def _pack_synonyms():
    """More ways people name columns, from the Excel agent's pack (so they grow without code changes)."""
    try:
        from app.mind import packs

        for k, ws in packs.get("office").get("column_synonyms", {}).items():
            SYNONYMS.setdefault(k, set()).update(ws)
            SYNONYMS[k].add(k)
    except Exception:  # noqa: BLE001
        pass


_pack_synonyms()
STOP = {"the", "a", "an", "of", "in", "on", "for", "to", "me", "mein", "mai", "ka", "ki", "ke", "ko", "se", "hai", "hain", "kya", "kitna",
        "kitne", "kitni", "batao", "bata", "do", "de", "dena", "chahiye", "karo", "kar", "is", "are", "what", "show", "give", "please", "pls",
        "and", "or", "aur", "sheet", "sheets", "data", "file", "excel", "wise", "har", "per", "each", "by", "total", "sum", "count",
        "average", "avg", "my", "mera", "meri", "mere", "this", "that", "yeh", "ye", "wo", "which", "kaun", "kaunsa", "kaunse", "sabse",
        "all", "sab", "saare", "list", "dikhao", "nikalo", "find", "how", "many", "much", "with", "jinka", "jinki", "jiska", "jo", "hua"}

AGG_WORDS = {
    "sum": {"total", "sum", "jod", "jodo", "kul", "sum", "totals", "jumla"},
    "count": {"count", "kitne", "kitni", "ginti", "number", "how many", "kitna"},
    "avg": {"average", "avg", "ausat", "mean", "aoast", "average"},
    "max": {"highest", "max", "maximum", "sabse zyada", "top", "most", "largest", "biggest", "zyada"},
    "min": {"lowest", "min", "minimum", "sabse kam", "least", "smallest", "kam"},
}
GROUP_RE = re.compile(r"\b(\w+)[- ]?wise\b|\bhar\s+(\w+)|\bper\s+(\w+)|\beach\s+(\w+)|\bby\s+(\w+)|\b(\w+)\s+ke\s+hisaab|\baccording to\s+(\w+)", re.I)
MONTH_RE = re.compile(r"\b(month|months|monthly|mahina|mahine|mahinewise|mahine ke hisaab|month[- ]?wise)\b", re.I)
TOPN_RE = re.compile(r"\b(?:top|pehle|first|sabse (?:zyada|upar))\s*(\d{1,3})\b|\b(\d{1,3})\s*(?:sabse|top|best)\b", re.I)


class NeedMore(Exception):
    """The question needs one more detail; options are what the person can pick."""

    def __init__(self, msg_hi, msg_en, options=()):
        super().__init__(msg_en)
        self.hi, self.en, self.options = msg_hi, msg_en, list(options)


# ------------------------------------------------------------------ matching words to sheets, columns and values
def _close(w, choices, cutoff=0.82):
    return bool(difflib.get_close_matches(w, list(choices), n=1, cutoff=cutoff)) if len(w) >= 4 else w in choices


def _header_words(h: str) -> list[str]:
    """Header words in the same spelling the question gets (Hindi script is written in Roman letters, like the question)."""
    ws = tokens(h) if re.search(r"[\u0900-\u097f]", h or "") else re.findall(r"[a-z0-9]+", (h or "").lower())
    return [w for w in ws if w not in {"of", "the", "no", "id", "practice"} or len(h) <= 3]


def _col_score(h: str, words: list[str]) -> float:
    words = [x for w in words for x in (tokens(w) if re.search(r"[\u0900-\u097f]", w) else [w])]
    hw = _header_words(h)
    if not hw:
        return 0.0
    hit, syn_hit = 0.0, False
    for x in hw:
        syns = [s for s in SYNONYMS.values() if x in s]
        if any(w == x or (len(w) >= 4 and _close(w, [x])) for w in words):
            hit += 1
        elif any(w in s for s in syns for w in words):
            hit += 0.6                                    # "sales" ~ "amount": weaker than the real word
            syn_hit = True
    score = hit / len(hw) + (0.01 * hit)
    if syn_hit and score < 0.52 and len(hw) <= 3:
        score = 0.52                                      # 'doctor' for 'Consultant', 'category' for 'Expense Head'
    return score


def pick_sheet(text: str, tables: list[Table]) -> Table | None:
    words = set(tokens(text))
    best, score = None, 0
    for t in tables:
        key = [w for w in re.findall(r"[a-z0-9]+", t.name.lower()) if w not in ("data", "sheet", "sheets")]
        if key and all(any(w == k or _close(w, [k]) for w in words) for k in key):
            if len(key) > score:
                best, score = t, len(key)
    return best


def named_sheet(text: str, tables: list[Table]) -> Table | None:
    """'the sales sheet', 'sheet Inventory', 'Sales Data mein', 'employee tab': the person named the sheet; use it."""
    low = (text or "").lower()
    hits = []
    for t in tables:
        if pick_sheet(text, [t]) is None:
            continue
        key = [w for w in re.findall(r"[a-z0-9]+", t.name.lower()) if w not in ("data", "sheet", "sheets")]
        k = r"\s+".join(re.escape(w) + r"\w*" for w in key)
        m = re.search(rf"\b{k}\s+(data\s+)?(sheet|tab|worksheet)s?\b|\b(sheet|tab)\s+(named\s+|called\s+)?{k}", low)
        if m:
            hits.append((0, -len(key), t))
        elif t.name.lower() in low and len(t.name.split()) >= 2:
            hits.append((1, -len(key), t))
    return min(hits, key=lambda x: x[:2])[2] if hits else None


def _value_not_sheet(text, t, tables):
    """'employees in Sales department': 'Sales' is a Department value in the Employee sheet, not the Sales sheet."""
    low = (text or "").lower()
    for o in tables:
        if o is t:
            continue
        for i, v in values_in(text, o):
            words = [w for w in _header_words(o.header[i]) if len(w) >= 4]
            if v.lower() in t.name.lower() and any(re.search(rf"\b{re.escape(v.lower())}\s+{re.escape(w)}", low) for w in words):
                return o
    return None


ROW_WORDS = {"record", "records", "row", "rows", "entries", "entry", "lines"}
ROWS_Q = re.compile(r"\b(how many|how much|kitne|kitni|kitna|number of|count|total)\b.{0,30}\b(records?|rows?|entries|entry|lines)\b|"
                    r"\b(records?|rows?|entries)\b.{0,20}\b(kitne|kitni|count|hain|hai)\b|(रिकॉर्ड|रो|एंट्री).{0,15}कितन", re.I)
NO_CHANGE = re.compile(r"\bbefore (mak\w*|doing|any|you make)\b.{0,12}\bchang|\b(don'?t|do not|dont)\s+(change|touch|modify|edit)|\bwithout chang|"
                       r"\bno changes?\b|\b(sirf|bas|just|only)\s+(batao|bata do|tell|show|dikhao)|\b(change|badlav|badal)\w*\s+mat\b|मत बदल", re.I)


def columns_in(text: str, t: Table, kinds=None) -> list[tuple[int, float]]:
    words = [w for w in tokens(text) if (w not in STOP or w in ("total", "sales", "net")) and w not in ROW_WORDS]
    if "total" in words and len(words) > 1:
        # 'total quantity': 'total' says add up, it isn't the column called Total, unless nothing else is named
        rest = [w for w in words if w != "total"]
        if any(h and _col_score(h, rest) >= 0.5 for i, h in enumerate(t.header) if not kinds or t.kind(i) in kinds):
            words = rest
    out = []
    for i, h in enumerate(t.header):
        if not h or (kinds and t.kind(i) not in kinds):
            continue
        s = _col_score(h, words)
        if s >= 0.5:
            if re.search(r"closing|current|final|balance", h, re.I):
                s += 0.005                                 # "stock" usually means the stock now
            out.append((i, s))
    return sorted(out, key=lambda x: -x[1])


def values_in(text: str, t: Table, skip_cols=()) -> list[tuple[int, str]]:
    """'East', 'pending', 'Priya' mentioned in the question -> (column, value) filters."""
    words = [w for w in tokens(text) if w not in STOP and len(w) >= 3]
    raw = (text or "").lower()
    found = []
    for i, h in enumerate(t.header):
        if i in skip_cols or t.kind(i) != "text":
            continue
        vals = {}
        for v in t.col(i):
            if isinstance(v, str) and v.strip():
                vals.setdefault(v.strip().lower(), v.strip())
            if len(vals) > 300:
                break
        if len(vals) > 300:
            continue
        for low, orig in vals.items():
            if len(low) < 3:
                continue
            if ((" " in low or "-" in low or "&" in low) and re.search(rf"(?<![\w]){re.escape(low)}(?![\w])", raw)) or \
                    any(w == low or (len(w) >= 5 and _close(w, [low], 0.85)) for w in words):
                found.append((i, orig))
        for low, orig in vals.items():                    # "IT", "HR": short values only when written in capitals
            if len(low) == 2 and re.search(rf"\b{re.escape(orig.upper())}\b", text or ""):
                found.append((i, orig))
    # "approved expenses", "paid bills": a Yes/No column named in the question means Yes (or No when negated)
    YES, NO = {"yes", "y", "haan", "true", "done", "ok"}, {"no", "n", "nahi", "false"}
    for i, h in enumerate(t.header):
        if i in skip_cols or not h or t.kind(i) != "text" or any(c == i for c, _ in found):
            continue
        vals = {str(v).strip() for v in t.col(i)[:5000] if v not in (None, "")}
        own = [w for w in words if w in set(_header_words(h)) and not any(
            j != i and hh and _col_score(hh, [w]) >= 0.5 for j, hh in enumerate(t.header))]
        if vals and {v.lower() for v in vals} <= YES | NO and own and _col_score(h, own) >= 0.5:
            neg = re.search(r"\b(not|nahi|nahin|bina|unapproved|unpaid|pending)\b", raw)
            pick = next((v for v in vals if v.lower() in (NO if neg else YES)), None)
            if pick:
                found.append((i, pick))
    # one value per column; prefer exact
    seen, out = set(), []
    for i, v in found:
        if i not in seen:
            seen.add(i)
            out.append((i, v))
    return out


def _number_filters(text, t, num_cols):
    """'18% gst wale': a number that is one of the values of a named number column -> filter on it."""
    out = []
    for m in re.finditer(r"(?<![\w.])(\d+(?:\.\d+)?)\s*(%?)", text or ""):
        n = float(m.group(1))
        for i, _ in num_cols:
            vals = {v for v in t.col(i)[:20000] if isinstance(v, (int, float))}
            for cand in (n, n / 100 if m.group(2) else None):
                if cand is not None and cand in vals and len(vals) <= 50:
                    out.append((i, int(cand) if float(cand).is_integer() else cand))
                    return out
    return out



# ------------------------------------------------------------------ the analyst's basics: conditions, dates, the main amount
_NUM = r"\d[\d,]*(?:\.\d+)?"
CMP_RE = re.compile(
    r"(?P<op>more than|greater than|bigger than|higher than|above|over|exceeding|exceeds|at least|atleast|minimum of|older than|"
    r"younger than|less than|lower than|fewer than|smaller than|below|under|at most|atmost|upto|up to|not more than|maximum of)\s*"
    r"(?:rs\.?|₹|inr)?\s*(?P<n>" + _NUM + r")\s*(?P<unit>%|k\b|lakh|lac|years?|yrs?|saal|sal|days?|din)?"
    r"|(?P<sym>>=|<=|>|<)\s*(?P<n2>" + _NUM + r")"
    r"|(?P<n3>" + _NUM + r")\s*(?P<unit3>%|k\b|lakh|lac|years?|yrs?|saal|sal|days?|din)?\s*(?:se|say|sey)\s*(?P<op3>zyada|jyada|jada|upar|adhik|kam|neeche|niche|kum)",
    re.I)
_OPWORD = {"more than": "gt", "greater than": "gt", "bigger than": "gt", "higher than": "gt", "above": "gt", "over": "gt", "exceeding": "gt",
           "exceeds": "gt", "older than": "gt", "at least": "ge", "atleast": "ge", "minimum of": "ge", "younger than": "lt", "less than": "lt",
           "lower than": "lt", "fewer than": "lt", "smaller than": "lt", "below": "lt", "under": "lt", "at most": "le", "atmost": "le",
           "upto": "le", "up to": "le", "not more than": "le", "maximum of": "le",
           "zyada": "gt", "jyada": "gt", "jada": "gt", "upar": "gt", "adhik": "gt", "kam": "lt", "neeche": "lt", "niche": "lt", "kum": "lt",
           ">": "gt", ">=": "ge", "<": "lt", "<=": "le"}
YEAR_RE = re.compile(r"(?:\b(?:in|of|during|year|saal|for|mein|me|se)\s+)?\b((?:19|20)\d{2})\b(?!\s*(?:se|-|to))", re.I)
MONTH_NAME_RE = re.compile(r"\b(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|(?:in|of|during)\s+may|may\s+(?:19|20)\d{2}|jun(?:e)?|jul(?:y)?|"
                           r"aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\b(?![- ]?wise)(?:\s+((?:19|20)\d{2}))?", re.I)
RECENT_RE = re.compile(r"\b(most recent(?:ly)?|recently|latest|newest|last (?:joined|added|one|entry|order|booking|visit)|sabse (?:naya|nayi|naye|latest)|"
                       r"youngest)\b", re.I)
OLDEST_RE = re.compile(r"\b(oldest|earliest|first (?:joined|one|entry|order|booking)|sabse (?:purana|purani|purane)|senior ?most|longest serving|eldest)\b", re.I)
MONEY_WORDS = re.compile(r"\b(sales|sale|revenue|turnover|income|earnings?|business|bikri|kamai|collection|takings|kitne ka|kitne ki)\b", re.I)
BEST_RE = re.compile(r"\b(?:best|top|biggest|highest|most valuable|worst|lowest|weakest)\s+(?:rated\s+|selling\s+|performing\s+)?(?!\d)(\w+(?:\s+\w+)?)", re.I)


def _scale(n: float, unit: str | None) -> float:
    u = (unit or "").lower()
    return n * 1000 if u == "k" else n * 100000 if u in ("lakh", "lac") else n


def _num_cols_of(t: Table) -> list[int]:
    from app.office import understand as U

    r = U.roles(t)
    return [i for i in range(len(t.header)) if t.header[i] and t.kind(i) == "number" and r.get(i, "amount") != "id"]


def _nearest_col(text: str, pos: int, t: Table, cols: list[int]) -> int | None:
    low = (text or "").lower()
    best, dist = None, 10 ** 9
    for i in cols:
        for w in _header_words(t.header[i]):
            if len(w) < 3:
                continue
            for m in re.finditer(rf"\b{re.escape(w)}\w*", low):
                d = abs(m.start() - pos)
                if d < dist:
                    best, dist = i, d
    return best if dist <= 45 else None


def _by_magnitude(n: float, t: Table, cols: list[int]) -> int | None:
    """'patients above 60': the number column whose usual values are like 60 (Age, not Bill Amount)."""
    import math
    import statistics

    best, score = None, 9.0
    for i in cols:
        vals = [v for v in t.col(i)[:5000] if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if len(vals) < 3:
            continue
        lo, hi, med = min(vals), max(vals), statistics.median(vals)
        if not (lo * 0.5 - 1 <= n <= hi * 1.5 + 1) or med <= 0 or n <= 0:
            continue
        d = abs(math.log10(n) - math.log10(med))
        if d < score:
            best, score = i, d
    return best if score < 0.8 else None


def _compare_filters(text: str, t: Table):
    """'more than 5 years', 'above 60', '50k se zyada', '> 1000' -> (column, ('cmp', op, n)); and the text without them."""
    out, spans = [], []
    cols = _num_cols_of(t)
    for m in CMP_RE.finditer(text or ""):
        if m.group("n") is not None:
            op, n, unit = _OPWORD[m.group("op").lower()], m.group("n"), m.group("unit")
        elif m.group("n2") is not None:
            op, n, unit = _OPWORD[m.group("sym")], m.group("n2"), None
        else:
            op, n, unit = _OPWORD[m.group("op3").lower()], m.group("n3"), m.group("unit3")
        val = _scale(float(n.replace(",", "")), unit)
        col = _nearest_col(text, m.start(), t, cols)
        if col is None and (unit or "").lower().startswith(("year", "yr", "saal")):
            col = next((i for i in reversed(cols) if re.search(r"age|years|experience|tenure|saal", t.header[i], re.I)), None)
        if col is None and m.group("op") and m.group("op").lower() in ("older than", "younger than"):
            col = next((i for i in reversed(cols) if re.search(r"\bage\b|umar", t.header[i], re.I)), None)
        if col is None:
            col = _by_magnitude(val, t, cols)
        if col is None:
            continue
        if unit == "%" and val > 1 and all(abs(v) <= 1 for v in t.col(col)[:200] if isinstance(v, (int, float))):
            val = val / 100
        out.append((col, ("cmp", op, int(val) if float(val).is_integer() else val)))
        spans.append(m.span())
    for a, b in reversed(spans):
        text = text[:a] + " " + text[b:]
    return out, text


def _date_cols(t: Table) -> list[int]:
    return [i for i in range(len(t.header)) if t.header[i] and t.kind(i) == "date"]


def _pick_date_col(text: str, t: Table, pos: int | None = None) -> int | None:
    dc = _date_cols(t)
    if not dc:
        return None
    if pos is not None:
        near = _nearest_col(text, pos, t, dc)
        if near is not None:
            return near
    low = (text or "").lower()
    for i in dc:
        if any(len(w) >= 4 and re.search(rf"\b{re.escape(w[:4])}", low) for w in _header_words(t.header[i])):
            return i
    return dc[0]


def _date_filters(text: str, t: Table):
    out, spans = [], []
    if not _date_cols(t):
        return out, text
    for m in MONTH_NAME_RE.finditer(text or ""):
        word = re.sub(r"^(in|of|during)\s+|\s+(19|20)\d{2}$", "", m.group(1).lower())
        mon = next((k + 1 for k, name in enumerate(MONTHS) if name.startswith(word[:3])), None)
        col = _pick_date_col(text, t, m.start())
        yr = m.group(2) or (re.search(r"(?:19|20)\d{2}", m.group(1)) or [None])[0]
        if mon and col is not None:
            out.append((col, ("month", mon, int(yr) if yr else None)))
            spans.append(m.span())
    if not any(v[0] == "month" for _, v in out):
        for m in YEAR_RE.finditer(text or ""):
            col = _pick_date_col(text, t, m.start())
            yr = int(m.group(1))
            if col is not None and not any(isinstance(v, (int, float)) and v == yr for v in t.col(col)[:50]):
                out.append((col, ("year", yr)))
                spans.append(m.span())
    for a, b in reversed(spans):
        text = text[:a] + " " + text[b:]
    return out, text


def _short_values(text: str, t: Table, taken) -> list:
    """'class 8', 'section A', 'room 2': a column's name followed by a short value of that column."""
    out = []
    low = text or ""
    for i, h in enumerate(t.header):
        if not h or i in taken or t.kind(i) == "date":
            continue
        cells = {}
        for v in t.col(i)[:5000]:
            if v not in (None, ""):
                cells.setdefault(str(v).strip().lower(), v)
            if len(cells) > 60:
                break
        if len(cells) > 60:
            continue
        names = set(_header_words(h))
        for x in list(names):
            for syn in SYNONYMS.values():
                if x in syn:
                    names |= {s for s in syn if " " not in s}
        for w in names:
            if len(w) < 3:
                continue
            m = re.search(rf"\b{re.escape(w)}\w*\s+([A-Za-z0-9]{{1,4}})\b", low, re.I)
            if m and m.group(1).lower() in cells:
                v = cells[m.group(1).lower()]
                out.append((i, v if isinstance(v, (int, float)) else str(v).strip()))
                break
    return out


def main_amount(t: Table):
    """The column a person means by 'sales', 'total', 'top bills': (column, derive-or-None).
    A unit price with a quantity next to it means quantity x price (worked out)."""
    from app.office import understand as U

    r = U.roles(t)
    nums = [i for i in _num_cols_of(t) if r.get(i) in ("amount", None)]
    if not nums:
        return None, None
    tot = [i for i in nums if re.search(r"amount|total|net|value|revenue|bill|fee|paid|salary|sales\b|sale\b|income", t.header[i], re.I)
           and not re.search(r"price|rate|cost|%", t.header[i], re.I)]
    if tot:
        return max(tot, key=lambda i: (bool(re.search(r"net|total", t.header[i], re.I)), -i)), None
    price = [i for i in nums if re.search(r"selling|sale|price|rate|mrp", t.header[i], re.I) and not re.search(r"cost|purchase|buy", t.header[i], re.I)]
    qty = [i for i, role in r.items() if role == "quantity"]
    if price and qty:
        return price[0], {"how": "line_total", "table": t.name, "arg": price[0]}
    return nums[0], None


def _agg_of(text: str) -> str | None:
    low = " " + (text or "").lower() + " "
    ws = set(tokens(text))
    has = lambda name: any((" " in w and f" {w} " in low) or w in ws for w in AGG_WORDS[name])  # noqa: E731
    if has("sum") and has("count") and not has("avg") and not re.search(r"\b(sabse|highest|lowest|max|min|top)\b", low):
        # "total sales kitni hai" is a total; "total orders kitne hain" is a count
        return "count" if re.search(r"\b(kitne|how many|count|ginti)\b", low) and not re.search(
            r"\b(sales|amount|salary|revenue|value|price|paisa|tax|bikri)\b", low) else "sum"
    for name in ("avg", "max", "min", "count", "sum"):
        for w in AGG_WORDS[name]:
            if (" " in w and f" {w} " in low) or w in ws:
                if name == "max" and w in ("zyada",) and "sabse" not in low and "top" not in low:
                    continue
                if name == "min" and w == "kam" and "sabse" not in low:
                    continue
                return name
    return None


# ------------------------------------------------------------------ understanding the question
class Skip(Exception):
    """This job doesn't fit after a closer look; try the next one."""


JOBS = [
    ("review", re.compile(r"\b(review|audit|jaanch|janch)\b|\b(check|dekho|dekh lo)\b.{0,15}\b(file|sheet|workbook|poori|puri)\b|\b(file|sheet)\b.{0,10}\b(check|review)\b", re.I)),
    ("reconcile", re.compile(r"\b(reconcil\w*|bank statement|milan|ledger\b.{0,25}\bbank|bank\b.{0,25}\bledger|books?\b.{0,15}\bbank)\b", re.I)),
    ("growth", re.compile(r"\b(growth|badhat|badhotri|month on month|mom|pichle mahine se|increase|decrease|ghata|badha)\b|"
                          r"(?<!before )(?<!before you )(?<!don't )(?<!dont )(?<!any )(?<!without )\bchange\b(?!s? (anything|nothing|mat))", re.I)),
    ("running", re.compile(r"\b(running total|cumulative|ab tak ka total|progressive total)\b", re.I)),
    ("missing", re.compile(r"\b(khaali|khali|blank|empty|missing|nahi bhara|bhara nahi|null)\b", re.I)),
    ("sort", re.compile(r"\b(sort|arrange|kram|ascending|descending|chhote se bade|bade se chhote|a to z|z to a|order mein laga)\b", re.I)),
    ("variance", re.compile(r"\b(variance|vs|versus|antar|farak|difference|budget)\b", re.I)),
    ("calc", re.compile(r"(\b(multiply|multiplied|guna|into|times|minus|divide|bhag|ratio|plus|full name|naam jodo|naam milao|line total|commission|gst amount|tax amount nikalo)\b|[×*/+]|\s x\s|\d+\s*%)", re.I)),
    ("aging", re.compile(r"\b(aging|ageing|overdue|kitne din se|days? (pending|unpaid|outstanding)|purane bill|bucket)\b", re.I)),
    ("lookup", re.compile(r"\b(lookup|look up|vlookup|xlookup|fetch|le ?aao|lao|laao|bring|pull|match karke|nikaal ke daalo)\b", re.I)),
    ("combine", re.compile(r"\b(combine|consolidate|merge|jodo|ek sheet|one sheet|stack|vstack|ek saath)\b.*\bsheets?\b|\bsheets?\b.*\b(combine|consolidate|merge|jodo|ek saath)\b", re.I)),
    ("highlight", re.compile(r"\b(highlight|rang|colou?r|mark|laal|red|yellow|peela)\b", re.I)),
    ("describe", re.compile(r"\b(kya hai is|kya kya hai|file mein kya|sheet mein kya|what.{0,10}(in|inside)|columns?|structure|samjhao|explain|kitni rows|how many rows|overview|summary of)\b", re.I)),
]


SHOW_RE = re.compile(r"\b(give me|show me|show|open|display|send me|de do|dedo|dikhao|dikha do|dikhado|chahiye|bhejo|view|see)\b|दिखाओ|दो\b", re.I)


def understand(text: str, tables: list[Table]) -> dict:
    """What the question asks, read against what the file really has. Hindi script is read in Roman letters
    (the way the file's headings are matched), with the Excel pack's word list for common Hindi words."""
    if re.search(r"[\u0900-\u097f]", text or ""):
        from app.mind import packs
        from app.ml.normalize import transliterate

        words = packs.get("office").get("romanised_words", {})
        roman = " ".join(words.get(w, w) for w in re.findall(r"[\w%.]+", transliterate(text).lower()))
        try:
            return _understand_en(roman, tables)
        except NeedMore:
            return _understand_en(text, tables)
    return _understand_en(text, tables)


def _understand_en(text: str, tables: list[Table]) -> dict:
    """What the question asks, read against what the file really has (app.office.understand)."""
    from app.office import understand as U

    if not tables:
        raise NeedMore("Is file mein padhne layak koi sheet nahi mili.", "No readable sheet in this file.")
    from app.office import edits

    if edits.MACRO_RE.search(text or ""):
        raise NeedMore("Macro / VBA likhna aur chalana ye agent nahi karta (file mein code chalana safe nahi). Jo kaam macro se karna tha, "
                       "wo seedha bolo: jaise \"blank rows hatao\", \"region wise alag sheets bana do\", \"header bold karo\".",
                       "This agent doesn't write or run macros / VBA (running code inside your file isn't safe). Tell me the job the macro "
                       "was for and I'll do it directly, e.g. \"remove blank rows\", \"one sheet per region\", \"make the header bold\".")
    order = edits.plan_edit(text, tables)                 # "make the header bold", "remove blank rows": do it on the file
    if order:
        return order
    if edits.PIVOT_RE.search(text or ""):
        plan = _understand_core_or_show(text, tables)
        if plan.get("job") == "aggregate" and plan.get("group") is not None:
            plan["pivot"] = True
        return plan
    show = _plan_show(text, tables)
    if show:
        return show
    named = named_sheet(text, tables) or pick_sheet(text, tables)
    for concept, vals in U.mentioned(text, tables).items():
        c = U._concepts()[concept]
        has_values = bool(c.get("values"))
        found = U.column_for(concept, [named] if named else tables) or (U.column_for(concept, tables) if named else None)
        if found and has_values:
            return _concept_plan(text, tables, found[0], found[1], concept, vals)
        if found:
            continue                                   # e.g. 'salary' and the file has a salary column: the usual reading
        # the question names a real column ('average science marks'): the concept word just describes it
        plain = _swap_words(text, c.get("words", []), " ")
        if not has_values and any(sc >= 1 for t in ([named] if named else tables) for i, sc in columns_in(plain, t, {"number"})
                                  if not re.search(r"\b(id|no|code)\b", t.header[i], re.I)):
            return _understand_core(plain, tables)
        how = c.get("derive")
        cands = [t for t in tables if how and U.can_derive(how, t)]
        if cands:
            t = named if named in cands else max(cands, key=lambda x: len(x.rows))
            t2, info = U.derive(how, t)
            tables2 = [t2 if x is t else x for x in tables]
            if has_values:
                plan = _concept_plan(text, tables2, t2, len(t.header), concept, vals)
            else:
                # 'kitna munafa hua' -> the worked-out column by its own name
                said = re.sub(r"\bolder than\b", f"{info['header']} more than", text, flags=re.I)
                said = re.sub(r"\byounger than\b", f"{info['header']} less than", said, flags=re.I)
                said = _swap_words(said, c.get("words", []), info["header"])
                plan = _understand_core(said, [t2])
            plan["derive"] = {"how": how, "table": t.name}
            return plan
        raise U.missing_message(concept, tables)
    return _understand_core(text, tables)


def _understand_core_or_show(text, tables):
    return _plan_show(text, tables) or _understand_core(text, tables)


def _swap_words(text: str, words, by: str) -> str:
    for w in sorted(words, key=len, reverse=True):
        text = re.sub(rf"(?<![\w]){re.escape(w)}(?![\w])", by, text or "", flags=re.I)
    return text


def _plan_show(text, tables):
    """'give me employee data', 'Sales Data dikhao': that sheet itself."""
    if not SHOW_RE.search(text or ""):
        return None
    t = named_sheet(text, tables) or pick_sheet(text, tables)
    if t is None and len(tables) == 1 and re.search(r"\b(data|sheet|file|all|sab|poora|puri|everything)\b|" + re.escape(tables[0].name.lower()),
                                                    (text or "").lower()):
        t = tables[0]
    if t is None:
        return None
    rest = _without_sheet_name(text, t)
    if _agg_of(rest) or ROWS_Q.search(rest) or GROUP_RE.search(rest) or values_in(rest, t) or TOPN_RE.search(rest) or \
            any(sc >= 1 for _, sc in columns_in(rest, t)):
        return None
    return {"job": "show", "table": t.name}


def _concept_plan(text, tables, t, col, concept, vals):
    """'how many girls and how many boys' -> count by the gender column; 'average salary of women' -> filter on it."""
    from app.office import understand as U

    c = U._concepts()[concept]
    stripped = text or ""
    for w in sorted(c.get("words", []) + [v for vs in c.get("values", {}).values() for v in vs if len(v) > 1], key=len, reverse=True):
        stripped = re.sub(rf"(?<![\w]){re.escape(w)}(?![\w])", " ", stripped, flags=re.I)
    agg = _agg_of(stripped) or "count"
    nums = columns_in(stripped, t, {"number"})
    measure = nums[0][0] if nums and agg != "count" else None
    if agg in ("sum", "avg", "max", "min") and measure is None:
        agg = "count"
    cmp_f, stripped = _compare_filters(stripped, t)
    date_f, stripped = _date_filters(stripped, t)
    filters = [(i, v) for i, v in values_in(stripped, t) if i != col]
    filters += [f for f in _short_values(stripped, t, {i for i, _ in filters} | {col})] + cmp_f + date_f
    group = None
    if len(vals) == 1:
        canon = next(iter(vals))
        cells = U.cell_values(t, col, concept, canon) or [canon]
        filters.append((col, cells[0]))
    else:
        group = col                                     # both named, or 'gender wise': count every value
    if group is None and re.search(r"\b(list|show|names?|who|kaun|kaunse|dikhao|which)\b", stripped, re.I) and not _agg_of(stripped):
        return {"job": "list", "table": t.name, "filters": filters, "measure": None, "topn": None}    # 'list of girls in class 8'
    return {"job": "aggregate", "table": t.name, "agg": agg, "measure": measure, "group": group, "filters": filters, "topn": None}


def _understand_core(text: str, tables: list[Table]) -> dict:
    if not tables:
        raise NeedMore("Is file mein padhne layak koi sheet nahi mili.", "No readable sheet in this file.")
    for job, rx in JOBS:
        if rx.search(text or ""):
            if job == "describe" and _agg_of(text) and not re.search(r"\b(columns?|structure|samjhao|explain|overview)\b", text, re.I):
                continue
            if job == "describe" and ROWS_Q.search(text or "") and not re.search(r"\b(columns?|structure|samjhao|explain|overview)\b", text, re.I):
                continue                                   # "how many rows in X": the count, not a description
            try:
                return globals()[f"_plan_{job}"](text, tables)
            except Skip:
                continue
    return _plan_question(text, tables)


SHEET_HINTS = {"sales": {"bika", "bike", "biki", "bikti", "sold", "bikri", "becha", "orders", "order"},
               "purchase": {"kharida", "khareeda", "bought", "purchase", "vendor", "vendors", "supplier"},
               "employee": {"employee", "employees", "staff", "karmchari", "salary", "tankhwah", "log", "people", "log", "joining"},
               "inventory": {"stock", "inventory", "reorder"}}


def _choose_table(text, tables, need_kinds=None) -> Table:
    named = named_sheet(text, tables)
    if named:
        return named
    t = pick_sheet(text, tables)
    if t and _value_not_sheet(text, t, tables):
        t = _value_not_sheet(text, t, tables)
    if t:
        if not any(sc >= 1 for _, sc in columns_in(text, t, {"number"})):
            # "kis customer ka taxable value": 'customer' names the Customers sheet, but the value lives in Invoices
            other = [(max((sc for _, sc in columns_in(text, o, {"number"})), default=0), o) for o in tables if o is not t]
            other = [x for x in other if x[0] >= 1]
            if other and _agg_of(text):
                return max(other, key=lambda x: x[0])[1]
        return t
    words = set(tokens(text))
    scored = []
    for t in tables:
        cols = columns_in(text, t, need_kinds)
        hint = sum(1 for k, hw in SHEET_HINTS.items() if k in t.name.lower() and words & hw)
        scored.append((sum(s for _, s in cols) + 0.5 * len(values_in(text, t)) + hint, len(t.rows), t))
    scored.sort(key=lambda x: (-x[0], -x[1]))
    if scored and scored[0][0] > 0 and (len(scored) == 1 or scored[0][0] > scored[1][0]):
        return scored[0][2]
    if scored and scored[0][0] > 0 and scored[0][1] >= 2 * max(scored[1][1], 1):
        return scored[0][2]                               # same score: the main sheet (many more rows), not a sample tab
    data = [t for t in tables if len(t.rows) >= 3]
    if len(data) == 1:
        return data[0]
    if RECENT_RE.search(text or "") or OLDEST_RE.search(text or "") or YEAR_RE.search(text or "") or MONTH_NAME_RE.search(text or ""):
        # 'who joined most recently': the sheet whose date column is the one meant ('join' -> Joining Date)
        low = (text or "").lower()
        hits = [t for t in data if any(any(len(w) >= 4 and re.search(rf"\b{re.escape(w[:4])}", low) for w in _header_words(t.header[i]))
                                       for i in _date_cols(t))]
        pool = hits or [t for t in data if _date_cols(t)]
        if len(pool) == 1 or (pool and len(hits) >= 1):
            return max(pool, key=lambda t: len(t.rows))
    raise NeedMore("Kis sheet mein dekhun?", "Which sheet should I look in?", [t.name for t in data][:8])


def _without_sheet_name(text, t):
    """Words that name the sheet aren't values to filter on: 'the sales sheet' is not Department = Sales."""
    for w in re.findall(r"[a-z0-9]+", t.name.lower()):
        if w not in ("data",):
            text = re.sub(rf"\b{re.escape(w)}\w*(\s+(data|sheet|tab))?\b", " ", text, flags=re.I)
    return re.sub(r"\b(sheet|tab|worksheet)s?\b", " ", text, flags=re.I)


HAVE_RE = re.compile(r"\b(have|has|having|got|with|filled)\b|\b(hai|hain|bhara|bhari|diya|di hai)\s*[?.!]*\s*$", re.I)
NOT_HAVE_RE = re.compile(r"\b(without|missing|nahi hai|nahi diya|nahi bhara|bina|blank|empty|khaali|khali)\b|\bno\s+\w+\s*$", re.I)


def _filled_filters(text: str, t: Table, taken) -> list:
    """'how many have a mobile number', 'kitne logon ka email hai' -> rows where it's filled in; 'without email' -> empty."""
    from app.office import understand as U

    out = []
    if not HAVE_RE.search(text or "") and not NOT_HAVE_RE.search(text or ""):
        return out
    low = (text or "").lower()
    for i, role in U.roles(t).items():
        if i in taken or role not in ("phone", "email", "text", "name"):
            continue
        words = [w for w in _header_words(t.header[i]) if len(w) >= 4 and w not in ("name", "naam")]
        if role in ("name", "text") and not re.search(r"\b(name|naam)\b", t.header[i], re.I):
            continue
        if words and any(re.search(rf"\b(have|has|having|got|with|without|no|ka|ki|ke)\s+(an?\s+|the\s+)?{re.escape(w)}", low) for w in words):
            out.append((i, ("blank",) if NOT_HAVE_RE.search(text or "") else ("filled",)))
    return out


def _group_col(text: str, t: Table, filters):
    """'doctor wise', 'by section', 'paid by ke hisaab se', 'room type wise': the column to group by, or None.
    'with paid by Ramesh' is a value of a column, not grouping."""
    from app.office import understand as U

    gr = U.roles(t)
    low = text or ""
    taken_vals = {str(v).lower() for _, v in filters if isinstance(v, str)}
    best = None
    for m in GROUP_RE.finditer(low):
        gword = next(g for g in m.groups() if g)
        if gword.lower() in taken_vals:
            continue
        pos = low.lower().find(gword.lower(), m.start())
        before = re.search(r"(\w+)\s*$", low[:pos])
        tries = ([f"{before.group(1)} {gword}"] if before else []) + [gword]
        for words in tries:
            ws = [w for w in words.lower().split() if w not in ("by", "per", "each", "har", "ke", "wise", "of", "the")] or words.lower().split()
            cands = [(sc, i) for i, h in enumerate(t.header) if h and (t.kind(i) == "text" or gr.get(i) == "category")
                     and gr.get(i) not in ("name", "id") and (sc := _col_score(h, ws)) >= 0.5]
            exact_elsewhere = any(set(ws) & set(_header_words(h)) for h in t.header if h)
            if exact_elsewhere:
                # a word that is literally another column's name ('net' -> Net Sales) isn't a synonym for this one
                cands = [(sc, i) for sc, i in cands if set(ws) & set(_header_words(t.header[i]))]
            if not cands:
                continue
            sc, i = max(cands)
            after = low[m.end():m.end() + 40].strip().lower()
            vals = {str(v).strip().lower() for v in t.col(i)[:3000] if v not in (None, "")}
            if any(after.startswith(v) for v in vals if v):
                break                                      # 'paid by Ramesh': that's a filter
            if best is None or sc > best[0]:
                best = (sc, i)
            break
    return best[1] if best else None


def _plan_order(text, t: Table, filters):
    """'who joined most recently', 'latest 5 orders', 'oldest patient': rows ordered by a date (or by age for people)."""
    rec, old = RECENT_RE.search(text or ""), OLDEST_RE.search(text or "")
    if not (rec or old):
        return None
    topn = TOPN_RE.search(text or "")
    n = int(next(g for g in topn.groups() if g)) if topn else (int(m.group(1)) if (m := re.search(r"\b(\d{1,3})\b", text or "")) else 1)
    age = next((i for i in _num_cols_of(t) if re.search(r"\bage\b|umar", t.header[i], re.I)), None)
    word = (rec or old).group(1).lower()
    if age is not None and word in ("oldest", "eldest", "youngest"):
        return {"job": "list", "table": t.name, "filters": filters, "measure": age, "topn": n, "order": "asc" if word == "youngest" else "desc"}
    col = _pick_date_col(text, t, (rec or old).start())
    if col is None:
        return None
    return {"job": "list", "table": t.name, "filters": filters, "measure": col, "topn": n, "order": "desc" if rec else "asc"}


def _plan_rows(text, tables):
    """'how many records in the sales sheet' -> the number of filled rows of that sheet (header and blank rows not counted)."""
    if not ROWS_Q.search(text or ""):
        return None
    t = named_sheet(text, tables) or pick_sheet(text, tables)
    if t is not None and not named_sheet(text, tables) and _value_not_sheet(text, t, tables):
        return None                                    # "records of Sales department": a filtered count
    if t is not None:
        rest = _without_sheet_name(text, t)
        if values_in(rest, t) or GROUP_RE.search(rest) or _number_filters(rest, t, columns_in(rest, t, {"number"})):
            return None                                # "how many rows of East region": a filtered count, not the sheet size
        return {"job": "rows", "table": t.name}
    if any(values_in(text, o) for o in tables) or GROUP_RE.search(text or ""):
        return None
    return {"job": "rows", "table": None}             # no sheet named: every sheet's count


SHEETS_Q = re.compile(r"\b(how many|kitni|kitne|number of|count( of)?|total)\b.{0,15}\b(sheets|tabs|worksheets)\b|"
                      r"\b(sheets|tabs)\b.{0,15}\b(kitni|kitne|how many|count)\b|शीट.{0,10}कितन|कितनी शीट|"
                      r"\b(sheets?|tabs?|worksheets?)\s*(names?|ke naam|ka naam|list)\b|\b(names?|list)\s+(of\s+)?(all\s+)?(the\s+)?(sheets|tabs|worksheets)\b|"
                      r"\b(which|what|kaun ?si|kaunsi|konsi|kya kya)\s+(sheets|tabs)\b|\b(sheets|tabs)\s+(kaun ?si|kaunsi|konsi|kya kya)\b|"
                      r"^\s*(sheets|tabs|all sheets|saari sheets)\s*\??\s*$|शीट.{0,6}नाम", re.I)


def _plan_sheets(text, tables):
    if not SHEETS_Q.search(text or ""):
        return None
    return {"job": "sheets"}


def _run_sheets(p, tables, wb, hi, n):
    ts = list(tables.values())
    names = ", ".join(f"{t.name} ({_fmt(len(t.rows))})" for t in ts)
    say = (f"Is file mein {len(ts)} sheets hain: {names}" if hi else f"This file has {len(ts)} sheets: {names}")
    return {"say": say, "formula": "", "sheet": None, "value": len(ts), "value_label": "sheets", "where": "this file",
            "sheet_links": [t.name for t in ts if t.rows],
            "chart": {"title": "Har sheet mein records" if hi else "Records in each sheet", "rows": [[t.name, len(t.rows)] for t in ts]}}


_ALL_TABLES = __import__("contextvars").ContextVar("all_tables", default=None)


def _plan_question(text, tables) -> dict:
    if len(tables) > 1 or _ALL_TABLES.get() is None:
        _ALL_TABLES.set(tables)
    sheets_plan = _plan_sheets(text, tables)
    if sheets_plan:
        return sheets_plan
    rows_plan = _plan_rows(text, tables)
    if rows_plan:
        return rows_plan
    agg = _agg_of(text)
    if agg is None and not pick_sheet(text, tables) and not MONEY_WORDS.search(text or "") and not MONTH_RE.search(text or "") \
            and not RECENT_RE.search(text or "") and not OLDEST_RE.search(text or "") and not CMP_RE.search(text or "") \
            and not any(columns_in(text, t) or values_in(text, t) for t in tables):
        # nothing in the question points at this file: say so plainly, don't ask which sheet
        names = ", ".join(t.name for t in tables if len(t.rows) >= 2)
        raise NeedMore(f"Ye poori tarah samajh nahi aaya. Is file mein ye sheets hain: {names}. Sheet ya column ke naam ke saath likho, jaise: "
                       "\"region wise total sales\", \"pending wale dikhao\", \"amount ka 5% nikalo\".",
                       f"I didn't fully get that. This file has: {names}. Ask using a sheet or column name, e.g. \"total sales by region\", "
                       "\"show pending rows\", \"5% of amount\".", [])
    t = _choose_table(text, tables)
    if named_sheet(text, tables) is t:
        text = _without_sheet_name(text, t)
    cmp_f, text = _compare_filters(text, t)                # 'above 60', '50k se zyada'
    date_f, text = _date_filters(text, t)                  # 'in 2024', 'march mein'
    special = _plan_order(text, t, cmp_f + date_f)
    if special:
        return special
    from app.office import understand as U

    _r = U.roles(t)
    num_cols = [(i, sc) for i, sc in columns_in(text, t, {"number"}) if _r.get(i) not in ("id", "phone", "category")]
    txt_cols = columns_in(text, t, {"text"})
    date_cols = [i for i in range(len(t.header)) if t.kind(i) == "date"]
    filters = values_in(text, t)
    filters += _short_values(text, t, {i for i, _ in filters})
    filters += cmp_f + date_f
    filters += _filled_filters(text, t, {i for i, _ in filters})
    group = _group_col(text, t, filters)
    named_group = group
    month_text = text or ""
    for i, _ in num_cols:                                  # "Monthly Salary" is a column, not "month-wise"
        month_text = re.sub(re.escape(t.header[i]), " ", month_text, flags=re.I)
    strong_txt = [i for i, sc in txt_cols if sc >= 0.5 and i not in [f for f, _ in filters]]
    if group is not None and not isinstance(group, tuple) and len(strong_txt) >= 2 and re.search(
            r"\b(aur|and|&|x|vs)\b", text or "", re.I) and re.search(r"\b(wise|har|per|by|hisaab)\b", text or "", re.I):
        low = (text or "").lower()
        pos = lambda i: min((low.find(w) for w in _header_words(t.header[i]) if low.find(w) >= 0), default=999)  # noqa: E731
        group, other = sorted({group, *strong_txt[:2]} if group in strong_txt[:2] else {group, strong_txt[0]}, key=pos)[:2]
        msr = num_cols[0][0] if num_cols else None
        return {"job": "pivot", "table": t.name, "rows": group, "cols": other, "measure": msr,
                "agg": "count" if msr is None else (_agg_of(text) if _agg_of(text) in ("sum", "avg", "count") else "sum"), "filters": filters}
    for i, v in filters:                                  # 'Monthly ka total fee': Monthly is a plan here, not month-wise
        if isinstance(v, str) and MONTH_RE.search(v):
            month_text = re.sub(re.escape(v), " ", month_text, flags=re.I)
    if group is None and MONTH_RE.search(month_text) and date_cols:
        group = ("month", (columns_in(text, t, {"date"}) or [(date_cols[0], 1)])[0][0])
    if group is None and agg and txt_cols and re.search(r"\b(wise|har|per|each|hisaab|according)\b", text or "", re.I):
        group = txt_cols[0][0]
    if group is not None and not isinstance(group, tuple):
        vals = [v for v in t.col(group) if v not in (None, "")]
        if vals and len({str(v).strip().lower() for v in vals}) > 0.95 * len(vals) and not (group == named_group and len(vals) <= 30):
            group = None                                   # "top 5 orders": each order is its own row, so rank rows
                                                           # (a small table asked 'by region' keeps its one row per region)
    filters = [(i, v) for i, v in filters if i != group] + _number_filters(text, t, num_cols)
    seen_f, uniq = set(), []
    for f in filters:
        key = (f[0], str(f[1]).lower())
        if key not in seen_f:
            seen_f.add(key)
            uniq.append(f)
    filters = uniq
    num_cols = [(i, sc) for i, sc in num_cols if i not in [f for f, _ in filters]]
    if group is not None and not isinstance(group, tuple) and re.search(r"\bby\s+\w+", text or "", re.I) and \
            any(i == group for i, _ in num_cols):
        group = None
    if group is None and _agg_of(text) in ("max", "min"):
        # "which vendor has the highest invoice total": per vendor, then the top
        strong = [i for i, sc in txt_cols if sc >= 1 and i not in [f for f, _ in filters]]
        wm0 = re.search(r"\b(?:which|kaun|kaunsa|kaunse|kis|kisne|kiska|kiski)\s+(\w+)", text or "", re.I)
        if wm0:
            strong += [i for i, sc in txt_cols if 0.5 <= sc < 1 and _col_score(t.header[i], [wm0.group(1).lower()]) >= 0.5]
        wm = re.search(r"\b(?:which|kaun|kaunsa|kaunse|kis|kisne|kiska|kiski)\s+(\w+)", text or "", re.I)
        if wm:
            named = [i for i in strong if _col_score(t.header[i], [wm.group(1).lower()]) >= 0.5]
            strong = named + [i for i in strong if i not in named]
        msr = num_cols[0][0] if num_cols else None
        if msr is None and strong and wm:
            sold = set(tokens(text)) & {"bika", "bike", "biki", "sold", "becha", "units", "pieces", "quantity", "qty", "sells", "sell", "selling",
                                        "bikta", "bikti", "bikte", "demand"}
            pat = r"qty|quantity|units" if sold else r"amount|total|net|sales|value"
            msr = next((i for i in range(len(t.header)) if t.kind(i) == "number" and re.search(pat, t.header[i], re.I)), None)
        if strong and wm and msr is not None:
            return {"job": "aggregate", "table": t.name, "agg": "sum", "measure": msr, "group": strong[0],
                    "filters": filters, "topn": 3, "ranked": _agg_of(text)}
    topn = TOPN_RE.search(text or "")
    topn = int(next(g for g in topn.groups() if g)) if topn else None
    measure = num_cols[0][0] if num_cols else None
    derive = None
    if agg == "count" and num_cols and num_cols[0][1] >= 1 and not re.search(r"\b(records?|rows?|entries|people|log|employees|students|patients|"
                                                                               r"customers|orders|bills|bookings|invoices|times|baar)\b", text or "", re.I):
        agg = "sum"                                        # 'how many nights in total', 'kitne rooms hain': add the column up
    if measure is None and (MONEY_WORDS.search(text or "") or (agg in ("sum", "max", "min") or topn) and agg != "avg"):
        if MONEY_WORDS.search(text or "") or agg in ("sum", "max", "min") or topn:
            m_col, derive = main_amount(t)
            if m_col is not None and (MONEY_WORDS.search(text or "") or agg in ("sum", "max", "min") or topn):
                measure = m_col if not derive else len(t.header)
    elif measure is not None and agg in ("sum", None) and re.search(r"price|rate|mrp", t.header[measure], re.I):
        m_col, d = main_amount(t)
        if d is not None and d["arg"] == measure:
            derive, measure = d, len(t.header)             # 'total sales' with Qty and a unit price: Qty x price
    best = BEST_RE.search(text or "")
    if group is None and best and agg in (None, "max", "min"):
        from app.office import understand as U

        _gr = U.roles(t)
        gcol = next((i for i, sc in columns_in(best.group(1), t, {"text"}) if sc >= 0.5 and _gr.get(i) == "category"
                     and not any(sc2 >= sc for _, sc2 in columns_in(best.group(1), t, {"number"}))), None)
        if gcol is not None:
            rated = re.search(r"\brated|rating|stars\b", text or "", re.I)
            rcol = next((i for i in _num_cols_of(t) if re.search(r"rating|stars|score", t.header[i], re.I)), None) if rated else None
            if rated and rcol is None and not getattr(t, "_tried_rated", False):
                for other in _ALL_TABLES.get() or []:
                    if other is not t and any(re.search(r"rating|stars|score", h or "", re.I) for h in other.header) and \
                            any(sc >= 0.5 for _, sc in columns_in(best.group(1), other, {"text"})):
                        other._tried_rated = True
                        try:
                            return _plan_question(text, [other])
                        finally:
                            other._tried_rated = False
            msr = rcol if rcol is not None else measure
            if msr is None:
                msr, derive = main_amount(t)
                msr = msr if not derive else len(t.header)
            worst = re.search(r"\b(worst|lowest|weakest)\b", text or "", re.I)
            plan = {"job": "aggregate", "table": t.name, "agg": "avg" if rcol is not None else ("sum" if msr is not None else "count"),
                    "measure": msr, "group": gcol, "filters": filters, "topn": 3, "ranked": "min" if worst else "max"}
            if derive and msr == len(t.header):
                plan["derive"] = derive
            return plan
    sold = set(tokens(text)) & {"bika", "bike", "biki", "sold", "becha", "units", "pieces", "nag"}
    if agg == "count" and sold:
        qty = next((i for i in range(len(t.header)) if t.kind(i) == "number" and re.search(r"qty|quantity|units", t.header[i], re.I)), None)
        if qty is not None:
            agg, measure = "sum", qty                    # "laptop kitne bike": units sold, not number of orders
    if agg is None:
        if topn and measure is not None:
            agg = "max"
        elif filters or re.search(r"\b(list|dikhao|show|kaun|kaunse|which|jinka|jinki|wale|wali)\b", text or "", re.I):
            return {"job": "list", "table": t.name, "filters": filters, "measure": measure, "topn": topn}
        elif group is not None and measure is not None and num_cols and (num_cols[0][1] >= 1 or (
                num_cols[0][1] >= 0.5 and re.search(r"amount|total|net|sales|revenue|value|salary|fee|bill|paid|balance|profit|cost|price|tax|amt|\brs\b|rupee|expense|kharch|rent",
                                                    t.header[measure], re.I))):
            agg = "sum"                                    # 'region wise sales amount', 'salary by department': that money column, added up
        elif group is not None:
            agg = "count"                                  # 'doctor wise patients', 'month wise bookings': count per group
        else:
            cols = ", ".join(h for h in t.header if h)
            raise NeedMore(f"Ye poori tarah samajh nahi aaya. {t.name} ke columns: {cols}. Inke naam ke saath likho, jaise: "
                           "\"region wise total sales\", \"amount ka 5% nikalo\", \"status pending wale dikhao\".",
                           f"I didn't fully get that. {t.name} has: {cols}. Ask using these names, e.g. \"total sales by region\", "
                           "\"5% of amount\", \"show pending rows\".", [h for h in t.header if h][:8])
    if agg in ("sum", "avg", "max", "min") and measure is None:
        cands = [i for i in range(len(t.header)) if t.kind(i) == "number" and t.header[i] and i not in [f for f, _ in filters]]
        amountish = [i for i in cands if re.search(r"amount|total|net|sales|salary|price|value|fee|fees|charges|revenue|bill", t.header[i], re.I)]
        if len(cands) == 1:
            measure = cands[0]
        elif len(amountish) == 1 or (amountish and agg == "sum"):
            measure = amountish[-1] if agg == "sum" else amountish[0]
        if measure is None and agg != "avg":
            m_col, derive = main_amount(t)
            measure = None if m_col is None else (m_col if not derive else len(t.header))
        if measure is None:
            raise NeedMore(f"{t.name} mein kis column ka {_AGG_HI[agg]} chahiye?", f"Which column in {t.name} should I {_AGG_EN[agg]}?",
                           [t.header[i] for i in cands][:8])
    plan = {"job": "aggregate", "table": t.name, "agg": agg, "measure": measure, "group": group, "filters": filters, "topn": topn}
    if derive and measure == len(t.header):
        plan["derive"] = derive
    return plan


_AGG_HI = {"sum": "total", "avg": "average", "max": "sabse zyada", "min": "sabse kam", "count": "kitne"}
_AGG_EN = {"sum": "total", "avg": "average", "max": "find the highest of", "min": "find the lowest of", "count": "count"}


def _plan_describe(text, tables):
    t = pick_sheet(text, tables)
    return {"job": "describe", "table": t.name if t else None}


def _plan_aging(text, tables):
    t = _choose_table(text, tables)
    dates = columns_in(text, t, {"date"}) or [(i, 1) for i in range(len(t.header)) if t.kind(i) == "date"]
    if not dates:
        raise NeedMore(f"{t.name} mein koi date column nahi mila; aging ke liye bill/invoice ki date chahiye.",
                       f"{t.name} has no date column; aging needs the invoice date.")
    nums = columns_in(text, t, {"number"})
    amount = nums[0][0] if nums else next((i for i in range(len(t.header)) if t.kind(i) == "number"
                                          and re.search(r"total|amount|value|net", t.header[i], re.I)), None)
    return {"job": "aging", "table": t.name, "date": dates[0][0], "amount": amount, "filters": values_in(text, t)}


def _plan_lookup(text, tables):
    named = [t for t in tables if pick_sheet(t.name, [t]) and pick_sheet(text, [t])]
    if len(named) < 2:
        raise NeedMore("Kis sheet se kis sheet mein laana hai? Dono sheets ke naam likho, jaise: \"Purchase Data se unit price Sales Data mein laao\".",
                       "From which sheet into which? Name both, e.g. \"bring unit price from Purchase Data into Sales Data\".",
                       [t.name for t in tables][:8])
    low = (text or "").lower()
    pos = {t.name: min((low.find(w) for w in re.findall(r"[a-z0-9]+", t.name.lower()) if w not in ("data", "sheet") and low.find(w) >= 0), default=999)
           for t in named}
    m = re.search(r"\b(\w[\w ]*?)\s+se\b", low)
    src = None
    for t in named:
        key = [w for w in re.findall(r"[a-z0-9]+", t.name.lower()) if w not in ("data", "sheet")]
        if re.search(r"\b(from|se)\b", low) and any(re.search(rf"\b(from\s+(the\s+)?{k}|{k}\w*\s+(data\s+)?(sheet\s+)?se)\b", low) for k in key):
            src = t
    if src is None:
        src = sorted(named, key=lambda t: pos[t.name])[-1]
    dst = next(t for t in named if t is not src)
    want = [i for i, _ in columns_in(text, src) if not any(_col_score(dst.header[j], _header_words(src.header[i])) >= 1 for j in range(len(dst.header)) if dst.header[j])]
    if not want:
        raise NeedMore(f"{src.name} se kaunsa column laana hai?", f"Which column should I bring from {src.name}?", [h for h in src.header if h][:8])
    key = _join_key(dst, src)
    if key is None:
        raise NeedMore(f"{dst.name} aur {src.name} mein koi milta-julta column nahi mila (jaise Item Code). Kaunse column se milaun?",
                       f"No matching column between {dst.name} and {src.name} (like Item Code). Which column should I match on?",
                       [h for h in dst.header if h][:8])
    return {"job": "lookup", "table": dst.name, "source": src.name, "key": key, "want": want[0]}


def _join_key(dst: Table, src: Table):
    best, score = None, 0.0
    for i, h in enumerate(dst.header):
        if not h or dst.kind(i) in ("number", "date", "empty"):
            continue
        dvals = {str(v).strip().lower() for v in dst.col(i) if v not in (None, "")}
        if not dvals:
            continue
        for j, h2 in enumerate(src.header):
            if not h2 or src.kind(j) in ("date", "empty"):
                continue
            svals = {str(v).strip().lower() for v in src.col(j) if v not in (None, "")}
            overlap = len(dvals & svals) / len(dvals)
            if overlap > score:
                best, score = (i, j), overlap
    return best if score >= 0.5 else None


def _plan_combine(text, tables):
    named = [t for t in tables if pick_sheet(text, [t])]
    groups = defaultdict(list)
    for t in (named or tables):
        groups[tuple(h.strip().lower() for h in t.header if h)].append(t)
    same = max(groups.values(), key=len)
    if len(same) < 2:
        raise NeedMore("Jodne ke liye kam se kam 2 sheets mein same columns chahiye; is file mein aisi sheets nahi hain.",
                       "Combining needs at least 2 sheets with the same columns; this file has none.")
    return {"job": "combine", "tables": [t.name for t in same]}


_LESS = re.compile(r"\b(kam|less|below|under|neeche|chhota|chota|lower|<)\b|<", re.I)
_MORE = re.compile(r"\b(zyada|jyada|more|above|over|upar|bada|greater|higher|>)\b|>", re.I)


def _plan_highlight(text, tables):
    t = _choose_table(text, tables, {"number"})
    nums = columns_in(text, t, {"number"})
    filters = values_in(text, t)
    op = "<" if _LESS.search(text or "") else (">" if _MORE.search(text or "") else None)
    number = re.findall(r"(?<![\w.])(\d[\d,]*(?:\.\d+)?)\s*(k|lakh|lac|hazar|hazaar|thousand)?\b", text or "")
    if filters and not op:
        return {"job": "highlight", "table": t.name, "filters": filters}
    if op and nums and number:
        n, unit = number[-1]
        val = float(n.replace(",", "")) * {"k": 1e3, "lakh": 1e5, "lac": 1e5, "hazar": 1e3, "hazaar": 1e3, "thousand": 1e3}.get(unit.lower(), 1)
        return {"job": "highlight", "table": t.name, "left": nums[0][0], "op": op, "right_val": val}
    if op and len(nums) >= 2:
        return {"job": "highlight", "table": t.name, "left": nums[0][0], "op": op, "right_col": nums[1][0]}
    if op and nums and number:
        n, unit = number[-1]
        val = float(n.replace(",", "")) * {"k": 1e3, "lakh": 1e5, "lac": 1e5, "hazar": 1e3, "hazaar": 1e3, "thousand": 1e3}.get(unit.lower(), 1)
        return {"job": "highlight", "table": t.name, "left": nums[0][0], "op": op, "right_val": val}
    raise NeedMore("Kaunsi rows highlight karun? Jaise: \"jahan closing stock reorder level se kam hai\" ya \"net sales 10000 se zyada\".",
                   "Which rows should I highlight? E.g. \"where closing stock is below reorder level\" or \"net sales above 10000\".")


# ------------------------------------------------------------------ doing it
def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _match_filters(row, filters):
    def ok(i, v):
        x = row[i] if i < len(row) else None
        if isinstance(v, (list, tuple)) and v and v[0] in ("filled", "blank"):
            empty = x is None or (isinstance(x, str) and not x.strip())
            return empty if v[0] == "blank" else not empty
        if isinstance(v, (list, tuple)) and v and v[0] in ("cmp", "year", "month"):
            if v[0] == "year":
                return isinstance(x, (datetime, date)) and x.year == v[1]
            if v[0] == "month":
                return isinstance(x, (datetime, date)) and x.month == v[1] and (len(v) < 3 or v[2] is None or x.year == v[2])
            op, n = v[1], v[2]
            if not isinstance(x, (int, float)) or isinstance(x, bool):
                return False
            return {"gt": x > n, "ge": x >= n, "lt": x < n, "le": x <= n}[op]
        if isinstance(v, (int, float)):
            return isinstance(x, (int, float)) and abs(x - v) < 1e-9
        return isinstance(x, str) and x.strip().lower() == str(v).lower()
    return all(ok(i, v) for i, v in filters)


OPS = {"gt": ">", "ge": ">=", "lt": "<", "le": "<="}
MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"]


def _fdesc(t: Table, filters) -> str:
    out = []
    for i, v in filters:
        h = t.header[i] if i < len(t.header) else "?"
        if isinstance(v, (list, tuple)) and v and v[0] == "cmp":
            out.append(f"{h} {OPS[v[1]]} {_fmt(v[2])}")
        elif isinstance(v, (list, tuple)) and v and v[0] == "year":
            out.append(f"{h} in {v[1]}")
        elif isinstance(v, (list, tuple)) and v and v[0] in ("filled", "blank"):
            out.append(f"{h} {'filled in' if v[0] == 'filled' else 'empty'}")
        elif isinstance(v, (list, tuple)) and v and v[0] == "month":
            out.append(f"{h} in {MONTHS[v[1] - 1].title()}" + (f" {v[2]}" if len(v) > 2 and v[2] else ""))
        else:
            out.append(f"{h} = {v}")
    return ", ".join(out)


def _fmt(n):
    if n is None:
        return "-"
    frac, neg = "", False
    if isinstance(n, float) and not n.is_integer():
        neg = n < 0
        whole, _, dec = f"{abs(round(n, 2)):.2f}".partition(".")
        dec = dec.rstrip("0")
        frac = f".{dec}" if dec else ""
        n = -int(whole) if neg else int(whole)
    if isinstance(n, float):
        n = int(round(n))
    s = str(abs(n))
    if isinstance(n, int) and len(s) > 3:                 # Indian grouping: 12,34,567
        head, tail = s[:-3], s[-3:]
        head = ",".join(re.findall(r"\d{1,2}(?=(?:\d{2})*$)", head))
        s = f"{head},{tail}"
    return ("-" if n < 0 or neg else "") + s + frac


def _label(t: Table):
    """The column that names a row: a filled 'Name' column, else first + last name, else people's names by their
    values ('Guest', 'Driver'), else the first text column."""
    from app.office import understand as U

    people = [i for i, role in U.roles(t).items() if role == "name" and not re.search(r"estimated|\(from ", t.header[i], re.I)]
    if people and not any(re.search(r"name", t.header[i] or "", re.I) for i in people):
        return people[0]
    names = [i for i, h in enumerate(t.header) if h and re.search(r"name", h, re.I) and t.kind(i) == "text"
             and not re.search(r"estimated|\(from ", h, re.I)]          # a worked-out column isn't a person's name
    first = next((i for i in names if re.search(r"first", t.header[i], re.I)), None)
    last = next((i for i in names if re.search(r"last|sur", t.header[i], re.I)), None)
    full = [i for i in names if i not in (first, last)]
    if full:
        return full[0]
    if first is not None:
        return (first, last)
    return next((i for i in range(len(t.header)) if t.kind(i) == "text" and t.header[i]), 0)


def _name(t, r, label) -> str:
    if isinstance(label, tuple):
        return " ".join(str(r[i]) for i in label if i is not None and i < len(r) and r[i])
    return str(r[label]) if label < len(r) else ""


def _crit(t: Table, filters):
    out = []
    for i, v in filters:
        if isinstance(v, (list, tuple)) and v and v[0] == "cmp":
            out.append(f',{t.ref(i)},"{OPS[v[1]]}{v[2]}"')
        elif isinstance(v, (list, tuple)) and v and v[0] == "year":
            out.append(f',{t.ref(i)},">="&DATE({v[1]},1,1),{t.ref(i)},"<="&DATE({v[1]},12,31)')
        elif isinstance(v, (list, tuple)) and v and v[0] in ("filled", "blank"):
            out.append(f',{t.ref(i)},"{"<>" if v[0] == "filled" else ""}"')
        elif isinstance(v, (list, tuple)):
            continue
        else:
            out.append(f',{t.ref(i)},{v}' if isinstance(v, (int, float)) else f',{t.ref(i)},"{v}"')
    return "".join(out)


def run(plan: dict, path: str, out: str, hi: bool, n: int) -> dict:
    """Do the plan on a copy of the file (new sheet / new column / colours). Returns {'say', 'formula', 'sheet'}."""
    if plan.get("job", "").startswith("word_"):
        from app.office import word

        return word.run(plan, path, out, hi)
    tables = {t.name: t for t in load_tables(path)}
    info = None
    if plan.get("derive"):
        from app.office import understand as U

        d = plan["derive"]
        tables[d["table"]], info = U.derive(d["how"], tables[d["table"]], d.get("arg"))
        plan = {**plan, "readonly": True}               # an estimate: shown, not written into their file
    job = plan["job"]
    answer_only = job in ("rows", "sheets", "describe", "review", "show") or (
        job == "aggregate" and plan.get("group") is None and plan.get("agg") not in ("max", "min")) or (
        job == "list" and plan.get("topn") and plan["topn"] <= 10)        # a few rows: the answer, not a new sheet
    if plan.get("readonly") and job in ("aggregate", "list", "describe", "rows", "pivot", "growth", "running", "missing", "variance", "aging", "review"):
        answer_only = True                                 # "before making changes, tell me...": answer only, the file stays as it is
    # a single number or a description adds nothing to the file: don't open, rewrite or save the workbook for it
    wb = None if answer_only else (load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
                                   if path.lower().endswith((".xlsx", ".xlsm")) else None)
    plan = {**plan, "_path": path}
    if job == "edit" and wb is None:
        return {"say": ("Ye badlav .xlsx file mein hi ho sakta hai; file ko Excel mein .xlsx save karke dobara upload karo" if hi else
                        "That change needs an .xlsx file; save it as .xlsx in Excel and upload it again"), "formula": "", "sheet": None}
    res = globals()[f"_run_{job}"](plan, tables, wb, hi, n)
    if info is not None:
        res = {**res, **_derived_note(info, hi, res)}
    if plan.get("pivot"):
        res = {**res, "say": res["say"] + (". (Excel ka asli PivotTable yahan nahi banta; ye wahi numbers wali summary table hai, jo data badalne pe "
                                           "formulas se update hoti hai)" if hi else
                                           ". (An Excel PivotTable itself can't be built here; this is a summary table with the same numbers, "
                                           "and its formulas update when the data changes)")}
    if wb is not None:
        wb.calculation.fullCalcOnLoad = True
        wb.save(out)
    return res


def _run_edit(p, tables, wb, hi, n):
    from app.office import edits

    res = edits.run_edit(p, tables, wb, hi, n)
    if res["changes"]:
        res["say"] += (". Nayi copy card pe Download karo; 'undo' likhoge to pehle jaisa ho jayega" if hi else
                       ". Download the new copy from the card; type 'undo' to put it back")
    return {"formula": "", **res, "sheet": None}


def _sample(t: Table, rows, k=5):
    """The real rows behind an answer, so the person can see what was counted."""
    keep = [i for i, h in enumerate(t.header) if h][:8]
    return {"sheet": t.name, "header": [t.header[i] for i in keep],
            "rows": [[_cell(r[i]) if i < len(r) else "" for i in keep] for r in rows[:k]], "total": len(rows)}


def _cell(v):
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, date):
        return v.strftime("%d-%m-%Y")
    return "" if v is None else v


def _run_rows(p, tables, wb, hi, n):
    allt = [t for t in tables.values() if t.rows]
    chart_rows = [[t.name, len(t.rows)] for t in allt]
    if p.get("table"):
        t = tables[p["table"]]
        say = (f"{t.name} mein {_fmt(len(t.rows))} records hain (heading aur khaali rows nahi gine)" if hi else
               f"{t.name} has {_fmt(len(t.rows))} records (the heading row and blank rows aren't counted)")
        if t.truncated:
            say += " (sirf pehli 5 lakh rows padhi)" if hi else " (read the first 5 lakh rows only)"
        return {"say": say, "formula": f"=COUNTA({t.ref(0)})", "sheet": None, "value": len(t.rows), "value_label": "records",
                "where": t.name, "chart": {"title": "Records in each sheet" if not hi else "Har sheet mein records", "rows": chart_rows,
                                            "highlight": t.name}, "sample": _sample(t, t.rows)}
    total = sum(len(t.rows) for t in allt)
    say = (f"Poori file mein {_fmt(total)} records: " if hi else f"{_fmt(total)} records in the file: ") + "; ".join(f"{t.name}: {_fmt(len(t.rows))}" for t in allt)
    return {"say": say, "formula": "", "sheet": None, "value": total, "value_label": "records", "where": "all sheets",
            "sheet_links": [t.name for t in allt],
            "chart": {"title": "Records in each sheet" if not hi else "Har sheet mein records", "rows": chart_rows}}


def _derived_note(info: dict, hi: bool, res: dict) -> dict:
    """Say it's derived, from what, and which rows couldn't be told."""
    unclear = info.get("unclear") or []
    if info["how"] == "gender_from_names":
        head = (f"Is file mein gender ka column nahi hai, isliye ye '{info['from']}' ke naamon se andaza hai. " if hi else
                f"This file has no gender column, so this is estimated from the names in '{info['from']}'. ")
        tail = ""
        if unclear:
            shown = ", ".join(unclear[:8]) + (" …" if len(unclear) > 8 else "")
            tail = (f" {len(unclear)} naamon se pata nahi chala ({shown}); unhe 'Unclear' mein rakha." if hi else
                    f" {len(unclear)} names couldn't be told ({shown}); they're counted as 'Unclear'.")
        return {"say": head + res["say"] + "." + tail, "estimate": True, "unclear": unclear[:50]}
    head = (f"Ye '{info['from']}' se nikala: " if hi else f"Worked out from '{info['from']}': ")
    return {"say": head + res["say"], "estimate": True}


def _run_show(p, tables, wb, hi, n):
    t = tables[p["table"]]
    cols = ", ".join(h for h in t.header if h)
    say = (f"{t.name}: {_fmt(len(t.rows))} records, columns: {cols}" if hi else f"{t.name}: {_fmt(len(t.rows))} records. Columns: {cols}")
    return {"say": say, "formula": "", "sheet": None, "value": len(t.rows), "value_label": "records", "where": t.name,
            "sample": _sample(t, t.rows, 20)}


def _answer_sheet(wb, n, title):
    name = f"Answer {n} - {title}"[:31]
    name = re.sub(r"[\[\]:*?/\\]", "", name)
    ws = wb.create_sheet(name, 0)
    return ws, name


def _write_table(ws, header, rows, formula_note=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    if formula_note:
        ws.append([])
        ws.append(["Excel formula:", formula_note])
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(h)) + 4))


def _run_aggregate(p, tables, wb, hi, n):
    t = tables[p["table"]]
    agg, m, g, filters, topn = p["agg"], p["measure"], p["group"], p["filters"], p["topn"]
    rows = [r for r in t.rows if _match_filters(r, filters)]
    mname = t.header[m] if m is not None and agg != "count" else "records"
    fdesc = _fdesc(t, filters)

    def calc(vals):
        nums = [x for x in (_num(v) for v in vals) if x is not None]
        if agg == "count":
            return len(vals)
        if not nums:
            return None
        return {"sum": sum(nums), "avg": sum(nums) / len(nums), "max": max(nums), "min": min(nums)}[agg]

    crit = _crit(t, filters)
    fn = {"sum": "SUMIFS" if filters else "SUM", "avg": "AVERAGEIFS" if filters else "AVERAGE", "max": "MAXIFS" if filters else "MAX",
          "min": "MINIFS" if filters else "MIN", "count": "COUNTIFS" if filters else "COUNTA"}[agg]
    if g is None:
        if agg in ("max", "min") and m is not None and (topn or True):
            ranked = sorted([r for r in rows if _num(r[m]) is not None], key=lambda r: r[m], reverse=agg == "max")[: (topn or 1)]
            label = _label(t)
            val = ranked[0][m] if ranked else None
            formula = f"={fn}({t.ref(m)}{crit})"
            if ranked and topn:
                formula = f"=TAKE(SORT(FILTER({t.ref(0)[:t.ref(0).index('!') + 1]}A{t.header_row + 1}:{get_column_letter(len(t.header))}{t.header_row + len(t.rows)},{t.ref(m)}<>\"\"),{m + 1},{-1 if agg == 'max' else 1}),{topn})"
            if wb is not None:
                ws, sname = _answer_sheet(wb, n, f"{'Top' if agg == 'max' else 'Lowest'} {mname}")
                _write_table(ws, t.header, ranked, formula)
            who = ", ".join(f"{_name(t, r, label)} ({_fmt(r[m])})" for r in ranked[:5])
            chart = {"title": mname, "rows": [[_name(t, r, label), r[m]] for r in ranked[:10]]}
            say = (f"{t.name}{' (' + fdesc + ')' if fdesc else ''}: {_AGG_HI[agg]} {mname} = {_fmt(val)}. {who}" if hi else
                   f"{t.name}{' (' + fdesc + ')' if fdesc else ''}: {'highest' if agg == 'max' else 'lowest'} {mname} = {_fmt(val)}. {who}")
            return {"say": say, "formula": formula, "sheet": sname if wb is not None else None, "chart": chart}
        val = calc([r[m] for r in rows] if m is not None else rows)
        target = t.ref(m) if m is not None else t.ref(0)
        formula = f"={fn}({target}{crit})" if filters or agg != "count" else f"=COUNTA({t.ref(0)})"
        if agg == "count" and filters:
            formula = f"=COUNTIFS({crit[1:]})"
        chart = None
        if filters:
            # the same number for every value of the filtered column, so the answer is seen next to the rest
            fi, fv = filters[0]
            others = filters[1:]
            per = defaultdict(list)
            for r in t.rows:
                k = r[fi] if fi < len(r) else None
                if k not in (None, "") and _match_filters(r, others):
                    per[str(k).strip()].append(r)
            if 1 < len(per) <= 40:
                rows_c = sorted(((k, calc([x[m] for x in v] if m is not None else v)) for k, v in per.items()), key=lambda x: -(x[1] or 0))
                chart = {"title": f"{'Records' if agg == 'count' else _AGG_EN[agg].split()[-1].title() + ' ' + mname} by {t.header[fi]}",
                         "rows": [[k, v or 0] for k, v in rows_c[:12]], "highlight": str(fv)}
        what = f"{_AGG_HI[agg]} {mname}" if hi else f"{_AGG_EN[agg]} {mname}".replace("count rows", "records").replace("total Total", "Total")
        say = f"{t.name}{' (' + fdesc + ')' if fdesc else ''}: {what} = {_fmt(val)}"
        if agg == "count":
            say = f"{t.name}{' (' + fdesc + ')' if fdesc else ''}: {_fmt(val)} records"
        if t.truncated:
            say += " (sirf pehli 5 lakh rows)" if hi else " (first 5 lakh rows only)"
        return {"say": say, "formula": formula, "sheet": None, "value": val, "value_label": what, "where": t.name + (f" ({fdesc})" if fdesc else ""),
                "chart": chart, "sample": _sample(t, rows)}

    # X-wise / month-wise
    buckets = defaultdict(list)
    if isinstance(g, tuple):
        gname = "Month"
        for r in rows:
            d = r[g[1]] if g[1] < len(r) else None
            if isinstance(d, (datetime, date)):
                buckets[d.strftime("%Y-%m")].append(r[m] if m is not None else 1)
        keys = sorted(buckets)
        formula = (f'=SUMPRODUCT((TEXT({t.ref(g[1])},"yyyy-mm")=A2)*{t.ref(m)})' if agg == "sum" and m is not None else
                   f'=SUMPRODUCT(--(TEXT({t.ref(g[1])},"yyyy-mm")=A2))')
    else:
        gname = t.header[g]
        for r in rows:
            k = r[g] if g < len(r) else None
            if k not in (None, ""):
                buckets[str(k).strip()].append(r[m] if m is not None else 1)
        keys = list(buckets)
        fn2 = {"sum": "SUMIFS", "avg": "AVERAGEIFS", "max": "MAXIFS", "min": "MINIFS", "count": "COUNTIFS"}[agg]
        formula = (f"=UNIQUE({t.ref(g)}) in A2, then in B2: " +
                   (f"={fn2}({t.ref(m)},{t.ref(g)},A2{crit})" if agg != "count" else f"=COUNTIFS({t.ref(g)},A2{crit})"))
    table = [(k, calc(buckets[k])) for k in keys]
    if not isinstance(g, tuple):
        table.sort(key=lambda x: -(x[1] or 0))
    if topn:
        table = table[:topn]
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, f"{mname} by {gname}")
        _write_table(ws, [gname, f"{_AGG_EN[agg].split()[-1].title()} {mname}"], table, formula)
        _chart(ws, len(table), f"{mname} by {gname}", line=isinstance(g, tuple))
    if p.get("ranked"):
        table.sort(key=lambda x: (x[1] or 0), reverse=p["ranked"] == "max")
        table = table[:3]
        best = table[0] if table else ("-", None)
        say = (f"{t.name}: sabse {'zyada' if p['ranked'] == 'max' else 'kam'} {mname} {best[0]} ka hai ({_fmt(best[1])}). "
               + "; ".join(f"{k}: {_fmt(v)}" for k, v in table) if hi else
               f"{t.name}: {best[0]} has the {'highest' if p['ranked'] == 'max' else 'lowest'} {mname} ({_fmt(best[1])}). "
               + "; ".join(f"{k}: {_fmt(v)}" for k, v in table))
        if wb is not None:
            ws, sname = _answer_sheet(wb, n, f"{mname} by {gname}")
            _write_table(ws, [gname, f"Total {mname}"], table, formula)
        return {"say": say, "formula": formula, "sheet": sname if wb is not None else None,
                "chart": {"title": f"{mname} by {gname}", "rows": [[k, v] for k, v in table]}}
    lines = "; ".join(f"{k}: {_fmt(v)}" for k, v in table[:8])
    more = f" (+{len(table) - 8})" if len(table) > 8 else ""
    say = (f"{t.name}{' (' + fdesc + ')' if fdesc else ''}, {gname} wise {_AGG_HI[agg]} {mname}: {lines}{more}" if hi else
           f"{t.name}{' (' + fdesc + ')' if fdesc else ''}, {_AGG_EN[agg]} {mname} by {gname}: {lines}{more}")
    return {"say": say, "formula": formula, "sheet": sname if wb is not None else None,
            "chart": {"title": f"{mname} by {gname}", "rows": [[k, v] for k, v in table[:12]], "line": isinstance(g, tuple)}}


def _run_list(p, tables, wb, hi, n):
    t = tables[p["table"]]
    rows = [r for r in t.rows if _match_filters(r, p["filters"])]
    if p.get("measure") is not None and p.get("topn"):
        mi = p["measure"]
        ok = [r for r in rows if mi < len(r) and (_num(r[mi]) is not None or isinstance(r[mi], (datetime, date)))]
        rows = sorted(ok, key=lambda r: r[mi], reverse=p.get("order", "desc") == "desc")[: p["topn"]]
    fdesc = _fdesc(t, p["filters"]) or ("sab" if hi else "all")
    span = f"{t.ref(0)[:t.ref(0).index('!') + 1]}A{t.header_row + 1}:{get_column_letter(len(t.header))}{t.header_row + len(t.rows)}"
    cond = "*".join(f'({t.ref(i)}="{v}")' for i, v in p["filters"]) or "TRUE"
    formula = f"=FILTER({span},{cond})"
    sname = None
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, "List " + (_fdesc(t, p["filters"]) or t.name))
        _write_table(ws, t.header, rows, formula)
    label = _label(t)
    mi = p.get("measure")
    if mi is not None and p.get("topn"):
        # ordered: say the value it was ordered by ('Aman Verma (12-10-2025)')
        names = ", ".join(f"{_name(t, r, label)} ({_cell(r[mi]) if not isinstance(r[mi], (int, float)) else _fmt(r[mi])})" for r in rows[:10])
    else:
        names = ", ".join(_name(t, r, label) for r in rows[:10])
    say = (f"{t.name} ({fdesc}): {len(rows)} rows. {names}{' …' if len(rows) > 10 else ''}" if hi else
           f"{t.name} ({fdesc}): {len(rows)} rows. {names}{' …' if len(rows) > 10 else ''}")
    if mi is not None and p.get("order"):
        by = t.header[mi]
        first = "sabse naya" if hi and p["order"] == "desc" else "sabse purana" if hi else "latest" if p["order"] == "desc" else "earliest"
        where = f" ({fdesc})" if p["filters"] else ""
        say = f"{t.name}{where}, {first} by {by}: {names}"
    return {"say": say, "formula": formula, "sheet": sname, "value": len(rows), "value_label": "rows", "where": f"{t.name} ({fdesc})",
            "sample": _sample(t, rows)}


def _run_describe(p, tables, wb, hi, n):
    ts = [tables[p["table"]]] if p.get("table") else list(tables.values())
    parts = []
    for t in ts:
        cols = ", ".join(f"{h} ({ {'number': 'number', 'date': 'date', 'text': 'text', 'empty': 'khaali' if hi else 'empty'}[t.kind(i)]})"
                         for i, h in enumerate(t.header) if h)
        parts.append(f"{t.name}: {len(t.rows)} rows; {cols}")
    return {"say": " | ".join(parts), "formula": "", "sheet": None}


def _run_aging(p, tables, wb, hi, n):
    t = tables[p["table"]]
    today = date.today()
    di, ai = p["date"], p["amount"]
    order = ["0-30", "31-60", "61-90", "90+"]
    rows, summary = [], {b: [0, 0] for b in order}
    for r in t.rows:
        if not _match_filters(r, p["filters"]):
            continue
        d = r[di] if di < len(r) else None
        if not isinstance(d, (datetime, date)):
            continue
        days = (today - (d.date() if isinstance(d, datetime) else d)).days
        b = "0-30" if days <= 30 else "31-60" if days <= 60 else "61-90" if days <= 90 else "90+"
        amt = _num(r[ai]) if ai is not None and ai < len(r) else None
        summary[b][0] += 1
        summary[b][1] += amt or 0
        rows.append(list(r) + [days, b])
    c = get_column_letter(di + 1)
    formula = f'=TODAY()-{c}{t.header_row + 1}  then  =IF(X2<=30,"0-30",IF(X2<=60,"31-60",IF(X2<=90,"61-90","90+")))'
    sname = None
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, "Aging")
        _write_table(ws, ["Bucket (days)", "Count", t.header[ai] if ai is not None else "Amount"], [(b, *summary[b]) for b in order])
        ws.append([])
        ws.append(t.header + ["Days", "Bucket"])
        for r in rows:
            ws.append(r)
        ws.append([])
        ws.append(["Excel formula:", formula])
    lines = "; ".join(f"{b} din: {summary[b][0]} ({_fmt(summary[b][1])})" if hi else f"{b} days: {summary[b][0]} ({_fmt(summary[b][1])})" for b in order)
    say = (f"{t.name} aging (aaj {today:%d-%b-%Y} tak, {t.header[di]} se): {lines}" if hi else
           f"{t.name} aging (as of {today:%d-%b-%Y}, from {t.header[di]}): {lines}")
    return {"say": say, "formula": formula, "sheet": sname}


def _run_lookup(p, tables, wb, hi, n):
    dst, src = tables[p["table"]], tables[p["source"]]
    ki, kj = p["key"]
    wj = p["want"]
    index = {}
    for r in src.rows:
        k = r[kj] if kj < len(r) else None
        if k not in (None, ""):
            index.setdefault(str(k).strip().lower(), r[wj] if wj < len(r) else None)
    found = missing = 0
    col_letter = get_column_letter(len(dst.header) + 1)
    k_letter = get_column_letter(ki + 1)
    first = dst.header_row + 1
    formula = (f"=IFERROR(INDEX({src.ref(wj)},MATCH({k_letter}{first},{src.ref(kj)},0)),\"Not found\")")
    if wb is not None:
        ws = wb[dst.name]
        hdr_cell = ws.cell(dst.header_row, len(dst.header) + 1, f"{src.header[wj]} (from {src.name})")
        hdr_cell.font = Font(bold=True)
        r_i = first
        for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
            k = row[ki].value if ki < len(row) else None
            if k in (None, ""):
                continue
            ws.cell(row[0].row, len(dst.header) + 1, formula.replace(f"{k_letter}{first}", f"{k_letter}{row[0].row}"))
            if str(k).strip().lower() in index:
                found += 1
            else:
                missing += 1
            r_i += 1
    say = (f"{dst.name} mein naya column {col_letter} \"{src.header[wj]} (from {src.name})\" jod diya, {dst.header[ki]} ko {src.name} ke "
           f"{src.header[kj]} se milake. {found} rows mein value mili, {missing} mein nahi (wahan \"Not found\")." if hi else
           f"Added column {col_letter} \"{src.header[wj]} (from {src.name})\" to {dst.name}, matching {dst.header[ki]} with {src.name}'s "
           f"{src.header[kj]}. Found for {found} rows, not found for {missing} (shown as \"Not found\").")
    return {"say": say, "formula": formula, "sheet": dst.name}


def _run_combine(p, tables, wb, hi, n):
    ts = [tables[name] for name in p["tables"]]
    total = sum(len(t.rows) for t in ts)
    formula = "=VSTACK(" + ",".join(f"{t.ref(0)[:t.ref(0).index('!') + 1]}A{t.header_row + 1}:{get_column_letter(len(t.header))}{t.header_row + len(t.rows)}" for t in ts) + ")"
    sname = None
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, "Combined")
        _write_table(ws, ts[0].header + ["From sheet"], [list(r) + [t.name] for t in ts for r in t.rows], formula)
    say = (f"{', '.join(p['tables'])} ko ek sheet \"{sname}\" mein jod diya: {total} rows, har row ke saath uski sheet ka naam." if hi else
           f"Combined {', '.join(p['tables'])} into one sheet \"{sname}\": {total} rows, each with the sheet it came from.")
    return {"say": say, "formula": formula, "sheet": sname}


def _run_highlight(p, tables, wb, hi, n):
    t = tables[p["table"]]
    first, last = t.header_row + 1, t.header_row + len(t.rows)
    end = get_column_letter(len(t.header))
    if "filters" in p:
        cond = "AND(" + ",".join(f'${get_column_letter(i + 1)}{first}="{v}"' for i, v in p["filters"]) + ")"
        hits = sum(1 for r in t.rows if _match_filters(r, p["filters"]))
        what = _fdesc(t, p["filters"])
    else:
        L = f"${get_column_letter(p['left'] + 1)}{first}"
        R = f"${get_column_letter(p['right_col'] + 1)}{first}" if "right_col" in p else _fmt(p["right_val"]).replace(",", "")
        cond = f"{L}{p['op']}{R}"
        hits = 0
        for r in t.rows:
            a = _num(r[p["left"]]) if p["left"] < len(r) else None
            b = _num(r[p["right_col"]]) if "right_col" in p and p["right_col"] < len(r) else p.get("right_val")
            if a is not None and b is not None and (a < b if p["op"] == "<" else a > b):
                hits += 1
        rname = t.header[p["right_col"]] if "right_col" in p else _fmt(p["right_val"])
        what = f"{t.header[p['left']]} {p['op']} {rname}"
    if wb is not None:
        wb[t.name].conditional_formatting.add(f"A{first}:{end}{last}",
                                              FormulaRule(formula=[cond], fill=PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")))
    say = (f"{t.name} mein {hits} rows highlight kar di (peela rang) jahan {what}. Ye rule live hai: data badlega to rang bhi badlega." if hi else
           f"Highlighted {hits} rows in {t.name} (yellow) where {what}. The rule is live: change the data and the colour follows.")
    return {"say": say, "formula": f"Conditional formatting rule: ={cond}", "sheet": t.name}


from app.office.ask_more import (_chart, _plan_calc, _plan_growth, _plan_missing, _plan_reconcile, _plan_review,  # noqa: E402,F401
                                 _plan_running, _plan_sort, _plan_variance, _run_calc, _run_growth, _run_missing, _run_pivot,
                                 _run_reconcile, _run_review, _run_running, _run_sort, _run_variance)
