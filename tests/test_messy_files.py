"""Messy everyday workbooks (sales, staff, stock, fees, expenses, rent) spoiled at random: every spoiled cell is put back
or flagged, every clean cell is left alone, nothing clean is flagged. The generator lives in tests/fuzz/messy.py."""

import pytest
from openpyxl import Workbook, load_workbook

from app import office
from tests.fuzz import messy


@pytest.mark.parametrize("seed", list(range(1, 41)) + [64, 144, 153, 159, 173, 439, 585])
def test_messy_file_is_fixed_or_flagged(seed, tmp_path):
    assert messy.check(seed, tmp_path) == []


def _fix(tmp_path, rows, ops=messy.OPS):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.save(tmp_path / "in.xlsx")
    res = office.fix(str(tmp_path / "in.xlsx"), str(tmp_path / "out.xlsx"), ops)
    return load_workbook(tmp_path / "out.xlsx").active, res["changes"]


def test_phone_numbers_stay_text_even_under_a_misspelt_heading(tmp_path):
    ws, ch = _fix(tmp_path, [["Name", "Moble"]] + [[f"P{i}", f"98{i:08d}"] for i in range(10)])
    assert ws["A1"].value == "Name" and ws["B1"].value == "Mobile"
    assert all(isinstance(ws.cell(r, 2).value, str) for r in range(2, 12))
    assert not ch.get("text_to_number")


def test_two_spellings_as_common_as_each_other_are_asked_about_not_guessed(tmp_path):
    rows = [["Item", "Group"]] + [[f"i{i}", g] for i, g in enumerate(["Snacks"] * 6 + ["Appliance", "Appliances", "Appliance", "Appliances"])]
    ws, ch = _fix(tmp_path, rows)
    assert ch.get("flag_maybe_same") == 4
    assert [ws.cell(r, 2).value for r in range(8, 12)] == ["Appliance", "Appliances", "Appliance", "Appliances"]


def test_a_rare_slip_next_to_the_right_spelling_is_fixed(tmp_path):
    rows = [["Item", "Group"]] + [[f"i{i}", g] for i, g in enumerate(["Internet"] * 5 + ["Travel"] * 5 + ["Intenret"])]
    ws, ch = _fix(tmp_path, rows)
    assert ws["B12"].value == "Internet" and ch.get("fix_categories") == 1


def test_similar_names_that_really_differ_are_left_alone(tmp_path):
    rows = [["Student", "Class"]] + [[f"s{i}", c] for i, c in enumerate(["Class A"] * 6 + ["Class B"] * 2 + ["Unit 1"] * 3 + ["Unit 2"])]
    ws, ch = _fix(tmp_path, rows)
    assert not ch.get("fix_categories") and not ch.get("flag_maybe_same")


def test_properly_written_form_wins_a_tie(tmp_path):
    rows = [["Store", "Region"]] + [[f"s{i}", g] for i, g in enumerate(["North"] * 4 + ["South"] * 3 + ["west", "West", "CENTRAL"])]
    ws, _ = _fix(tmp_path, rows)
    assert [ws.cell(r, 2).value for r in range(9, 12)] == ["West", "West", "Central"]
