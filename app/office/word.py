"""Word documents (.docx): read, describe, answer questions on the text and the tables, and make the everyday changes
(find & replace, fix spacing / empty lines, heading case) as a new version so Undo works.

Tables in the document go through the same question engine as Excel sheets ("Table 2: total by region").
Text questions are answered from the document itself: counts, outline, a short summary made of its own sentences,
where something is mentioned, what it says about a topic, and every date / amount / email / phone in it. No AI."""

import re
from collections import Counter

from docx import Document

from app.office.sheetdata import Table, _value, _without_totals

HEADING_RE = re.compile(r"^(heading|title|subtitle)\s*\d*$", re.I)
WORD_RE = re.compile(r"[\w'’-]+", re.UNICODE)
SENT_RE = re.compile(r"(?<=[.!?।])\s+(?=[A-Z0-9\"'(ऀ-ॿ])")

# what the person wants from a document (English / Hinglish)
COUNT_RE = re.compile(r"\b(how many|kitne|kitni|kitna|count|number of|total)\b.{0,25}\b(words?|shabd|pages?|panne|paragraphs?|para|"
                      r"lines?|tables?|headings?|sections?|sentences?|images?|pictures?|photos?)\b|\bword count\b|\blength of\b", re.I)
OUTLINE_RE = re.compile(r"\b(outline|headings|sections|table of contents|contents|toc|structure|index|kya kya hai|kya hai is|"
                        r"what.{0,15}(in|inside) (this|the|my) (doc|document|file)|describe|samjhao|explain (this|the) (doc|document|file))\b", re.I)
SUMMARY_RE = re.compile(r"\b(summari[sz]e|summary|summery|saar|saransh|gist|overview|short mein|brief|key points|main points|tl;?dr|"
                        r"important points|zaroori baatein|mukhya)\b", re.I)
FIND_RE = re.compile(r"\b(find|search|where|kahan|kis jagah|locate|mentioned|mention|occurrences?|appear|aaya|likha hai)\b", re.I)
ABOUT_RE = re.compile(r"\b(what does it say about|what is (said|written) about|tell me about|about the|regarding|ke baare mein|ke bare me|"
                      r"ke bare mein|kya likha hai|kya bola hai|details of|explain)\b", re.I)
EXTRACT = {
    "email": re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+"),
    "phone": re.compile(r"(?<!\d)(?:\+91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}(?!\d)|(?<!\d)0\d{2,4}[\s-]\d{6,8}(?!\d)"),
    "date": re.compile(r"\b\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}\b|\b\d{1,2}(?:st|nd|rd|th)?\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)"
                       r"[a-z]*,?\s+\d{4}\b|\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\s+\d{1,2},?\s+\d{4}\b", re.I),
    "amount": re.compile(r"(?:₹|rs\.?|inr|\$|usd|€|£)\s?\d[\d,]*(?:\.\d+)?(?:\s?(?:lakh|lakhs|crore|cr|k|million|mn))?|"
                         r"\b\d[\d,]*(?:\.\d+)?\s?(?:lakh|lakhs|crore|rupees|dollars)\b", re.I),
    "link": re.compile(r"https?://\S+|www\.\S+", re.I),
}
EXTRACT_WORDS = {
    "email": r"emails?|e-mails?|mail ids?", "phone": r"phones?|mobiles?|contact numbers?|phone numbers?|numbers? to call",
    "date": r"dates?|deadlines?|tarikh", "amount": r"amounts?|prices?|money|rupees|costs?|fees?|paise|rakam|figures",
    "link": r"links?|urls?|websites?",
}
LIST_RE = re.compile(r"\b(list|all|every|saare|saari|sab|nikalo|extract|show|give|dikhao|batao|which|kaunse|kaun se)\b", re.I)
REPLACE_RE = [
    re.compile(r"\breplace\s+(?:all\s+)?(?:the\s+)?(?:word\s+)?['\"“]?(.+?)['\"”]?\s+with\s+['\"“]?(.*?)['\"”]?\s*$", re.I),
    re.compile(r"\bchange\s+(?:all\s+)?(?:the\s+)?(?:word\s+)?['\"“]?(.+?)['\"”]?\s+to\s+['\"“]?(.*?)['\"”]?\s*$", re.I),
    re.compile(r"['\"“]?(.+?)['\"”]?\s+ko\s+['\"“]?(.+?)['\"”]?\s+(?:se|mein|me)\s+(?:badlo|badal do|replace karo|kar do|karo)\s*$", re.I),
]
SPACING_RE = re.compile(r"\b(fix|clean|remove|hatao|hata do|theek|thik|sudharo|saaf)\b.{0,30}\b(spac\w*|double spaces?|extra spaces?|"
                        r"empty lines?|blank lines?|khaali lines?|gaps?|formatting)\b|\b(spacing|extra spaces?|double spaces?|blank lines?)\b.{0,20}"
                        r"\b(theek|thik|hatao|fix|remove|saaf)\b|\bclean (up )?(the |this |my )?(doc|document|file)\b", re.I)
HEAD_CASE_RE = re.compile(r"\b(headings?|titles?)\b.{0,30}\b(upper ?case|capital|caps|title case|proper case|first letter capital)\b|"
                          r"\b(upper ?case|capital|caps|title case|proper case)\b.{0,30}\b(headings?|titles?)\b", re.I)
STOP = set("""a an the is are was were be been of to in on for and or but with at by from this that these those it its as into about
what which who whom how why where when does do did can could should would will shall may might me my our your their there here
tell say says said written write document doc file please give show find search mention mentioned kya hai hain ka ki ke ko se
mein me par pe aur ya bhi is us ye wo yeh woh likha bola baare bare batao dikhao""".split())


class Doc:
    """The document as plain parts: paragraphs (style, text), its tables, headings."""

    def __init__(self, path):
        d = Document(path)
        self.paras = [(p.style.name if p.style is not None else "", p.text) for p in d.paragraphs]
        self.tables = d.tables
        self.images = len(d.inline_shapes)
        self.sections = len(d.sections)

    @property
    def text_paras(self):
        return [(i, s, t) for i, (s, t) in enumerate(self.paras) if t.strip()]

    @property
    def headings(self):
        return [(s, t.strip()) for s, t in self.paras if t.strip() and (HEADING_RE.match(s or "") or _looks_heading(s, t))]

    def words(self):
        n = sum(len(WORD_RE.findall(t)) for _, t in self.paras)
        for tb in self.tables:
            for row in tb.rows:
                for c in row.cells:
                    n += len(WORD_RE.findall(c.text))
        return n


def _looks_heading(style, text):
    t = text.strip()
    return bool(not HEADING_RE.match(style or "") and 0 < len(t) <= 70 and not t.endswith((".", ",", ";", ":")) and
                (t.isupper() and len(t) > 3) and len(t.split()) <= 10)


_docs: dict = {}


def load(path: str) -> Doc:
    import os

    key = (path, os.path.getmtime(path))
    if key not in _docs:
        if len(_docs) > 16:
            _docs.clear()
        _docs[key] = Doc(path)
    return _docs[key]


# ------------------------------------------------------------------ tables, for the Excel question engine
def tables(path: str) -> list[Table]:
    out = []
    d = load(path)
    title = _table_titles(path)
    for k, tb in enumerate(d.tables, start=1):
        rows = []
        for row in tb.rows:
            cells, last = [], None
            for c in row.cells:
                if c._tc is last:                       # merged cells repeat: keep one
                    continue
                last = c._tc
                cells.append(c.text.strip())
            rows.append(cells)
        rows = [r for r in rows if any(v for v in r)]
        if len(rows) < 2:
            continue
        header = [h or f"Column {i + 1}" for i, h in enumerate(rows[0])]
        data = _without_totals([[_value(v) if v != "" else None for v in r] for r in rows[1:]])
        name = f"Table {k}" + (f" ({title[k - 1]})" if title.get(k - 1) else "")
        out.append(Table(name[:60], header, data, 1))
    return out


def _table_titles(path):
    """The line just before each table ('Price list:'), to name it."""
    from docx.text.paragraph import Paragraph

    d = Document(path)
    body = d.element.body
    titles, last_text, k = {}, "", 0
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            txt = Paragraph(el, d).text.strip()
            if txt:
                last_text = txt
        elif tag == "tbl":
            if last_text and len(last_text) <= 40:
                titles[k] = last_text.rstrip(":")
            k += 1
            last_text = ""
    return titles


# ------------------------------------------------------------------ the check on arrival
def scan(path: str) -> dict:
    import os

    d = load(path)
    tps = d.text_paras
    double = sum(1 for _, _, t in tps if "  " in t.strip() or t != t.strip())
    empties, run = 0, 0
    for _, t in d.paras:
        if not t.strip():
            run += 1
            if run >= 2:
                empties += 1
        else:
            run = 0
    tbls = tables(path)
    words = d.words()
    summary = {"sheets": len(tbls), "rows": sum(len(t.rows) for t in tbls), "formulas": 0, "words": words,
               "paragraphs": len(tps), "headings": len(d.headings), "tables": len(d.tables), "images": d.images,
               "pages": max(1, round(words / 450)) if words else 0, "double_spaces": double, "empty_lines": empties}
    return {"kind": "docx", "size": os.path.getsize(path), "sheets": [], "summary": summary,
            "outline": [t for _, t in d.headings][:40]}


# ------------------------------------------------------------------ understanding what they asked
def plan(text: str, path: str) -> dict | None:
    """A plan for a question / order about the document's text, or None (then its tables are tried)."""
    low = (text or "").strip()
    for rx in REPLACE_RE:
        m = rx.search(low)
        if m and m.group(1).strip() and not re.search(r"\b(spacing|spaces|format)\b", m.group(1), re.I):
            return {"job": "word_replace", "find": m.group(1).strip(), "repl": (m.group(2) or "").strip()}
    if SPACING_RE.search(low):
        return {"job": "word_spacing"}
    if HEAD_CASE_RE.search(low):
        how = "upper" if re.search(r"upper|capital|caps", low, re.I) and not re.search(r"first letter", low, re.I) else "title"
        return {"job": "word_heading_case", "how": how}
    for kind, words in EXTRACT_WORDS.items():
        plural = re.search(rf"\b({words})\b", low, re.I) and re.search(r"s\b", re.search(rf"\b({words})\b", low, re.I).group(0))
        if re.search(rf"\b({words})\b", low, re.I) and (LIST_RE.search(low) or plural or kind in ("email", "phone", "link") and
                                                         re.search(r"\b(kya|what|kaun|which)\b", low, re.I)) \
                and not re.search(r"\b(table|column|total|sum|average)\b", low, re.I):
            return {"job": "word_extract", "what": kind}
    m = COUNT_RE.search(low)
    if m and not re.search(r"\b(rows?|records?|column|sum|average|amount|sales|price)\b", low, re.I):
        return {"job": "word_count", "what": m.group(2).lower() if m.group(2) else "words"}
    if SUMMARY_RE.search(low):
        return {"job": "word_summary"}
    if OUTLINE_RE.search(low) and not re.search(r"\btable \d", low, re.I):
        return {"job": "word_outline"}
    term = _quoted(low)
    if FIND_RE.search(low) and (term or _topic(low)):
        t = term or _topic(low)
        m = re.search(re.escape(t), text or "", re.I)
        return {"job": "word_find", "term": m.group(0) if m else t}       # as they typed it ('Mumbai', not 'mumbai')
    if ABOUT_RE.search(low) and _topic(low):
        return {"job": "word_about", "topic": _topic(low)}
    return None


def fallback(text: str, path: str) -> dict | None:
    """Nothing else fitted: what the document says closest to the question, if anything is close."""
    topic = _topic(text)
    if not topic:
        return None
    if not _best_paras(load(path), text, k=1, min_score=0.12):
        return None
    return {"job": "word_about", "topic": topic, "q": text}


def _quoted(text):
    m = re.search(r"['\"“‘]([^'\"”’]{1,60})['\"”’]", text or "")
    return m.group(1).strip() if m else None


def _topic(text):
    ws = [w for w in WORD_RE.findall((text or "").lower()) if w not in STOP and len(w) > 1]
    ws = [w for w in ws if not re.match(r"^(word|words|para|paragraph|line|lines|section|page|pages|times|baar|kitni|kitne|"
                                        r"how|many|where|kahan|search|find|locate|mentioned|mention|appear|occurrences?|aaya)$", w)]
    return " ".join(ws[:8]) or None


# ------------------------------------------------------------------ answering
def run(p: dict, path: str, out: str, hi: bool) -> dict:
    job = p["job"]
    d = load(path)
    if job == "word_count":
        return _count(d, p["what"], hi)
    if job == "word_outline":
        return _outline(d, path, hi)
    if job == "word_summary":
        return _summary(d, hi)
    if job == "word_find":
        return _find(d, p["term"], hi)
    if job == "word_about":
        return _about(d, p.get("q") or p["topic"], p["topic"], hi)
    if job == "word_extract":
        return _extract(d, p["what"], hi)
    if job in ("word_replace", "word_spacing", "word_heading_case"):
        return _edit(p, path, out, hi)
    raise ValueError(job)


def _res(say, **kw):
    return {"say": say, "formula": "", "sheet": None, **kw}


def _count(d, what, hi):
    w = what.rstrip("s")
    if w in ("word", "shabd"):
        n, label = d.words(), ("shabd" if hi else "words")
    elif w in ("page", "panne"):
        n = max(1, round(d.words() / 450)) if d.words() else 0
        pg = "page" if n == 1 else "pages"
        return _res((f"Lagbhag {n} {pg} (har page ~450 shabd maan ke; asli ginti Word mein font aur margin pe nirbhar hai)" if hi else
                     f"About {n} {pg} (at ~450 words a page; the exact count in Word depends on font and margins)"),
                    value=n, value_label="pages (approx.)" if not hi else "pages (lagbhag)")
    elif w in ("paragraph", "para"):
        n, label = len(d.text_paras), "paragraphs"
    elif w in ("line", "sentence"):
        n = sum(len([s for s in SENT_RE.split(t) if s.strip()]) for _, _, t in d.text_paras)
        label = "sentences"
    elif w == "table":
        n, label = len(d.tables), "tables"
    elif w in ("heading", "section"):
        n, label = len(d.headings), "headings"
    elif w in ("image", "picture", "photo"):
        n, label = d.images, "images"
    else:
        n, label = d.words(), "words"
    label = label[:-1] if n == 1 and label.endswith("s") else label
    return _res((f"Is document mein {n:,} {label} hain" if hi else f"This document has {n:,} {label}"), value=n, value_label=label)


def _outline(d, path, hi):
    heads = d.headings
    tbls = tables(path)
    words = d.words()
    bits = []
    if heads:
        shown = "; ".join(t for _, t in heads[:15]) + (" …" if len(heads) > 15 else "")
        bits.append((f"{len(heads)} headings: {shown}" if not hi else f"{len(heads)} headings: {shown}"))
    else:
        first = next((t for _, _, t in d.text_paras), "")
        bits.append(("Headings nahi hain; shuruaat: " if hi else "No headings; it starts: ") + f"\"{first[:140]}\"")
    if tbls:
        bits.append((f"{len(tbls)} tables: " if not hi else f"{len(tbls)} tables: ") +
                    "; ".join(f"{t.name} ({len(t.rows)} rows: {', '.join(h for h in t.header[:5] if h)})" for t in tbls[:6]))
    pages = max(1, round(words / 450))
    head = (f"Document: {words:,} shabd, {len(d.text_paras)} paragraphs, lagbhag {pages} page{'s' if pages > 1 else ''}. " if hi else
            f"The document: {words:,} words, {len(d.text_paras)} paragraphs, about {pages} page{'s' if pages > 1 else ''}. ")
    sample = {"sheet": "Outline", "header": ["Level", "Heading"],
              "rows": [[_level(s), t] for s, t in heads[:12]], "total": len(heads)} if heads else None
    return _res(head + ". ".join(bits), **({"sample": sample} if sample else {}))


def _level(style):
    m = re.search(r"(\d+)", style or "")
    return int(m.group(1)) if m else (0 if re.match(r"title", style or "", re.I) else 1)


def _sentences(d):
    out = []
    for i, s, t in d.text_paras:
        if HEADING_RE.match(s or "") or _looks_heading(s, t):
            continue
        for sent in SENT_RE.split(t.strip()):
            sent = sent.strip()
            if 6 <= len(sent.split()) <= 60:
                out.append((i, sent))
    return out


def _summary(d, hi):
    """The document in its own words: the sentences that carry the most of what it keeps talking about."""
    sents = _sentences(d)
    if not sents:
        return _res("Is document mein summary banane layak text nahi hai" if hi else "There isn't enough text in this document to summarise")
    freq = Counter(w for _, s in sents for w in WORD_RE.findall(s.lower()) if w not in STOP and len(w) > 2)
    top = {w for w, _ in freq.most_common(40)}
    n = len(sents)

    def score(k, s):
        ws = [w for w in WORD_RE.findall(s.lower()) if w not in STOP and len(w) > 2]
        if not ws:
            return 0
        base = sum(freq[w] for w in ws if w in top) / (len(ws) ** 0.6)
        return base * (1.25 if k < max(3, n // 10) else 1)          # openings usually state the point

    k = 3 if n <= 12 else 5 if n <= 60 else 7
    pick = sorted(sorted(range(n), key=lambda j: -score(j, sents[j][1]))[:k])
    points = [sents[j][1] for j in pick]
    heads = [t for _, t in d.headings[:8]]
    say = ("Mukhya baatein (document ke apne vaakya): " if hi else "Key points, in the document's own sentences: ") + " • ".join(points)
    if heads:
        say += (". Sections: " if not hi else ". Sections: ") + "; ".join(heads)
    return _res(say, sample={"sheet": "Summary", "header": ["#", "Point"], "rows": [[i + 1, p] for i, p in enumerate(points)],
                             "total": len(points)})


def _context(t, start, end, width=70):
    a, b = max(0, start - width), min(len(t), end + width)
    return ("…" if a else "") + t[a:b].strip() + ("…" if b < len(t) else "")


def _find(d, term, hi):
    rx = re.compile(r"(?<!\w)" + re.escape(term).replace(r"\ ", r"\s+") + r"(?!\w)", re.I)
    hits = []
    for i, s, t in d.text_paras:
        for m in rx.finditer(t):
            hits.append((i, _context(t, m.start(), m.end())))
    for k, tb in enumerate(d.tables, start=1):
        for r_i, row in enumerate(tb.rows, start=1):
            for c in row.cells:
                if rx.search(c.text):
                    hits.append((f"Table {k} row {r_i}", c.text.strip()[:140]))
                    break
    if not hits:
        words = [w for w in term.split() if len(w) > 2]
        near = _best_paras(d, term, k=2, min_score=0.15) if words else []
        say = (f"\"{term}\" is document mein nahi mila" if hi else f"\"{term}\" isn't in this document")
        if near:
            say += (". Sabse milta-julta: " if hi else ". Closest match: ") + " | ".join(f"\"{t[:160]}\"" for _, t in near)
        return _res(say, value=0, value_label="matches")
    paras = len({h[0] for h in hits})
    where = "; ".join(f"{_where(h[0], hi)}: \"{h[1]}\"" for h in hits[:4])
    say = (f"\"{term}\" {len(hits)} baar aaya ({paras} jagah). {where}" if hi else
           f"\"{term}\" appears {len(hits)} times, in {paras} places. {where}")
    return _res(say, value=len(hits), value_label="matches",
                sample={"sheet": "Matches", "header": ["Where", "Text"], "rows": [[_where(w, hi), t] for w, t in hits[:8]], "total": len(hits)})


def _where(i, hi):
    return i if isinstance(i, str) else (f"Paragraph {i + 1}" if not hi else f"Paragraph {i + 1}")


def _best_paras(d, query, k=3, min_score=0.12, sentences=False):
    from sklearn.feature_extraction.text import TfidfVectorizer

    paras = [(i, t) for i, s, t in d.text_paras if len(t.split()) >= 3]
    if sentences:                                       # the sentence that answers, not the whole paragraph
        paras = [(i, x.strip()) for i, t in paras for x in SENT_RE.split(t) if len(x.split()) >= 3]
    for kk, tb in enumerate(d.tables, start=1):
        for r_i, row in enumerate(tb.rows, start=1):
            txt = " | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
            if txt:
                paras.append((f"Table {kk} row {r_i}", txt))
    if not paras:
        return []
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), sublinear_tf=True)
    mat = vec.fit_transform([t.lower() for _, t in paras])
    q = " ".join(w for w in WORD_RE.findall((query or "").lower()) if w not in STOP) or query
    sims = (mat @ vec.transform([q.lower()]).T).toarray().ravel()
    order = sims.argsort()[::-1][:k]
    floor = max(min_score, 0.6 * sims[order[0]]) if len(order) else min_score
    return [paras[j] for j in order if sims[j] >= floor]       # the clear matches, not whatever comes next


def _about(d, q, topic, hi):
    best = _best_paras(d, q, k=2, min_score=0.12, sentences=True) or _best_paras(d, q, k=3, min_score=0.12)
    if not best:
        return _res((f"Document mein \"{topic}\" ke baare mein kuch nahi mila" if hi else f"The document doesn't say anything about \"{topic}\""),
                    value=0, value_label="matches")
    lead = (f"Document mein \"{topic}\" ke baare mein ye likha hai: " if hi else f"What the document says about \"{topic}\": ")
    say = lead + " | ".join(f"{_where(w, hi)}: \"{t[:300]}{'…' if len(t) > 300 else ''}\"" for w, t in best)
    return _res(say, sample={"sheet": "From the document", "header": ["Where", "Text"], "rows": [[_where(w, hi), t[:300]] for w, t in best],
                             "total": len(best)})


def _extract(d, what, hi):
    rx = EXTRACT[what]
    found = []
    texts = [t for _, _, t in d.text_paras] + [c.text for tb in d.tables for row in tb.rows for c in row.cells]
    for t in texts:
        for m in rx.finditer(t):
            v = m.group(0).strip().rstrip(".,;)")
            if v not in found:
                found.append(v)
    label = {"email": "emails", "phone": "phone numbers", "date": "dates", "amount": "amounts", "link": "links"}[what]
    if not found:
        return _res((f"Is document mein koi {label} nahi mila" if hi else f"No {label} in this document"), value=0, value_label=label)
    shown = ", ".join(found[:20]) + (" …" if len(found) > 20 else "")
    one = label[:-1] if len(found) == 1 else label
    return _res((f"{len(found)} {one}: {shown}" if hi else f"{len(found)} {one} in the document: {shown}"), value=len(found), value_label=label,
                sample={"sheet": label.capitalize(), "header": ["#", label.capitalize()], "rows": [[i + 1, v] for i, v in enumerate(found[:12])],
                        "total": len(found)})


# ------------------------------------------------------------------ changes (saved as a new version)
def _all_paragraphs(doc):
    yield from doc.paragraphs
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                yield from c.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            if part is not None:
                yield from part.paragraphs


def _replace_in(par, rx, repl) -> int:
    runs = par.runs
    full = "".join(r.text for r in runs)
    want = len(rx.findall(full))
    if not want:
        return 0
    if sum(len(rx.findall(r.text)) for r in runs) == want:
        for r in runs:                                  # each match inside one run: the formatting stays exactly as it was
            if rx.search(r.text):
                r.text = rx.sub(repl, r.text)
        return want
    runs[0].text = rx.sub(repl, full)                   # a match split across runs (Word does that): one run carries the text
    for r in runs[1:]:
        r.text = ""
    return want


def _has_objects(par):
    return bool(par._element.xpath(".//w:drawing | .//w:pict | .//w:sectPr | .//w:br[@w:type='page'] | .//w:object"))


def _edit(p, path, out, hi):
    doc = Document(path)
    job = p["job"]
    n = 0
    if job == "word_replace":
        f, rp = p["find"], p["repl"]
        rx = re.compile(r"(?<!\w)" + re.escape(f) + r"(?!\w)" if re.match(r"^\w.*\w$|^\w$", f) else re.escape(f), re.I)
        for par in _all_paragraphs(doc):
            n += _replace_in(par, rx, rp.replace("\\", "\\\\"))
        say = ((f"\"{f}\" ko \"{rp}\" kar diya, {n} jagah" if n else f"\"{f}\" document mein nahi mila; kuch nahi badla") if hi else
               (f"Replaced \"{f}\" with \"{rp}\" in {n} places" if n else f"\"{f}\" isn't in the document; nothing changed"))
    elif job == "word_spacing":
        spaces = empties = 0
        for par in _all_paragraphs(doc):
            runs = par.runs
            for r in runs:
                new = re.sub(r"(?<=\S) {2,}(?=\S)", " ", r.text)
                if new != r.text:
                    spaces += 1
                    r.text = new
            if runs and runs[0].text != runs[0].text.lstrip(" ") and par.style is not None and not re.search(r"code|pre", par.style.name, re.I):
                runs[0].text = runs[0].text.lstrip(" ")
                spaces += 1
            if runs and runs[-1].text != runs[-1].text.rstrip(" "):
                runs[-1].text = runs[-1].text.rstrip(" ")
                spaces += 1
        prev_empty = False
        for par in list(doc.paragraphs):
            empty = not par.text.strip() and not _has_objects(par)
            if empty and prev_empty:
                par._element.getparent().remove(par._element)
                empties += 1
                continue
            prev_empty = empty
        n = spaces + empties
        say = (f"Spacing theek ki: {spaces} jagah extra space hataye, {empties} faltu khaali lines hatayi" if hi else
               f"Fixed the spacing: removed extra spaces in {spaces} places and {empties} extra blank lines") if n else \
              ("Spacing pehle se theek hai; kuch nahi badla" if hi else "The spacing is already clean; nothing changed")
    else:
        for par in doc.paragraphs:
            st = par.style.name if par.style is not None else ""
            if par.text.strip() and (HEADING_RE.match(st) or _looks_heading(st, par.text)):
                for r in par.runs:
                    new = r.text.upper() if p["how"] == "upper" else _title(r.text)
                    if new != r.text:
                        r.text = new
                n += 1
        say = (f"{n} headings {'CAPITAL' if p['how'] == 'upper' else 'Title Case'} mein kar di" if hi else
               f"Changed {n} headings to {'UPPER CASE' if p['how'] == 'upper' else 'Title Case'}") if n else \
              ("Document mein headings nahi mili" if hi else "No headings found in the document")
    if n:
        doc.save(out)
        say += (". Nayi copy card pe Download karo; 'undo' likhoge to pehle jaisa ho jayega" if hi else
                ". Download the new copy from the card; type 'undo' to put it back")
    return _res(say, changes=n)


SMALL = {"a", "an", "the", "and", "or", "of", "in", "on", "to", "for", "at", "by", "with"}


def _title(s):
    words = s.split(" ")
    return " ".join(w if (w.isupper() and len(w) > 1) else (w.lower() if i and w.lower() in SMALL else w[:1].upper() + w[1:].lower())
                    for i, w in enumerate(words))
