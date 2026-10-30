"""Microsoft agent end to end: chat -> card -> upload in pieces -> automatic check -> fix -> download -> compare -> cleanup."""

import pytest
import re
import io
import time

from openpyxl import Workbook, load_workbook

from tests.test_ml import say


def _xlsx(rows, broken=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Invoice No", "Customer", "Qty", "Rate", "Amount"])
    for i, (inv, cust, q, r) in enumerate(rows, start=2):
        ws.append([inv, cust, q, r, f"=C{i}*D{i}"])
    if broken:
        ws["E4"] = "=C4*#REF!"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _upload(client, agent_id, data, name="sales.xlsx"):
    f = client.post("/api/office/files", json={"name": name, "size": len(data), "agent_id": agent_id}).json()
    half = len(data) // 2
    client.put(f"/api/office/files/{f['id']}/chunk?offset=0", content=data[:half])
    client.put(f"/api/office/files/{f['id']}/chunk?offset=0", content=data[:half])   # a retried piece is ignored
    r = client.put(f"/api/office/files/{f['id']}/chunk?offset={half}", content=data[half:]).json()
    assert r["received"] == len(data)
    client.post(f"/api/office/files/{f['id']}/done")
    return _wait(client, f["id"])


def _wait(client, token):
    for _ in range(200):
        f = client.get(f"/api/office/files/{token}").json()
        if f["status"] in ("ready", "error"):
            return f
        time.sleep(0.05)
    raise AssertionError("still working")


ROWS = [(f"INV{i}", f" Cust  {i} ", i, "1,200" if i == 3 else 100) for i in range(1, 9)]


def test_chat_to_fixed_file(client, ml_mode):
    r = say(client, "meri excel sheet mein #REF! aa raha hai row delete karne ke baad")
    agent = next(a for a in r["agents"] if a.get("office"))
    assert agent["brain"]["id"] == "excel-ref-error"
    assert "Samjha" in r["reply"] or "Understood" in r["reply"]

    f = _upload(client, agent["id"], _xlsx(ROWS, broken=True))
    assert f["status"] == "ready", f
    s = f["report"]["summary"]
    assert s["broken_refs"] == 1 and s["text_numbers"] == 1 and s["spaces"] >= 8

    client.post(f"/api/office/files/{f['id']}/fix", json={"ops": ["trim_spaces", "text_to_number", "fix_broken_refs"]})
    f = _wait(client, f["id"])
    assert f["download"] and f["result"]["changes"]["fix_broken_refs"] == 1
    d = client.get(f"/api/office/files/{f['id']}/download")
    ws = load_workbook(io.BytesIO(d.content))["Sales"]
    assert ws["E4"].value == "=C4*D4" and ws["D4"].value == 1200 and ws["B2"].value == "Cust 1"


def test_compare_two_uploads(client, ml_mode):
    r = say(client, "do excel sheet compare karni hai, data match nahi ho raha")
    agent = next(a for a in r["agents"] if a.get("office"))
    a = _upload(client, agent["id"], _xlsx(ROWS), "old.xlsx")
    new_rows = [row if row[0] != "INV2" else ("INV2", " Cust  2 ", 2, 999) for row in ROWS[1:]] + [("INV99", "New", 1, 1)]
    b = _upload(client, agent["id"], _xlsx(new_rows), "new.xlsx")
    client.post(f"/api/office/files/{a['id']}/compare", json={"other": b["id"]})
    a = _wait(client, a["id"])
    res = a["result"]
    assert res["key"] == "Invoice No" and (res["only_in_a"], res["only_in_b"], res["changed_rows"]) == (1, 1, 1)
    assert a["compare_download"]


def test_files_are_private_and_cleaned(client, ml_mode):
    r = say(client, "excel file ka data saaf karna hai")
    agent = next(a for a in r["agents"] if a.get("office"))
    f = _upload(client, agent["id"], _xlsx(ROWS))
    other = client.__class__(client.app)
    other.cookies.set("ge_vid", "someone-else-000000000000000000")
    assert other.get(f"/api/office/files/{f['id']}").status_code == 400
    assert client.post("/api/office/files", json={"name": "virus.exe", "size": 10}).status_code == 400
    client.post("/api/privacy/delete-all")
    from app.db import _SessionLocal
    from app.models import OfficeFile

    with _SessionLocal() as db:
        assert db.query(OfficeFile).filter(OfficeFile.token == f["id"]).count() == 0


def test_old_files_removed_after_an_hour(client, ml_mode):
    from datetime import timedelta

    from app.agents import office_agent
    from app.db import _SessionLocal
    from app.models import OfficeFile, utcnow

    r = say(client, "csv file mein duplicate hatane hain")
    agent = next(a for a in r["agents"] if a.get("office"))
    f = _upload(client, agent["id"], b"Name,Phone\nA,1\nA,1\n", "list.csv")
    with _SessionLocal() as db:
        row = db.query(OfficeFile).filter(OfficeFile.token == f["id"]).one()
        path = row.path
        assert office_agent.maintain(db, utcnow() + timedelta(hours=2)) >= 1
    import os

    assert not os.path.exists(path)


def test_work_cut_off_by_a_restart_carries_on(client, ml_mode):
    from datetime import timedelta

    from app.agents import office_agent
    from app.db import _SessionLocal
    from app.models import OfficeFile, utcnow

    r = say(client, "excel file ke formula check karo")
    agent = next(a for a in r["agents"] if a.get("office"))
    f = _upload(client, agent["id"], _xlsx(ROWS))
    with _SessionLocal() as db:                       # pretend the server restarted in the middle of a fix
        row = db.query(OfficeFile).filter(OfficeFile.token == f["id"]).one()
        row.status, row.task, row.args = "working", "fix", '{"ops": ["trim_spaces"]}'
        row.updated_at = utcnow() - timedelta(minutes=1)
        db.commit()
    assert office_agent.resume_interrupted(utcnow()) == 1
    f = _wait(client, f["id"])
    assert f["status"] == "ready" and f["result"]["changes"]["trim_spaces"] >= 8


def test_only_one_process_resumes(tmp_path):
    import subprocess
    import sys

    code = ("from app.agents import office_agent as o; import pathlib; o.DATA = pathlib.Path(r'%s');"
            "print(o._resume_owner())") % tmp_path
    from app.agents import office_agent as o
    import fcntl

    tmp_path.mkdir(exist_ok=True)
    held = open(tmp_path / ".resume.lock", "w")
    fcntl.flock(held, fcntl.LOCK_EX | fcntl.LOCK_NB)          # another worker already owns it
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=str(__import__("pathlib").Path(o.__file__).parents[2]))
    assert out.stdout.strip() == "False", out.stderr
    held.close()


def test_one_upload_many_tasks_from_chat_and_undo(client, ml_mode):
    r = say(client, "meri excel file mein extra spaces hain, hata do")
    goal_id = r["goal"]["id"]
    agent = next(a for a in r["agents"] if a.get("office"))
    rows = [(f"INV{i}", f" Cust  {i} ", i, "1,200" if i == 3 else 100) for i in range(1, 9)] + [("INV1", " Cust  1 ", 1, 100)]
    f = _upload(client, agent["id"], _xlsx(rows))
    token = f["id"]

    r = say(client, "extra space hata do", goal_id)
    assert "theek kar raha" in r["reply"] or "Fixing" in r["reply"], r["reply"]
    f = _wait(client, token)
    assert f["result"]["changes"]["trim_spaces"] >= 9 and f["report"]["summary"]["spaces"] == 0

    r = say(client, "ab number wale text bhi theek karo", goal_id)
    f = _wait(client, token)
    assert len(f["steps"]) == 2 and f["report"]["summary"]["text_numbers"] == 0

    d = client.get(f"/api/office/files/{token}/download")
    wb = load_workbook(io.BytesIO(d.content))
    ws = wb["Sales"]
    assert ws["B2"].value == "Cust 1" and ws["D4"].value == 1200          # both tasks are in one file
    changes = [r[0] for r in wb["Changes by Goal Engine"].iter_rows(min_row=2, values_only=True)]
    assert any(c.startswith("Extra space") for c in changes) and any(c.startswith("Text ko number") for c in changes)

    r = say(client, "undo karo", goal_id)
    f = _wait(client, token)
    assert len(f["steps"]) == 1 and f["report"]["summary"]["text_numbers"] == 1


def test_find_in_a_named_sheet_with_typos_then_yes(client, ml_mode):
    """Chetan's own flow: open file, 'find dupicate entries from slaes sheet' -> exact answer, then 'yes' fixes it."""
    from pathlib import Path

    r = say(client, "I need to work on excel file")
    assert "won't open" not in r["reply"] and "corrupt" not in r["reply"].lower()
    goal_id = r["goal"]["id"]
    agent = next(a for a in r["agents"] if a.get("office"))
    data = (Path(__file__).parent / "data" / "practice.xlsx").read_bytes()
    f = _upload(client, agent["id"], data, "Excel_Practice.xlsx")

    r = say(client, "can you find dupicate entries", goal_id)
    card = next(a for a in r["agents"] if a.get("office"))
    assert card["formula"] is None                                  # no formula homework while the file is open
    f = _wait(client, f["id"])
    res = f["result"]
    assert res["task"] == "extract" and res["rows"] == 2 and res["values"] == 0   # Item Code / Invoice No. repeat normally there
    assert f["report"]["summary"]["duplicate_rows"] == 2           # listing changed nothing
    wb = load_workbook(io.BytesIO(client.get(f"/api/office/files/{f['id']}/download").content))
    assert wb.sheetnames[0] == "Duplicates (Goal Engine)"
    listed = list(wb["Duplicates (Goal Engine)"].iter_rows(min_row=2, values_only=True))
    assert ("Messy Data", 6, 2) in [r[:3] for r in listed]

    r = say(client, "yes", goal_id)
    assert "Fixing" in r["reply"] or "theek kar raha" in r["reply"], r["reply"]
    f = _wait(client, f["id"])
    assert f["result"]["changes"].get("remove_duplicates") == 2 and f["report"]["summary"]["duplicate_rows"] == 0


def test_only_the_named_sheet_changes(client, ml_mode):
    from pathlib import Path

    r = say(client, "excel file check karni hai")
    goal_id = r["goal"]["id"]
    agent = next(a for a in r["agents"] if a.get("office"))
    data = (Path(__file__).parent / "data" / "practice.xlsx").read_bytes()
    f = _upload(client, agent["id"], data, "Excel_Practice.xlsx")
    say(client, "messy sheet mein capital small letters theek karo", goal_id)
    f = _wait(client, f["id"])
    assert f["result"]["changes"].get("fix_case") == 4
    say(client, "kya galat hai is file mein?", goal_id)


def test_chetans_hinglish_requests_never_delete(client, ml_mode):
    from pathlib import Path

    r = say(client, "muje excel sheet me kaam krna hai")
    goal_id = r["goal"]["id"]
    agent = next(a for a in r["agents"] if a.get("office"))
    f = _upload(client, agent["id"], (Path(__file__).parent / "data" / "practice.xlsx").read_bytes(), "Excel_Practice.xlsx")
    r = say(client, "muje sales sheet me duplicates dhund k do", goal_id)
    assert "Sales Data mein koi duplicate nahi" in r["reply"] and "Messy Data: 2" in r["reply"], r["reply"]
    for text in ("muej alag se sheet me do dalke muje chekc krna hai kon c values duplicate hai", "poori file ke duplicate dikhao"):
        r = say(client, text, goal_id)
        assert "2 duplicate mile" in r["reply"] and "Duplicates (Goal Engine)" in r["reply"], (text, r["reply"])
        f = _wait(client, f["id"])
        assert f["result"]["task"] == "extract" and f["report"]["summary"]["duplicate_rows"] == 2, text
    r = say(client, "duplicate hatao", goal_id)
    f = _wait(client, f["id"])
    assert f["report"]["summary"]["duplicate_rows"] == 0


def test_formula_requests_get_a_formula_for_their_sheet(client, ml_mode):
    from pathlib import Path

    r = say(client, "muje excel sheet me kaam krna hai")
    goal_id = r["goal"]["id"]
    agent = next(a for a in r["agents"] if a.get("office"))
    f = _upload(client, agent["id"], (Path(__file__).parent / "data" / "practice.xlsx").read_bytes(), "Excel_Practice.xlsx")
    for text in ("give me duplicate formula", "I am asking muje formula do dupliate delete krne ka"):
        r = say(client, text, goal_id)
        assert "COUNTIFS(" in r["reply"] and "Messy Data" in r["reply"], (text, r["reply"])
        card = next(a for a in r["agents"] if a.get("office"))
        assert card["formula"] and "COUNTIFS(" in card["formula"]["excel"]
        f = _wait(client, f["id"])
        assert f["report"]["summary"]["duplicate_rows"] == 2, text        # nothing was deleted
    r = say(client, "extra space hatane ka formula batao", goal_id)
    assert "TRIM(" in r["reply"], r["reply"]


def test_files_too_big_for_the_ram_are_refused_before_upload(client, ml_mode):
    from app.agents import office_agent as oa
    r = say(client, "excel file ke formula check karo")
    agent = next(a for a in r["agents"] if a.get("office"))
    mb = 1024 * 1024
    big = client.post("/api/office/files", json={"name": "huge.xlsx", "size": 17 * mb, "agent_id": agent["id"]})
    assert big.status_code == 400 and "Excel files up to 16 MB" in big.json()["detail"] and "CSV" in big.json()["detail"]
    ok = client.post("/api/office/files", json={"name": "rows.csv", "size": 20 * mb, "agent_id": agent["id"]})
    assert ok.status_code == 200
    assert oa.heavy(type("F", (), {"name": "a.xlsx", "size": 4 * mb})()) and not oa.heavy(type("F", (), {"name": "a.xlsx", "size": 2 * mb})())


@pytest.mark.parametrize("text,what", [
    ("my powerpoint file won't open", "PowerPoint"), ("outlook not sending emails error 0x800CCC0E", "Outlook"),
    ("my excel macro shows run-time error 1004", "macros"),
    ("I closed excel without saving, recover my file", "never saved"), ("convert this pdf to excel", "PDF"),
])
def test_what_it_cant_do_yet_is_said_plainly(client, ml_mode, text, what):
    reply = say(client, text)["reply"]
    assert "can't help with" in reply and what in reply and "fix it, and give you a fixed copy" not in reply, reply


def test_unknown_excel_task_gets_no_false_promise(client, ml_mode):
    reply = say(client, "excel mein GST calculate karna hai")["reply"]
    assert "fixed copy milegi" not in reply and "expert" in reply
    sci = say(client, "aadhaar numbers showing as 1.23E+11")["reply"]       # an Excel display problem, not Aadhaar advice
    assert "Drop the file" in sci and "Seva Kendra" not in sci


def test_charts_are_not_called_coming_soon(client, ml_mode):
    reply = say(client, "create a bar chart from my excel data")["reply"]
    assert "can't help with" not in reply and "charts" in reply


def test_old_answers_can_still_be_viewed_and_the_sheet_stays_in_focus(client, ml_mode):
    data = open("tests/data/practice.xlsx", "rb").read()
    r = say(client, "excel file ke formula check karo")
    goal, agent = r["goal"]["id"], next(a for a in r["agents"] if a.get("office"))
    _upload(client, agent["id"], data, name="practice.xlsx")
    first = say(client, "employee data sheet mein department wise total salary", goal)["reply"]
    say(client, "location wise?", goal)
    assert "Employee Data mein 40 records" in say(client, "how many rows?", goal)["reply"]      # stays on that sheet
    name = re.search(r'"(Answer \d+ - [^"]+)"', first).group(1)
    v = client.get(f"/api/office/answer?goal_id={goal}&name={name}").json()
    assert v["name"] == name and v["rows"] and "Sales" in str(v["rows"])                      # the first answer, after newer ones
    assert client.get(f"/api/office/answer?goal_id={goal}&name=Sales Data").status_code == 400   # only answer sheets
    dup = say(client, "show duplicates", goal)["reply"]
    assert "Answer" not in dup.split("Duplicates (Goal Engine)")[0]                            # its own sheets aren't data


def test_one_column_compared_with_another(client, ml_mode):
    data = open("tests/data/practice.xlsx", "rb").read()
    r = say(client, "excel file ke formula check karo")
    goal, agent = r["goal"]["id"], next(a for a in r["agents"] if a.get("office"))
    _upload(client, agent["id"], data, name="practice.xlsx")
    ask = say(client, "inventory sheet mein kaunse items reorder level se neeche hain", goal)
    assert "Reorder Level" in ask["reply"] and any("Closing Stock below Reorder Level" == o for o in ask.get("options", []))
    got = say(client, "Closing Stock below Reorder Level", goal)["reply"]
    assert "Printer" in got and "Router" in got and "Laptop" not in got


def _messy_without_calc_settings():
    """A workbook like ones made by other tools: no calculation settings in it (openpyxl's wb.calculation is None)."""
    import zipfile
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Invoice", "Region", "Qty", "Price", "Total"])
    rows = [("INV-1", "North", 2, 100), ("INV-2", "north", 1, "₹1,499"), ("INV-3", "NORTH ", 3, 50), ("INV-4", " South", 1, "200")]
    for i, (inv, reg, q, p) in enumerate(rows, start=2):
        ws.append([inv, reg, q, p, f"=C{i}*D{i}"])
    raw = io.BytesIO(); wb.save(raw)
    src, out = zipfile.ZipFile(io.BytesIO(raw.getvalue())), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(rb"<calcPr[^>]*/>", b"", data)
            z.writestr(item, data)
    return out.getvalue()


def test_clean_up_list_on_a_file_without_calc_settings(client, ml_mode):
    r = say(client, "excel file ke formula check karo")
    goal, agent = r["goal"]["id"], next(a for a in r["agents"] if a.get("office"))
    f = _upload(client, agent["id"], _messy_without_calc_settings(), name="sales.xlsx")
    assert f["status"] == "ready"
    reply = say(client, "Your task is to clean and correct this file: - Duplicate records - Extra spaces in text - Region values like "
                        "North, north, NORTH - Numbers stored as text - Currency symbol inside numeric data - Mixed date formats", goal)["reply"]
    assert "column widths" not in reply.lower() and "AttributeError" not in reply
    assert reply.startswith("Done in sales.xlsx") and "Extra spaces" in reply and "Text-numbers" in reply   # English, what really changed
    assert "Needs your look" in reply or "Dates written one way" in reply or "Text-numbers" in reply
    assert "Duplicate rows: not deleted" in reply or "Duplicate rows" not in reply
    assert say(client, "make a pivot of qty by region", goal)["reply"] and "AttributeError" not in say(client, "total qty by region", goal)["reply"]
