"""The Microsoft agent: the person gives a file, the agent finds what's wrong and fixes it (a copy), or compares two.

Files go to data/office/ in chunks (resumable, up to 1.5 GB), work runs in the background one big file at a
time, and the privacy clean-up deletes everything after an hour. No AI: app/officeknow (what's wrong) +
app/office (doing it).
"""

import os
import json
import re
import secrets
import shutil
import threading
from datetime import timedelta
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import office
from app.mind import packs
from app.models import Agent, Goal, OfficeFile, dumps, loads, utcnow

DATA = Path(__file__).resolve().parents[2] / "data" / "office"
MAX_SIZE = 1536 * 1024 * 1024
CHUNK_MAX = 8 * 1024 * 1024
FILE_TTL = timedelta(hours=1)          # after the last thing done with the file
FILE_MAX_AGE = timedelta(hours=8)      # whatever happens
BIG = 50 * 1024 * 1024
OPEN_FILES_PER_VISITOR = 6
class _ServerLane:
    """One big file at a time on the WHOLE server (all app processes), so a few huge uploads can't use up the RAM
    the server shares with other sites: a thread lock inside this process + a file lock across processes."""

    def __init__(self):
        self._threads = threading.Lock()
        self._fh = None

    def acquire(self):
        import fcntl

        self._threads.acquire()
        try:
            DATA.mkdir(parents=True, exist_ok=True)
            self._fh = open(DATA / ".big-file.lock", "a")
            fcntl.flock(self._fh, fcntl.LOCK_EX)        # waits while another process works on a big file
        except Exception:  # noqa: BLE001 - no lock file: still one at a time within this process
            self._fh = None

    def release(self):
        import fcntl

        try:
            if self._fh is not None:
                fcntl.flock(self._fh, fcntl.LOCK_UN)
                self._fh.close()
        finally:
            self._fh = None
            self._threads.release()


_big_lane = _ServerLane()                   # one big file at a time, so the VPS stays responsive
FILE_WORDS = packs.words("office", "about")            # the field's words come from its knowledge pack
ALLOWED = (".xlsx", ".xlsm", ".csv", ".txt")
STRONG_FILE = packs.words("office", "strong")
PROBLEM_WORDS = re.compile(r"\b(issues?|problems?|errors?|kharab|broken|galat|wrong|not working|nahi chal|help|fix|theek|check|messy|dikkat|gadbad)\b", re.I)


class OfficeAgentError(Exception):
    pass


# ------------------------------------------------------------------ chat side
def wants_office(goal: Goal | None, text: str, db: Session | None = None) -> bool:
    if goal is None:
        return False
    if db is not None:                                # a file is already open in this conversation: follow-up tasks go to it
        agent = goal_agent(db, goal)
        if agent is not None and open_file(db, agent) is not None and (
                UNDO_RE.search(text or "") or ALL_RE.search(text or "") or ops_in(text) or REPORT_RE.search(text or "")
                or (pending_offer(agent) and YES_RE.search(text or ""))):
            return True
    from app.officeknow import brain

    if any(packs.mentions(name, text) for name, p in packs.all_packs().items() if "office" in p.get("outranks", [])):
        return False                                  # "database se excel chahiye": that agent's words win (its pack says so)
    rec = brain.recognize(text)
    if rec["confidence"] >= 0.99:
        return True                                   # an Excel error value like #REF!
    if STRONG_FILE.search(text or ""):
        return True                                   # "my excel sheet has issues": it names Excel, that's enough
    if FILE_WORDS.search(text or "") and PROBLEM_WORDS.search(text or ""):
        return True                                   # "meri sheet kharab hai"
    return bool(FILE_WORDS.search(text or "")) and rec["confidence"] >= 0.2


def goal_agent(db: Session, goal: Goal) -> Agent | None:
    return db.scalar(select(Agent).where(Agent.goal_id == goal.id, Agent.recipe == "office").order_by(Agent.id.desc()))


def create_agent(db: Session, goal: Goal) -> Agent:
    agent = Agent(visitor_id=goal.visitor_id, goal_id=goal.id, recipe="office", status="needs_input", inputs=dumps({}))
    db.add(agent)
    db.flush()
    return agent


SURE = 0.5


def plan_reply(rec: dict, hi: bool) -> str:
    from app.officeknow import brain

    seen = brain.describe(rec, hi) if rec.get("confidence", 0) >= SURE else ""
    lead = (f"Samjha: {seen}. " if hi else f"Understood: {seen}. ") if seen else ""
    if rec.get("kind") == "formula" and seen:
        return lead + ("Formula card mein hai. File daaloge to main khud laga dunga." if hi
                       else "The formula is on the card. Give me the file and I'll put it in for you.")
    return lead + ("File card mein daalo. Main khud check karke theek kar dunga, aur aapko fixed copy milegi." if hi
                   else "Drop the file on the card. I'll check it myself, fix it, and give you a fixed copy.")


def office_dict(db: Session, agent: Agent, hi: bool) -> dict:
    inputs = loads(agent.inputs, {}) or {}
    files = list(db.scalars(select(OfficeFile).where(OfficeFile.agent_id == agent.id, OfficeFile.status != "removed")
                            .order_by(OfficeFile.id.desc()).limit(6)))
    has_file = any(f.status in ("ready", "queued", "working") for f in files)
    if has_file:
        inputs = {k: v for k, v in inputs.items() if k != "brain"}     # the agent works on the file; no generic formula
    return {"id": agent.id, "lang": "hi" if hi else "en", "recipe": "office", "office": True, "status": agent.status,
            "name": "Excel / file ka kaam" if hi else "Excel and file work", "access": "upload", "available": True,
            "blocks": [], "runs": [], "brain": _brain_view(inputs.get("brain"), hi), "formula": inputs.get("file_formula") if has_file and inputs.get("file_formula") else _formula_view(inputs.get("brain"), hi),
            "files": [file_dict(f) for f in files]}


def _brain_view(rec, hi):
    if not rec or not rec.get("kind"):
        return None
    from app.officeknow import brain

    return {**rec, "text": brain.describe(rec, hi)}


def _formula_view(rec, hi):
    if not rec or rec.get("kind") != "formula":
        return None
    from app import officeknow

    f = next((x for x in officeknow.load("formulas") if x["id"] == rec["id"]), None)
    if not f:
        return None
    return {"title": f["title_hi" if hi else "title_en"], "explain": f["explain_hi" if hi else "explain_en"],
            "excel": f["templates"].get("excel365", ""), "sheets": f["templates"].get("google_sheets", ""), "needs": f["needs"]}


# ------------------------------------------------------------------ files
def steps(f: OfficeFile) -> list[dict]:
    return loads(f.versions, None) or [{"path": f.path, "ops": [], "changes": {}}]


def file_dict(f: OfficeFile) -> dict:
    st = steps(f)
    return {"id": f.token, "name": f.name, "size": f.size, "received": f.received, "status": f.status, "task": f.task,
            "report": loads(f.report, None), "result": loads(f.result, None), "error": f.error,
            "steps": [{"ops": x["ops"], "changes": x["changes"]} for x in st[1:]], "before": st[0].get("summary"),
            "download": len(st) > 1 and bool(f.path and os.path.exists(f.path)),
            "compare_download": bool(f.result_path and os.path.exists(f.result_path)), "result_name": f.result_name}


def start_upload(db: Session, visitor_id: str, agent_id: int | None, name: str, size: int) -> OfficeFile:
    name = os.path.basename(name or "file.xlsx")[:150]
    if not name.lower().endswith(ALLOWED):
        raise OfficeAgentError("Excel (.xlsx) ya CSV file daalo. Purani .xls ho to Excel mein Save As .xlsx karke daalo.")
    if not 0 < size <= MAX_SIZE:
        raise OfficeAgentError("File 1.5 GB tak ki ho sakti hai.")
    DATA.mkdir(parents=True, exist_ok=True)
    if shutil.disk_usage(DATA).free < size * 3 + 2 * 1024**3:
        raise OfficeAgentError("Server abhi itni badi file nahi le sakta. Thodi der baad try karo.")
    open_now = db.scalars(select(OfficeFile).where(OfficeFile.visitor_id == visitor_id, OfficeFile.status.not_in(("removed",)))).all()
    if len(open_now) >= OPEN_FILES_PER_VISITOR:
        raise OfficeAgentError("Ek saath 6 files tak. Purani files 1 ghante mein apne aap hat jaati hain.")
    token = secrets.token_urlsafe(24)
    f = OfficeFile(token=token, visitor_id=visitor_id, agent_id=agent_id, name=name, size=size,
                   path=str(DATA / f"{token}{os.path.splitext(name)[1].lower()}"))
    db.add(f)
    db.commit()
    open(f.path, "wb").close()
    return f


def own_file(db: Session, token: str, visitor_id: str) -> OfficeFile:
    f = db.scalar(select(OfficeFile).where(OfficeFile.token == token))
    if not f or f.visitor_id != visitor_id or f.status == "removed":
        raise OfficeAgentError("File nahi mili (1 ghante baad files apne aap hat jaati hain).")
    return f


def add_chunk(db: Session, f: OfficeFile, offset: int, data: bytes) -> OfficeFile:
    if f.status != "uploading":
        raise OfficeAgentError("Upload pehle hi poora ho chuka hai.")
    if offset != f.received:
        return f                                     # the browser resumes from f.received
    if len(data) > CHUNK_MAX or f.received + len(data) > f.size:
        raise OfficeAgentError("Galat hissa aaya. Dobara upload karo.")
    with open(f.path, "r+b") as out:
        out.seek(offset)
        out.write(data)
    f.received += len(data)
    db.commit()
    return f


def finish_upload(db: Session, f: OfficeFile) -> OfficeFile:
    if f.received != f.size:
        raise OfficeAgentError("File poori nahi aayi. Upload dobara chalu karo, wahin se aage badhega.")
    try:
        office.kind_of(f.path)
    except office.OfficeError as exc:
        f.status, f.error = "error", str(exc)
        db.commit()
        return f
    return _start(db, f, "scan", {})


def start_fix(db: Session, f: OfficeFile, ops: list[str], sheets: list[str] | None = None) -> OfficeFile:
    ops = [o for o in ops if o in office.FIX_OPS]
    if f.status not in ("ready", "error") or not ops:
        raise OfficeAgentError("Pehle check poora hone do.")
    return _start(db, f, "fix", {"ops": ops, "sheets": sheets or []})


def start_compare(db: Session, a: OfficeFile, b: OfficeFile, key: str = "") -> OfficeFile:
    if a.status not in ("ready", "error") or b.status not in ("ready", "error"):
        raise OfficeAgentError("Dono files ka upload aur check poora hone do.")
    return _start(db, a, "compare", {"other": b.id, "key": key})


def _start(db: Session, f: OfficeFile, task: str, args: dict) -> OfficeFile:
    f.status, f.task, f.error, f.args = "queued", task, None, dumps(args)
    db.commit()
    threading.Thread(target=_work, args=(f.id, task, args), daemon=True).start()
    return f


def _work(file_id: int, task: str, args: dict) -> None:
    from app.db import _SessionLocal, get_engine

    get_engine()
    with _SessionLocal() as db:
        f = db.get(OfficeFile, file_id)
        big = f.size > BIG or task == "compare"
        lane = _big_lane if big else None
        if lane:
            lane.acquire()
        try:
            f.status = "working"
            db.commit()
            if task == "scan":
                rep = office.scan(f.path)
                f.report = dumps(rep)
                if not f.versions:                       # remember how the file looked on arrival, for before/after
                    f.versions = dumps([{"path": f.path, "ops": [], "changes": {}, "summary": rep["summary"]}])
            elif task == "fix":
                # every fix builds on the latest version; the file stays open for the next task
                st = steps(f)
                root, ext = os.path.splitext(st[0]["path"])
                out = f"{root}.v{len(st)}{ext}"
                res = office.fix(f.path, out, args["ops"], args.get("sheets") or None)
                st.append({"path": out, "ops": args["ops"], "changes": res["changes"]})
                f.versions, f.path = dumps(st), out
                f.result = dumps({"task": "fix", **res})
                f.report = dumps(office.scan(out))
            elif task == "extract":
                st = steps(f)
                root, ext = os.path.splitext(st[0]["path"])
                res = office.extract_duplicates(f.path, f"{root}.v{len(st)}{ext}", args.get("sheets") or None)
                out = res.pop("out")
                st.append({"path": out, "ops": ["list_duplicates"], "changes": {"list_duplicates": res["count"]}})
                f.versions, f.path = dumps(st), out
                f.result = dumps({"task": "extract", **res})
                f.report = dumps(office.scan(out))
            elif task == "ask":
                from app.office import ask as asker

                st = steps(f)
                n = 1 + sum(1 for x in st if "answer" in x["ops"])
                root, ext = os.path.splitext(st[0]["path"])
                out = f"{root}.v{len(st)}{ext}"
                res = asker.run(args["plan"], f.path, out, args.get("hi", True), n)
                if os.path.exists(out):
                    st.append({"path": out, "ops": ["answer"], "changes": {"answer": 1}})
                    f.versions, f.path = dumps(st), out
                f.result = dumps({"task": "ask", **res})
            elif task == "undo":
                st = steps(f)
                if len(st) > 1:
                    last = st.pop()
                    _rm(last["path"])
                    f.versions, f.path = dumps(st), st[-1]["path"]
                    f.result = dumps({"task": "undo", "undone": last["ops"]})
                f.report = dumps(office.scan(f.path))
            elif task == "compare":
                other = db.get(OfficeFile, args["other"])
                out = f.path + ".compare.xlsx"
                summary = office.compare(f.path, other.path, out, key=args.get("key", ""))
                f.result = dumps({"task": "compare", "other": other.name, **summary})
                _rm(f.result_path)
                f.result_path, f.result_name = out, _named(f.name, "compare", ".xlsx")
            f.status = "ready"
        except office.OfficeError as exc:
            f.status, f.error = "error", str(exc)
        except Exception as exc:  # noqa: BLE001 - a broken file must not crash the worker
            f.status, f.error = "error", f"File padhi nahi ja saki ({type(exc).__name__}). Excel mein khol kar Save As .xlsx karke dobara daalo."
        finally:
            if lane:
                lane.release()
            db.commit()


def _named(name: str, suffix: str, ext: str = "") -> str:
    base, e = os.path.splitext(name)
    return f"{base} ({suffix}){ext or e}"


def _rm(p) -> None:
    try:
        if p:
            os.remove(p)
    except OSError:
        pass


def start_undo(db: Session, f: OfficeFile) -> OfficeFile:
    if f.status not in ("ready", "error") or len(steps(f)) < 2:
        raise OfficeAgentError("Wapas lene ke liye koi badlav nahi hai.")
    return _start(db, f, "undo", {})


def remove_file(f: OfficeFile) -> None:
    from app.office.sheetdata import forget

    for p in {f.path, f.result_path, *(x["path"] for x in steps(f))}:
        if p:
            forget(p)
        _rm(p)
    f.path = f.result_path = None
    f.report = f.result = None
    f.status = "removed"


def maintain(db: Session, now=None) -> int:
    now = now or utcnow()
    n = 0
    from sqlalchemy import or_

    for f in db.scalars(select(OfficeFile).where(OfficeFile.status != "removed",
                                                 or_(OfficeFile.updated_at < now - FILE_TTL, OfficeFile.created_at < now - FILE_MAX_AGE))):
        if f.status in ("queued", "working") and f.created_at > now - FILE_MAX_AGE:
            continue                                  # still working on a very big file
        remove_file(f)
        n += 1
    for f in db.scalars(select(OfficeFile).where(OfficeFile.status == "removed", OfficeFile.created_at < now - timedelta(days=1))):
        db.delete(f)
    db.commit()
    return n


_resume_lock = None


def _resume_owner() -> bool:
    """Only one server process resumes work: the first to take this lock keeps it for its whole life."""
    global _resume_lock
    if _resume_lock is not None:
        return True
    import fcntl

    DATA.mkdir(parents=True, exist_ok=True)
    fh = open(DATA / ".resume.lock", "w")
    try:
        fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        fh.close()
        return False
    _resume_lock = fh
    return True


def resume_interrupted(boot) -> int:
    """After a restart (an update, a crash), carry on with work that was running, in one process only."""
    from sqlalchemy import update

    from app.db import _SessionLocal, get_engine

    if not _resume_owner():
        return 0
    get_engine()
    n = 0
    with _SessionLocal() as db:
        ids = list(db.scalars(select(OfficeFile.id).where(OfficeFile.status.in_(("queued", "working")), OfficeFile.updated_at < boot)))
        for fid in ids:
            claimed = db.execute(update(OfficeFile).where(OfficeFile.id == fid, OfficeFile.status.in_(("queued", "working")),
                                                          OfficeFile.updated_at < boot).values(status="queued", updated_at=boot)).rowcount
            db.commit()
            if claimed:
                f = db.get(OfficeFile, fid)
                threading.Thread(target=_work, args=(f.id, f.task or "scan", loads(f.args, {}) or {}), daemon=True).start()
                n += 1
    return n


# ------------------------------------------------------------------ tasks on the open file, typed in the chat
# Keywords are matched loosely, so typos still work ("dupicate", "slaes sheet").
OP_WORDS = {
    "trim_spaces": ["space", "spaces", "spacing", "jagah"],
    "remove_duplicates": ["duplicate", "duplicates", "duplicated", "repeat", "repeated", "dobara", "dohra"],
    "fix_case": ["capital", "capitals", "lowercase", "uppercase", "case", "spelling"],
    "text_to_number": ["numbers", "number", "numeric"],
    "fix_broken_refs": ["#ref", "ref", "broken"],
}
FIND_WORDS = {"find", "show", "check", "list", "dhoondo", "dhundo", "dhoond", "dhund", "dikhao", "dikha", "batao", "bata", "kitne",
              "kitni", "kaun", "kaunse", "konse", "kahan", "where", "which", "how", "many", "any", "pata", "highlight", "search"}
DO_WORDS = {"remove", "delete", "hatao", "hata", "hatado", "hataye", "hatana", "fix", "theek", "thik", "saaf", "clean", "correct",
            "sudharo", "badlo", "drop", "mitao", "delete karo"}
YES_RE = re.compile(r"^\s*(haan|ha|haa|han|yes|yep|ok|okay|theek hai|thik hai|kar do|karo|hata do|hatao|do it|go ahead|please)\b", re.I)
FORMULA_RE = re.compile(r"\b(formula|formule|formulae|fourmula|farmula|function)\b", re.I)
UNDO_RE = re.compile(r"\b(undo|wapas|vapas|pehle jaisa|revert|purana wala)\b", re.I)
ALL_RE = re.compile(r"\b(sab|saari|saare|everything|all)\b.{0,20}\b(theek|thik|fix|saaf|clean)|\b(fix|clean|correct)\b.{0,12}\b(everything|all|it all|them all|the file|my file)\b", re.I)
REPORT_RE = re.compile(r"\b(kya galat|kya problem|what.{0,10}wrong|check (karo|my|the|this)|problems?|errors?|galti|galtiyan|issues?)\b", re.I)
_OP_NAMES = {"trim_spaces": ("extra space", "extra spaces"), "remove_duplicates": ("duplicate rows", "duplicate rows"),
             "fix_case": ("capital/small letters", "capital letters"), "text_to_number": ("text wale numbers", "text-numbers"),
             "fix_broken_refs": ("toote formula", "broken formulas")}


def _words(text: str) -> list[str]:
    from app.ml.normalize import tokens

    return tokens(text or "")


def _close(word: str, choices) -> bool:
    import difflib

    return len(word) >= 3 and bool(difflib.get_close_matches(word, list(choices), n=1, cutoff=0.8 if len(word) > 4 else 0.9))


def ops_in(text: str) -> list[str]:
    ws = _words(text)
    found = [op for op, kws in OP_WORDS.items() if any(w in kws or _close(w, kws) for w in ws)]
    if "text_to_number" in found and not re.search(r"\b(text|number|numbers|total|sum|jod)\b", text or "", re.I):
        found.remove("text_to_number")
    if "fix_broken_refs" in found and not re.search(r"(#?ref|formula|broken)", text or "", re.I):
        found.remove("fix_broken_refs")
    return found


def sheets_in(text: str, report: dict) -> list[str]:
    """'from the sales sheet' -> ['Sales Data'] (typos allowed)."""
    ws = set(_words(text))
    out = []
    for sh in report.get("sheets", []):
        name = sh["name"]
        key = [w for w in re.findall(r"[a-z0-9]+", name.lower()) if w not in ("data", "sheet", "sheets", "the", "and", "of")]
        if key and all(any(w == k or _close(w, [k]) for w in ws) for k in key):
            out.append(name)
    return out


def wants_do(text: str) -> bool:
    """Only a clear 'remove / fix / hatao' changes the data. Anything else just finds and shows."""
    ws = _words(text)
    order = re.search(r"\b(ek jaise|ek jaisa|same kar|kar do|kardo|kar dijiye|badal do|badlo|convert|bana do|banao|sudhar)\b", text or "", re.I)
    return (bool(order) or any(w in DO_WORDS or _close(w, DO_WORDS) for w in ws)) and not re.search(
        r"\b(alag|separate|new sheet|nayi sheet|list|report|dikhao|check)\b", text or "", re.I)


def wants_find(text: str) -> bool:
    ws = _words(text)
    finding = any(w in FIND_WORDS for w in ws) or "?" in (text or "") or bool(re.search(r"\bhai kya\b|\bhain kya\b", text or "", re.I))
    doing = any(w in DO_WORDS or _close(w, DO_WORDS) for w in ws)
    return finding and not doing


def _findings(report: dict, op: str, sheets: list[str], hi: bool) -> tuple[str, int]:
    """Plain words for what the scan found, for one kind of problem, in the asked sheets."""
    lines, total = [], 0
    for sh in report.get("sheets", []):
        if sheets and sh["name"] not in sheets:
            continue
        if op == "remove_duplicates":
            n, cells = sh.get("duplicate_rows", 0), [f"row {r} ({'row' if not hi else 'row'} {m} {'jaisi' if hi else 'again'})"
                                                      for r, m in sh.get("duplicate_list", [])[:6]]
        elif op == "trim_spaces":
            n, cells = sh.get("spaces", 0), sh.get("space_cells", [])[:8]
        elif op == "fix_case":
            n, cells = sh.get("case_mixed", 0), sh.get("case_cells", [])[:8]
        elif op == "text_to_number":
            n, cells = sh.get("text_numbers", 0), sh.get("text_number_cells", [])[:8]
        else:
            n, cells = sh.get("broken_ref_count", 0), sh.get("broken_refs", [])[:8]
        total += n
        if n:
            lines.append(f"{sh['name']}: {n} ({', '.join(cells)}{' …' if n > len(cells) else ''})")
        elif sheets:
            lines.append(f"{sh['name']}: " + ("koi nahi" if hi else "none"))
    return "; ".join(lines), total


def open_file(db: Session, agent: Agent) -> OfficeFile | None:
    return db.scalar(select(OfficeFile).where(OfficeFile.agent_id == agent.id, OfficeFile.status.in_(("ready", "queued", "working")))
                     .order_by(OfficeFile.updated_at.desc()))


def pending_offer(agent: Agent) -> dict | None:
    return (loads(agent.inputs, {}) or {}).get("offer")


def _set_offer(agent: Agent, offer: dict | None) -> None:
    inputs = loads(agent.inputs, {}) or {}
    if offer:
        inputs["offer"] = offer
    else:
        inputs.pop("offer", None)
    agent.inputs = dumps(inputs)


def chat_task(db: Session, agent: Agent, text: str, hi: bool) -> str | None:
    """Work on the file that's already open, from what the person typed. Returns the reply, or None if it isn't a file task."""
    f = open_file(db, agent)
    if f is None:
        return None
    if f.status != "ready":
        return "Pichla kaam abhi chal raha hai, khatam hote hi agla karta hoon. Card pe dikhega." if hi else "Still working on the last task; the card shows when it's done."
    report = loads(f.report, {}) or {}
    offer = pending_offer(agent)
    from app import mind

    if offer and YES_RE.search(text or "") and not ops_in(text):
        mind.set_pattern("fix")
        _set_offer(agent, None)
        f = start_fix(db, f, offer["ops"], offer.get("sheets"))
        return _doing(f, offer["ops"], offer.get("sheets"), hi)
    if UNDO_RE.search(text or ""):
        mind.set_pattern("undo")
        try:
            start_undo(db, f)
        except OfficeAgentError as exc:
            return str(exc)
        return "Aakhri badlav wapas le raha hoon." if hi else "Undoing the last change."
    ops = ops_in(text)
    sheets = sheets_in(text, report)
    if FORMULA_RE.search(text or "") and ops:
        # "duplicate ka formula do": they want the formula for THEIR file, not a change
        mind.set_pattern("formula")
        return formula_reply(agent, f, report, ops or ["remove_duplicates"], sheets, hi)
    if not ops and ALL_RE.search(text or ""):
        s = report.get("summary", {})
        ops = [op for key, op in (("spaces", "trim_spaces"), ("text_numbers", "text_to_number"), ("broken_refs", "fix_broken_refs"),
                                  ("case_mixed", "fix_case"), ("duplicate_rows", "remove_duplicates")) if s.get(key)]
        mind.set_pattern("fix")
        if not ops:
            return "Is file mein theek karne layak kuch nahi bacha." if hi else "Nothing left to fix in this file."
        return _doing(start_fix(db, f, ops, sheets), ops, sheets, hi)
    if not ops and REPORT_RE.search(text or ""):
        ops = [op for op in _OP_NAMES]          # "kya galat hai?" -> everything the scan knows
        mind.set_pattern("review")
        return _answer(agent, report, ops, sheets, hi, f)
    if not ops:
        return ask_task(db, agent, f, text, hi)
    mind.set_pattern("fix" if wants_do(text) else "dups" if ops == ["remove_duplicates"] else "review")
    if ops == ["remove_duplicates"] and not wants_do(text):
        # "duplicates dhund ke do", "alag sheet mein daalo": list them in a new sheet, delete nothing
        _set_offer(agent, {"ops": ["remove_duplicates"], "sheets": sheets})
        where = f" ({', '.join(sheets)})" if sheets else ""
        if f.size <= BIG:                        # small file: do it now and answer with the result
            f.status, f.task, f.args = "queued", "extract", dumps({"sheets": sheets})
            db.commit()
            _work(f.id, "extract", {"sheets": sheets})
            db.refresh(f)
            r = loads(f.result, {}) or {}
            if f.status == "ready" and r.get("task") == "extract":
                return _extract_reply(f, r, sheets, report, hi)
        _start(db, f, "extract", {"sheets": sheets})
        return (f"{f.name}{where} ke saare duplicate ek nayi sheet \"Duplicates (Goal Engine)\" mein daal raha hoon: kaunsi row, "
                "kis row jaisi, aur kaunsa ID/invoice/email do baar aaya. Aapka data nahi badlega. Ready hote hi card pe Download karo. "
                "Baad mein hatane hain to 'duplicate hatao' likho." if hi else
                f"Putting every duplicate in {f.name}{where} into a new sheet \"Duplicates (Goal Engine)\": which row, which row it repeats, "
                "and any ID/invoice/email used twice. Your data isn't changed. Download it from the card when it's ready. "
                "To remove them later, type 'remove duplicates'.")
    if not wants_do(text):
        return _answer(agent, report, ops, sheets, hi, f)
    return _doing(start_fix(db, f, ops, sheets), ops, sheets, hi)


def _answer(agent, report, ops, sheets, hi, f) -> str:
    """Find, don't change: say exactly where, then offer to fix it."""
    parts, fixable = [], []
    for op in ops:
        text, n = _findings(report, op, sheets, hi)
        if n or sheets:
            parts.append(f"{_OP_NAMES[op][0 if hi else 1].capitalize()}: {text or ('koi nahi' if hi else 'none')}")
        if n:
            fixable.append(op)
    where = f" ({', '.join(sheets)})" if sheets else ""
    if not fixable:
        others = []
        if sheets:                               # nothing in the asked sheet; maybe elsewhere
            for op in ops:
                t, n = _findings(report, op, [], hi)
                if n:
                    others.append(f"{_OP_NAMES[op][0 if hi else 1]}: {t}")
        what = ", ".join(_OP_NAMES[o][0 if hi else 1] for o in ops) if len(ops) < 5 else ("koi galti" if hi else "problems")
        place = ", ".join(sheets) if sheets else f.name
        msg = (f"{place} mein ab koi {what} nahi hai." if hi else f"No {what} in {place}.")
        if others:
            msg += (" Baaki sheets mein: " if hi else " In other sheets: ") + "; ".join(others) + "."
            _set_offer(agent, {"ops": [o for o in ops if _findings(report, o, [], hi)[1]], "sheets": []})
            msg += " Theek kar doon? 'haan' likho." if hi else " Shall I fix those? Type 'yes'."
        return msg
    _set_offer(agent, {"ops": fixable, "sheets": sheets})
    return ("Mila" if hi else "Found") + f"{where}. " + ". ".join(parts) + ". " + (
        "Theek kar doon? 'haan' likho, ya card pe Theek karo dabao." if hi else "Shall I fix it? Type 'yes', or press Fix it on the card.")


def _doing(f, ops, sheets, hi) -> str:
    what = ", ".join(_OP_NAMES[o][0 if hi else 1] for o in ops)
    where = f" ({', '.join(sheets)})" if sheets else ""
    return (f"{f.name}{where} mein {what} theek kar raha hoon. Pichle badlav bhi saath rahenge." if hi
            else f"Fixing {what} in {f.name}{where}. Earlier changes stay in.")


def cant_yet(f: OfficeFile, hi: bool) -> str:
    return (f"{f.name} khuli hai. Is pe main abhi ye kar sakta hoon: extra space, duplicate rows, capital/small letters, text wale numbers, "
            "toote formula dhoondna aur theek karna, aur doosri file se compare. Aapka ye kaam main abhi seekh raha hoon; "
            "tab tak inmein se koi kaam bolo." if hi else
            f"{f.name} is open. Right now I can find and fix extra spaces, duplicate rows, mixed capitals, numbers saved as text, "
            "broken formulas, and compare with another file. I'm still learning this one; meanwhile ask for any of those.")


def _extract_reply(f: OfficeFile, r: dict, sheets: list[str], report: dict, hi: bool) -> str:
    where = ", ".join(sheets) if sheets else f.name
    if not r["count"]:
        msg = f"{where} mein koi duplicate nahi mila." if hi else f"No duplicates in {where}."
        if sheets:                                           # say where they ARE, if anywhere
            other = [f"{s['name']}: {s['duplicate_rows']}" for s in report.get("sheets", []) if s.get("duplicate_rows") and s["name"] not in sheets]
            if other:
                msg += (f" Doosri sheets mein hain ({', '.join(other)}). 'poori file ke duplicate dikhao' likho." if hi
                        else f" Other sheets have some ({', '.join(other)}). Type 'show duplicates in the whole file'.")
        return msg
    parts = []
    for g in r["groups"][:5]:
        parts.append((f"{g['sheet']}: {g['count']} poori rows do baar" if hi else f"{g['sheet']}: {g['count']} whole rows repeated")
                     if g["what"] == "row" else
                     (f"{g['sheet']}: {g['count']} baar same {g['what']}" if hi else f"{g['sheet']}: same {g['what']} {g['count']} times"))
    return ((f"{r['count']} duplicate mile ({'; '.join(parts)}). Sab ek nayi sheet \"Duplicates (Goal Engine)\" mein daal diye: kaunsi row, "
             "kis row jaisi. Aapka data nahi badla. Card pe Download karo. Hatane hain to 'duplicate hatao' likho.") if hi else
            (f"Found {r['count']} duplicates ({'; '.join(parts)}). They're all in a new sheet \"Duplicates (Goal Engine)\": which row, "
             "and which row it repeats. Your data wasn't changed. Download it from the card. To remove them, type 'remove duplicates'."))


# ------------------------------------------------------------------ "formula do": a formula written for the person's own sheet
def _sheet_with(report: dict, key: str, sheets: list[str]) -> dict | None:
    cands = [sh for sh in report.get("sheets", []) if not sheets or sh["name"] in sheets]
    return max(cands, key=lambda sh: sh.get(key, 0) if not isinstance(sh.get(key), list) else len(sh[key]), default=None)


def _col_of(cell: str) -> str:
    return re.match(r"[A-Z]+", cell).group(0) if cell else "B"


def formula_reply(agent: Agent, f: OfficeFile, report: dict, ops: list[str], sheets: list[str], hi: bool) -> str:
    from openpyxl.utils import get_column_letter

    op = ops[0]
    sh = _sheet_with(report, {"remove_duplicates": "duplicate_rows", "trim_spaces": "spaces", "fix_case": "case_mixed",
                              "text_to_number": "text_numbers"}.get(op, "rows"), sheets)
    if sh is None:
        return None
    name, last = sh["name"], max(sh["rows"], 2)
    header = sh.get("header", [])
    q = f"'{name}'!" if re.search(r"\W", name) else f"{name}!"
    if op == "remove_duplicates":
        from app.office import _ROWNO

        cols = [get_column_letter(i + 1) for i, h in enumerate(header) if h and not _ROWNO.match(str(h))][:6] or ["A"]
        new_col = get_column_letter(len(header) + 1)
        parts = ",".join(f"${c}$2:${c}2,{c}2" for c in cols)
        excel = f'=IF(COUNTIFS({parts})>1,"Duplicate","")'
        names = ", ".join(header[ord(c) - 65] for c in cols if len(c) == 1 and ord(c) - 65 < len(header))
        unique = f"=UNIQUE({q}{cols[0]}1:{cols[-1]}{last})"            # without a row-number column, or nothing looks repeated
        title = (f"Duplicate pehchanne ka formula ({name})" if hi else f"Formula to mark duplicates ({name})")
        explain = (f"{new_col}2 mein daalo aur {new_col}{last} tak neeche kheencho. Jo row upar wali kisi row jaisi hai ({names} sab same), "
                   f"uspe \"Duplicate\" likha aayega; pehli wali nahi. Phir filter lagao, \"Duplicate\" chuno aur wo rows delete karo. "
                   f"Excel 365 mein bina duplicate ki poori list ek nayi sheet mein: {unique}" if hi else
                   f"Put it in {new_col}2 and drag down to {new_col}{last}. A row that repeats an earlier row ({names} all the same) "
                   f"shows \"Duplicate\"; the first one doesn't. Then filter on \"Duplicate\" and delete those rows. "
                   f"In Excel 365, the whole list without duplicates on a new sheet: {unique}")
        tail = (" Excel mein formula khud kuch delete nahi karta; main ye aapke liye kar sakta hoon: 'duplicate hatao' likho." if hi
                else " A formula can't delete rows by itself in Excel; I can do it for you: type 'remove duplicates'.")
    else:
        cell = {"trim_spaces": sh.get("space_cells"), "fix_case": sh.get("case_cells"), "text_to_number": sh.get("text_number_cells")}.get(op) or ["B2"]
        col = _col_of(cell[0])
        fx = {"trim_spaces": f"=TRIM({col}2)", "fix_case": f"=PROPER({col}2)", "text_to_number": f'=VALUE(SUBSTITUTE(SUBSTITUTE({col}2,",",""),"₹",""))',
              "fix_broken_refs": None}[op]
        if fx is None:
            return ("Toote formula (#REF!) ka koi ek formula nahi hota; main inhe khud theek kar deta hoon: 'formula theek karo' likho." if hi
                    else "There's no single formula for broken #REF! formulas; I can rebuild them for you: type 'fix formulas'.")
        excel = fx
        title = {"trim_spaces": ("Extra space hatane ka formula", "Formula to remove extra spaces"),
                 "fix_case": ("Ek jaise capital letters ka formula", "Formula for consistent capitals"),
                 "text_to_number": ("Text ko number banane ka formula", "Formula to turn text into numbers")}[op][0 if hi else 1] + f" ({name}, {col})"
        explain = (f"Ek khaali column mein row 2 pe daalo aur {last} tak kheencho; phir Copy > Paste Special > Values se {col} column pe daal do." if hi
                   else f"Put it in an empty column on row 2 and drag down to row {last}; then Copy > Paste Special > Values over column {col}.")
        tail = (" Ya main seedha file mein kar deta hoon: 'theek karo' likho." if hi else " Or I can do it in the file directly: type 'fix it'.")
    inputs = loads(agent.inputs, {}) or {}
    inputs["file_formula"] = {"title": title, "excel": excel, "explain": explain, "sheets": "", "needs": []}
    agent.inputs = dumps(inputs)
    return f"{title}: {excel}  {explain}{tail}"


# ------------------------------------------------------------------ any other question about the open file
TEACH = ("Agent ko sikhao", "Teach the agent")


LAST_FORMULA_RE = re.compile(r"^\s*(give me |show me |send me |what is |what's )?(the |its |iska |uska |is ka )?(excel )?"
                             r"(formula|formule|fourmula|farmula)( (for|of) (this|that|it|the (last )?answer))?( do| dena| batao| bhi do| kya hai| kya hoga)?\s*[?.!]*\s*$"
                             r"|^\s*(iska|uska|is ka) formula\b", re.I)


def ask_task(db: Session, agent: Agent, f: OfficeFile, text: str, hi: bool):
    """Totals, counts, X-wise, lists, lookups, aging, combining, highlighting... on the person's own sheets.
    Returns the reply, or (question, options) when one detail is missing. Never a random answer.

    Before answering it checks the plan against the question and the person: the sheet they named, whether they
    want changes or only the answer, formulas only when they want them, and a different reading when the same
    question comes back because the last answer missed."""
    from app import mind
    from app.office import mind as mo
    from app.office import ask as asker
    from app.office.sheetdata import load_tables

    m = mind.now()
    sig, prof = m.get("signals", {}), m.get("profile", {})
    _set_offer(agent, None)
    inputs = loads(agent.inputs, {}) or {}
    pending = inputs.pop("pending_q", None)
    last_plan = inputs.get("mind_plan")
    if (text or "").strip() in TEACH and pending:
        from app import learning

        learning.ask_to_learn(db, "office", pending)
        agent.inputs = dumps(inputs)
        db.commit()
        return ("Note kar liya. Ye kaam agent ki learning list mein chala gaya; seekhte hi ye bhi karega. Tab tak koi aur kaam bolo."
                if hi else "Noted. It's on the agent's learning list; once learned it'll do this too. Meanwhile, ask for something else.")
    # "give me the formula" right after an answer: the formula of that answer
    last = loads(f.result, {}) or {}
    if LAST_FORMULA_RE.search(text or "") and last.get("task") == "ask" and last.get("formula"):
        f.result = dumps({**last, "show_formula": True})
        agent.inputs = dumps(inputs)
        db.commit()
        mind.set_pattern("formula", last_plan)
        return (f"Pichle jawab ka formula: {last['formula']}" if hi else f"The formula for the last answer: {last['formula']}")
    if pending and len((text or "").split()) <= 5:
        text = f"{pending} {text}"                        # the answer to "which sheet / which column?"
    tables = load_tables(f.path)
    m["tables"] = tables
    from app.office import understand as U

    try:
        plan = asker.understand(text, tables)
    except U.Missing as miss:
        # not in this file and can't be worked out from it: say so plainly, with what it does have
        agent.inputs = dumps(inputs)
        db.commit()
        mind.set_pattern("missing", None)
        try:
            from app import learning

            learning.ask_to_learn(db, "office", text, guess="not in file: " + miss.concept)
        except ValueError:
            pass
        return miss.hi if hi else miss.en
    except asker.NeedMore as need:
        inputs["pending_q"] = text
        agent.inputs = dumps(inputs)
        db.commit()
        opts = list(need.options)
        msg = need.hi if hi else need.en
        try:
            # it couldn't answer this directly: the sentence (scrubbed, nobody attached) goes to the learning list,
            # so the agent learns to answer it straight away next time
            from app import learning

            learning.ask_to_learn(db, "office", text, guess="needed: " + (need.en[:60] if need.en else ""))
        except ValueError:
            pass
        if need.hi.startswith("Ye poori tarah samajh nahi aaya"):
            # not understood: offer the closest things this file can answer, not a dead end
            near = mo.closest_options(text, tables, hi)
            if near:
                msg = ("Ye poori tarah samajh nahi aaya. Is file pe inme se kuch chahiye?" if hi
                       else "I didn't fully get that. Did you mean one of these?")
                opts = near
            opts = opts[:6] + ["Agent ko sikhao" if hi else "Teach the agent"]
        mind.set_pattern("other", None)
        mind.asked_back()
        return msg, opts
    # the named sheet is the sheet (a check on the plan, whatever job it became)
    named = asker.named_sheet(text, tables)
    if named is not None and plan.get("table") not in (named.name, None) and plan.get("job") not in ("lookup", "combine", "reconcile"):
        try:
            plan = asker.understand(text, [named])
        except asker.NeedMore:
            pass
    # the same question again after an answer: don't repeat that answer; offer the other readings
    if (sig.get("repeat") or sig.get("frustrated")) and last_plan and mo._same(plan, last_plan) and not plan.get("derive"):
        alts = mo.other_readings(text, tables, plan, hi)
        try:                                       # the last answer missed: this sentence goes to the learning list
            from app import learning

            learning.ask_to_learn(db, "office", text, guess="missed: " + mo.label_of(plan))
        except ValueError:
            pass
        if alts:
            agent.inputs = dumps(inputs)
            db.commit()
            mind.set_pattern(mo.label_of(plan), plan)
            mind.asked_back()
            return (("Pichli baar maine ise aise samjha tha, shayad galat. Inme se kaunsa chahiye?" if hi
                     else "Last time I read this one way and it may have been wrong. Which of these did you mean?"), alts)
    if asker.NO_CHANGE.search(text or "") or prof.get("control") == "ask_first":
        plan["readonly"] = True                            # "before making changes, tell me...": answer only
    show_fx = bool(sig.get("wants_formula") or (prof.get("formula") is True and not sig.get("no_formula")) or re.search(
        r"\b(in excel|excel mein|excel me)\b.{0,20}\b(kaise|how)\b|\bhow (do|can|to) i\b.{0,30}\b(calculate|count|sum|find)\b", text or "", re.I))
    other = None if plan.get("derive") else mo.sheet_or_value(text, tables, plan, hi)
    inputs["mind_plan"] = json.loads(json.dumps(plan, default=str))
    agent.inputs = dumps(inputs)
    mind.set_pattern(mo.label_of(plan), plan)
    args = {"plan": plan, "hi": hi}
    if f.size <= BIG:
        f.status, f.task, f.args = "queued", "ask", dumps(args)
        db.commit()
        _work(f.id, "ask", args)
        db.refresh(f)
        r = loads(f.result, {}) or {}
        if f.status == "ready" and r.get("task") == "ask":
            tail = ""
            if r.get("sheet"):
                tail = (f" Nayi sheet \"{r['sheet']}\" file mein jod di; card pe Download karo." if hi and r["sheet"].startswith("Answer") else
                        f" Ye \"{r['sheet']}\" sheet mein hai; card pe Download karo." if hi else
                        f" Added to the file as \"{r['sheet']}\"; download it from the card.")
            # the answer itself; the formula only when they asked for one (it's always in Details on the card)
            fx = (f" Formula: {r['formula']}" if r.get("formula") and show_fx else "")
            if show_fx != bool(r.get("show_formula")):
                f.result = dumps({**r, "show_formula": show_fx})
                db.commit()
            if plan.get("readonly") and not r.get("sheet"):
                tail = " Nothing in your file was changed." if not hi else " File mein kuch nahi badla."
            reply = f"{r['say']}.{fx}{tail}".replace("..", ".")
            if other:
                # answered the likelier reading; the other one is one tap away
                reply += (" Agar aapka matlab kuch aur tha, neeche chuno." if hi else " If you meant something else, pick it below.")
                return reply, [other]
            return reply
        return f.error or ("Ye kaam poora nahi ho saka." if hi else "Couldn't finish that.")
    _start(db, f, "ask", args)
    return "Badi file hai, kaam chal raha hai; card pe jawab aayega." if hi else "Big file, working on it; the answer will show on the card."


# ------------------------------------------------------------------ hooks the brain calls (named in the office pack)
def _outranked(text: str) -> bool:
    """Words of a field whose pack outranks this one (e.g. 'export my database to excel' while a sheet is open)."""
    return any(packs.mentions(name, text) for name, p in packs.all_packs().items() if "office" in p.get("outranks", []))


def holds_conversation(db: Session, goal: Goal, text: str, any_text: bool = False) -> Agent | None:
    """A file is open in this conversation: what the person types is about that file."""
    agent = goal_agent(db, goal)
    if agent is None or open_file(db, agent) is None:
        return None
    if not any_text and _outranked(text):
        return None
    return agent


def continue_turn(db: Session, goal: Goal, agent: Agent, text: str):
    from app.ml.normalize import HINGLISH_MARKERS, detect_language, tokens

    lang = detect_language(text)
    if lang == "English" and goal.detected_language in ("Hinglish", "Hindi") and (
            len(text.split()) < 4 or any(w in HINGLISH_MARKERS for w in tokens(text))):
        lang = goal.detected_language
    hi = lang in ("Hinglish", "Hindi")
    reply = chat_task(db, agent, text, hi)
    options = []
    if isinstance(reply, tuple):
        reply, options = reply
    if not reply:
        reply = cant_yet(open_file(db, agent), hi)
    return reply, options, hi


def wants_turn(db: Session, goal: Goal, text: str, relation: str, held: bool) -> bool:
    if relation in ("safety", "chit_chat"):
        return False
    return (held and not _outranked(text)) or wants_office(goal, text, db)


def take_turn(db: Session, goal: Goal, text: str, turn, visitor_id: str):
    """Excel / file work: the agent works on the file itself and gives back a fixed copy."""
    from app.officeknow import brain as officebrain

    rec = officebrain.recognize(text)
    agent = goal_agent(db, goal) or create_agent(db, goal)
    task_reply = chat_task(db, agent, text, turn.hi)
    inputs = loads(agent.inputs, {}) or {}
    if rec["confidence"] >= SURE and not task_reply and open_file(db, agent) is None:
        inputs["brain"] = rec
        goal.title = officebrain.describe(rec, turn.hi)
    elif not goal.title:
        goal.title = "Excel / file ka kaam" if turn.hi else "Excel and file work"
    agent.inputs = dumps(inputs)
    goal.missing_info = "[]"
    goal.pending = None
    goal.next_action_type = "execute"
    goal.next_action = "Give the file to the agent"
    db.commit()
    opened = open_file(db, agent)
    if not task_reply and opened is not None:
        task_reply = cant_yet(opened, turn.hi)
    options = []
    if isinstance(task_reply, tuple):
        task_reply, options = task_reply
    return (task_reply or plan_reply(rec, turn.hi)), options, "agent_plan"
