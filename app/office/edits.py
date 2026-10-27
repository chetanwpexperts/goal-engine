"""Everyday Excel orders the agent carries out on the file itself (a new version, so Undo works):
bold header, freeze the top row, fit column widths, filter buttons, remove blank rows, delete / rename a column,
serial numbers, split text into columns, a dropdown list, find & replace, date / money / decimal formats,
UPPER / Proper / lower case, one sheet per value of a column, charts, highlighting.

Understood by meaning (app/office/meaning.py), never by fixed sentences: the words say what kind of order it is,
the person's own file says which column, which value, which sheet.

plan_edit(text, tables) -> a plan {"job": "edit", "op": ..., "table": ..., ...} or None (not an order).
run_edit(plan, tables, wb, hi, n) -> {"say", "sheet", "changes"}.
English / Hinglish / Hindi. No AI."""

import re
from datetime import date, datetime

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.office import meaning as M

# Indian grouping (1,23,45,678.00): Excel has no built-in code for it, so three ranges
INDIAN = '[>=10000000]"₹"##\\,##\\,##\\,##0.00;[>=100000]"₹"##\\,##\\,##0.00;"₹"##,##0.00'
COLOURS = M.lexicon()["colours"]
WHOLE_FILE = ("bold_header", "freeze", "autofit", "filter", "remove_blank_rows", "replace")


class NotAnEdit(Exception):
    pass


def _ask():
    from app.office import ask
    return ask


def wants_macro(text: str) -> bool:
    return M.Reading(text).has("macro")


def wants_pivot(text: str) -> bool:
    return M.Reading(text).has("pivot")


EDGE = ("act", "polite", "to", "with", "into", "name", "rename", "column", "change", "replace", "all", "header", "sheet", "file")


def _edge_words() -> set[str]:
    """Order words that can sit around the person's own words ('rename X to Y karo'), never part of them."""
    lex = M.lexicon()["_words"]
    order = {w for w, cs in lex.items() if cs & set(EDGE)}
    return order | _ask().STOP | {"the", "of", "a", "an", "ka", "ki", "ke", "ko", "jagah", "wala", "wali", "wale", "likho"}


def _plain(text: str, drop=("column", "header", "sheet", "file", "act", "polite")) -> str:
    """The person's words without the order words, so column names stand out."""
    lex = M.lexicon()["_words"]
    toks = M.Reading(text).tokens
    return " ".join(t for t in toks if not (lex.get(t, set()) & set(drop)))


def _named(text, tables):
    return _ask().named_sheet(text, tables)


def _target(text, tables):
    a = _ask()
    named = a.named_sheet(text, tables)
    if named is not None:
        return named
    # no sheet named: the sheet that has the column they named ('make region uppercase')
    best, best_sc = None, 0.0
    plain = _plain(text)
    for t in tables:
        if not t.rows:
            continue
        sc = max((x for _, x in a.columns_in(plain, t)), default=0)
        if sc > best_sc:
            best, best_sc = t, sc
    scores = sorted((max((x for _, x in a.columns_in(plain, t)), default=0) for t in tables if t.rows), reverse=True)
    clear = len(scores) < 2 or scores[0] - scores[1] >= 0.2
    if best is not None and (best_sc >= 1 or (best_sc >= 0.6 and clear)):
        return best
    picked = a.pick_sheet(text, tables)
    if picked is not None:
        return picked
    data = [t for t in tables if t.rows]
    return max(data, key=lambda t: len(t.rows)) if data else (tables[0] if tables else None)


def _all_sheets(r: M.Reading) -> bool:
    return r.has("all") and r.has("sheet", "file")


def _col(text, t, kinds=None, exclude=(), need=0.5):
    a = _ask()
    for name in M.QUOTED.findall(text or ""):
        for i, h in enumerate(t.header):
            if h and h.strip().lower() == name.strip().lower() and i not in exclude:
                return i
    cands = [(sc, i) for i, sc in a.columns_in(_plain(text), t, kinds) if i not in exclude]
    return max(cands)[1] if cands and max(cands)[0] >= need else None


def _with_file(r: M.Reading, t, tables=None) -> M.Reading:
    """Add what the file knows: a column was named, a cell value was named, the value to find is in the data."""
    if t is None:
        return r
    a = _ask()
    text = r.text
    for x in tables or []:
        text = re.sub(rf"\b{re.escape(x.name)}\b", " ", text, flags=re.I)    # the sheet's own name isn't a value
    if a.columns_in(_plain(text), t) and max(sc for _, sc in a.columns_in(_plain(text), t)) >= 0.5:
        r.at.setdefault("COLUMN", []).append(-1)
    if any(x.rows and a.values_in(text, x) for x in (tables or [t])):
        r.at.setdefault("VALUE", []).append(-1)
    find, _ = _find_repl(r.text)
    if find and (M.QUOTED.search(r.text) or _in_data(find, t)):
        r.at.setdefault("FIND_VALUE", []).append(-1)
    doing = r.has("act", "remove", "change", "replace", "highlight", "freeze", "rename", "split", "round", "bold", "polite", "fit")
    if r.has("show", "question") and not doing:
        r.at.setdefault("show_only", []).append(-1)
    return r


def _in_data(value: str, t) -> bool:
    v = value.strip().lower()
    if not v:
        return False
    for i in range(len(t.header)):
        for x in t.col(i)[:5000]:
            if isinstance(x, str) and v in x.strip().lower():
                return True
    return False


def is_order(text: str) -> bool:
    """Could this be an order at all (before the file is read)? The file then decides (plan_edit)."""
    r = M.Reading(text)
    if r.asking():
        return False
    for c in ("COLUMN", "VALUE", "FIND_VALUE"):
        r.at.setdefault(c, []).append(-1)
    return bool(M.orders(r))


def _splittable(said, t, col):
    """The column whose values really have two parts ('Full Name', not 'Last Name')."""
    a = _ask()

    def parts(i):
        vals = [v for v in t.col(i)[:500] if isinstance(v, str) and v.strip()]
        return sum(1 for v in vals if len(v.split()) > 1) / len(vals) if vals else 0

    if col is not None and parts(col) >= 0.5:
        return col
    cands = [(sc + parts(i), i) for i, sc in a.columns_in(_plain(said), t, {"text"}) if parts(i) >= 0.5]
    if cands:
        return max(cands)[1]
    cands = [(parts(i), i) for i in range(len(t.header)) if t.kind(i) == "text" and parts(i) >= 0.5]
    if len(cands) == 1:
        return cands[0][1]
    _no_column(said, t, "split")


def _no_column(text, t, what):
    a = _ask()
    cols = ", ".join(h for h in t.header if h)[:300]
    raise a.NeedMore(f"{t.name} mein ye column nahi mila. Columns: {cols}. Kaunsa {what}?",
                     f"I couldn't find that column in {t.name}. It has: {cols}. Which one should I {what}?",
                     [h for h in t.header if h][:8])


def _number(text):
    ns = M.numbers(text)
    return ns[0] if ns else None


# ------------------------------------------------------------------ the person's own words
SEPARATORS = ("with", "to", "ko", "jagah", "into", "as", "by")


def _strip_edges(s: str) -> str:
    edge = _edge_words()
    words = s.split()
    while words and re.sub(r"\W", "", words[0].lower()) in edge:
        words.pop(0)
    while words and re.sub(r"\W", "", words[-1].lower()) in edge:
        words.pop()
    return " ".join(words).strip(" '\"“”‘’,.")


def _find_repl(text: str):
    """'replace X with Y', 'change X to Y', 'X ko Y se badlo', 'X ki jagah Y': (X, Y). Quotes win."""
    q = M.QUOTED.findall(text or "")
    if len(q) >= 2:
        return q[0], q[1]
    for sep in SEPARATORS:
        m = re.search(rf"\s{sep}\s", f" {text} ", re.I)
        if m:
            left, right = (f" {text} ")[:m.start()], (f" {text} ")[m.end():]
            find, repl = _strip_edges(left), _strip_edges(right)
            if find and repl:
                return find, repl
    return None, None


def _new_name(text: str, t, col: int) -> str | None:
    q = M.QUOTED.findall(text or "")
    if len(q) >= 2:
        return q[1]
    head = t.header[col] or ""
    m = re.search(re.escape(head), text, re.I) if head else None
    for part in ([text[m.end():], text[:m.start()]] if m else [text]):
        for sep in ("to", "as", "into"):
            mm = re.search(rf"\s{sep}\s", f" {part} ", re.I)
            if mm:
                part = (f" {part} ")[mm.end():]
                break
        name = _strip_edges(part)
        if name and name.lower() != head.lower():
            return name[:80]
    return None


def _list_values(text: str, r: M.Reading) -> list[str]:
    """'with Paid, Pending, Overdue' / 'options: A / B / C' -> the values."""
    words = M.lexicon()["concepts"]
    intro = "|".join(re.escape(w) for w in words["options"] + words["with"] if " " not in w)
    m = re.search(rf"(?:\b(?:{intro})\b|:)\s*(.+)$", text, re.I)
    if not m:
        return []
    stop = _ask().STOP | {"the", "a", "an"}

    def clean(v):
        w = v.split()
        while w and w[0].lower() in stop:
            w.pop(0)
        while w and w[-1].lower() in stop:
            w.pop()
        return " ".join(w).strip(" '\"“”‘’,.")
    vals = [clean(v) for v in re.split(r",|/|;|\|\s|\bor\b|\band\b|\baur\b|\bya\b", m.group(1), flags=re.I)]
    return [v for v in vals if v and len(v) <= 40]


# ------------------------------------------------------------------ understanding the order
def _mask_values(text: str, tables) -> str:
    """'Green Park', 'Red Fort': a value from the file is the person's data, not an order word (green, red)."""
    a = _ask()
    out = text
    for t in tables:
        if not t.rows:
            continue
        lex = M.lexicon()["_words"]
        for _, v in a.values_in(text, t):
            words = re.findall(r"\w+", v.lower())
            # 'Green Park' is data; 'Data Validation' is all order words, so it is what they're asking for
            if len(words) > 1 and not all(w in lex for w in words):
                out = re.sub(re.escape(v), " datavalue ", out, flags=re.I)
    return out


def _file_words(tables) -> set[str]:
    from app.ml.normalize import tokens

    return {w for t in tables for h in list(t.header) + [t.name] if h for w in tokens(str(h))}


def _fix_typos(text: str, tables) -> str:
    """'Regon' -> 'Region', 'Discoutn' -> 'Discount': the person's misspelt column names, put right from the file."""
    import difflib

    heads = {w.lower(): w for t in tables for h in t.header if h for w in re.findall(r"[A-Za-z]{4,}", h)}
    lex = M.lexicon()["_words"]
    out = []
    for w in re.split(r"(\s+)", text or ""):
        low = w.lower()
        if len(low) >= 4 and low.isalpha() and low not in heads and low not in lex:
            close = difflib.get_close_matches(low, list(heads), n=1, cutoff=0.8)
            if close:
                w = heads[close[0]]
        out.append(w)
    return "".join(out)


def plan_edit(text: str, tables) -> dict | None:
    text = _fix_typos(text, tables)
    r = M.Reading(_mask_values(text, tables), _file_words(tables))
    r.text = text
    if r.asking():
        return None                                          # a question about the file, not an order
    t = _target(text, tables)
    if t is None:
        return None
    _with_file(r, t, tables)
    for _, op in M.orders(r):
        tt = t
        if op == "conditional_format" and r.has("VALUE") and not r.has("above", "below", "between", "equal", "NUMBER", "scale", "top",
                                                                      "bottom", "duplicate", "blank", "overdue") \
                and not _ask().values_in(text, t):
            with_value = [c for c in tables if c.rows and _ask().values_in(text, c)]
            if with_value:
                tt = max(with_value, key=lambda c: len(_ask().values_in(text, c)))
        plan = {"job": "edit", "op": op, "table": tt.name,
                "all": _all_sheets(r) or (op in WHOLE_FILE and _named(text, tables) is None and len([x for x in tables if x.rows]) > 1)}
        try:
            plan.update(_details(op, text, r, tt, tables))
        except NotAnEdit:
            continue
        return plan
    return None


def _details(op, text, r: M.Reading, t, tables) -> dict:
    a = _ask()
    if op == "chart":
        return _chart_details(text, r, t)
    if op == "conditional_format":
        return _conditional_details(text, r, t)
    if op == "freeze":
        row = r.has("FIRST_ROW", "header") or (r.has("row") and not r.has("FIRST_COL"))
        col = r.has("FIRST_COL") or (r.has("column") and not r.has("row", "header"))
        return {"cell": "B2" if row and col else ("B1" if col else "A2")}
    if op in ("delete_column", "case", "round", "money_format", "date_format", "dropdown", "split_sheets", "split"):
        kinds = {"date_format": {"date"}, "money_format": {"number"}, "round": {"number"}, "split_sheets": {"text"}}.get(op)
        col = _col(text, t, kinds, need=1 if op == "delete_column" else 0.5)
        if op == "split":
            col = _splittable(text, t, col)
        if op == "date_format" and col is None:
            col = next((i for i in range(len(t.header)) if t.kind(i) == "date"), None)
        if op in ("money_format", "round") and col is None:
            col = a.main_amount(t)[0]
        if col is None and op == "delete_column":
            _no_column(text, t, "delete")
        if col is None and op == "split_sheets":
            _no_column(text, t, "split by")
        if col is None and op not in ("date_format",):
            raise NotAnEdit
        out = {"col": col}
        if op == "money_format":
            said = set(r.tokens)
            only_commas = bool(said & {"comma", "commas"}) and not (said & {"rupee", "rupees", "rs", "inr", "currency", "rupaye", "rupiya", "rupay"})
            out["rupee"] = not only_commas
        if op == "case":
            first = min((r.first(c), c) for c in ("upper", "lower", "proper") if r.has(c))[1]
            if r.has("first") and r.has("letter") and r.has("upper"):
                first = "proper"
            out["how"] = first
        if op == "round":
            small = [n for n in r.numbers if isinstance(n, int) and 0 <= n <= 6]
            out["places"] = 0 if r.has("zero") else (small[0] if small else 2)
        if op == "date_format":
            out["fmt"] = r.mask.group(1).lower() if r.mask else "dd-mm-yyyy"
        if op == "dropdown":
            vals = [v for v in _list_values(text, r) if v.lower() != (t.header[col] or "").lower()]
            if len(vals) < 2:
                vals = sorted({str(v).strip() for v in t.col(col)[:5000] if v not in (None, "")})[:50]
            out["values"] = vals[:50]
        if op == "split":
            q = [x for x in M.QUOTED.findall(text) if len(x) <= 3]
            sep = q[0] if q else ("," if r.has("comma_sep") else "-" if r.has("hyphen_sep") else "/" if r.has("slash_sep") else " ")
            out["sep"] = sep
            out["names"] = bool(re.search(r"name|naam", t.header[col] or "", re.I)) and sep == " "
        return out
    if op == "rename":
        q = M.QUOTED.findall(text)
        col = _col(f'"{q[0]}"', t) if q else _col(text, t)
        if col is None:
            _no_column(text, t, "rename")
        new = _new_name(text, t, col)
        if not new:
            raise NotAnEdit
        return {"col": col, "new": new}
    if op == "replace":
        find, repl = _find_repl(text)
        if not find:
            raise NotAnEdit
        col = None
        m = re.search(r"\s(?:in|mein|me)\s+(?:the\s+)?(.+?)\s*$", f" {repl} ", re.I)
        if m:
            c = _col(m.group(1), t, need=1)
            if c is not None:
                col, repl = c, _strip_edges((f" {repl} ")[:m.start()])
        return {"find": find, "repl": repl, "col": col}
    return {}


def _group_named(r: M.Reading, t):
    """The word next to by / per / wise / ke hisaab: that is the grouping. A word there that isn't a column -> ask."""
    a = _ask()
    lex = M.lexicon()["_words"]
    stop = a.STOP | {"the", "a"}
    for p in r.at.get("group", []):
        if p < 0:
            continue
        after = r.tokens[p] in ("by", "per", "each", "every", "har", "according")
        near = [r.tokens[i] for i in ((p + 1, p + 2) if after else (p - 1, p - 2)) if 0 <= i < len(r.tokens)]
        for w in near:
            if w in stop or (w in lex and not (lex[w] & {"column"})) or w.isdigit() or M._fuzzy(w):
                continue
            hits = [(sc, i) for i, sc in a.columns_in(w, t) if sc >= 0.5 and t.kind(i) in ("text", "date")]
            if hits:
                return max(hits)[1]
            if w not in lex:
                _no_column(r.text, t, "use for the chart groups")
    return None


def _chart_details(text, r: M.Reading, t):
    a = _ask()
    lex = M.lexicon()["_words"]
    rowish = any("row" in lex.get(w, ()) and "line_kind" not in lex.get(w, ()) for w in r.tokens)
    kind = "pie" if r.has("pie") else "line" if r.has("line_kind") and not rowish else "bar"
    direction = "bar" if "bar" in r.tokens or "bars" in r.tokens else "col"
    month = r.has("month")
    if month:
        dates = a.columns_in(_plain(text), t, {"date"})
        group = dates[0][0] if dates else next((i for i in range(len(t.header)) if t.kind(i) == "date"), None)
    else:
        group = _group_named(r, t)
        if group is None:
            named = [(sc, i) for i, sc in a.columns_in(_plain(text), t, {"text"}) if sc >= 0.5]
            group = max(named)[1] if named else a._group_col(text, t, [])
    if group is None:
        _no_column(text, t, "use for the chart groups")
    named = [(sc, i) for i, sc in a.columns_in(_plain(text), t, {"number"}) if i != group and sc >= 0.5]
    measure = max(named)[1] if named else a.main_amount(t)[0]
    if measure is None:
        _no_column(text, t, "chart")
    group_name = "Month" if month else t.header[group]
    return {"chart_type": kind, "direction": direction, "group_col": group, "value_col": measure,
            "month": month, "title": f"{t.header[measure]} by {group_name}", "next_to_data": r.has("beside")}


def _conditional_details(text, r: M.Reading, t):
    a = _ask()
    out = {"colour": M.colour(r, "yellow")}
    num_col = lambda: _col(text, t, {"number"}) if _col(text, t, {"number"}) is not None else a.main_amount(t)[0]  # noqa: E731
    if r.has("scale"):
        rule, col = "color_scale", num_col()
    elif r.has("top", "bottom") and r.has("highlight", "red", "green", "yellow", "orange", "blue", "pink", "grey") and \
            not r.has("above", "below", "between"):
        rule, col = ("bottom" if r.has("bottom") and not r.has("top") else "top10"), num_col()
        small = [n for n in r.numbers if isinstance(n, int) and 1 <= n <= 1000]
        out["rank"] = small[0] if small else 10
    elif r.has("duplicate"):
        rule, col = "duplicate", _col(text, t)
        if col is None:
            _no_column(text, t, "check for repeats")
    elif r.has("blank"):
        rule, col = "blank", _col(text, t)
        if col is None:
            col = "all"
    elif r.has("overdue"):
        rule = "overdue"
        col = _col(text, t, {"date"})
        if col is None:
            dates = [i for i in range(len(t.header)) if t.kind(i) == "date"]
            due = [i for i in dates if re.search(r"due|delivery|expiry|last", t.header[i] or "", re.I)]
            col = due[0] if due else (dates[0] if len(dates) == 1 else None)
        if col is None:
            _no_column(text, t, "check for overdue dates")
        out["colour"] = M.colour(r, "red")
    elif r.has("between") and len(r.numbers) >= 2:
        rule, col, out["values"] = "between", num_col(), sorted(r.numbers[:2])
    elif r.has("below") and r.numbers:
        rule, col, out["value"] = "below", num_col(), r.numbers[0]
    elif r.has("above") and r.numbers:
        rule, col, out["value"] = "above", num_col(), r.numbers[0]
    elif r.has("equal") and r.numbers:
        rule, col, out["value"] = "equal", num_col(), r.numbers[0]
    elif r.has("VALUE") and a.values_in(text, t):
        col, value = a.values_in(text, t)[0]
        rule, out["value"] = "text_equal", value
    elif r.numbers:
        rule, col, out["value"] = "equal", num_col(), r.numbers[0]
    else:
        raise NotAnEdit
    if col is None:
        _no_column(text, t, "highlight")
    # "... wale", "rows where ...", "orders above ...": the whole row lights up, not just the cell
    out["rows"] = r.has("row") and col != "all" and rule != "color_scale"
    return {"col": col, "rule": rule, **out}


# ------------------------------------------------------------------ doing it
def _sheets(p, wb, tables=None):
    every = p.get("all") and not (p["op"] == "replace" and p.get("col") is not None)
    names = [n for n in wb.sheetnames if tables is None or (tables.get(n) is not None and tables[n].rows)] if every else [p["table"]]
    return [wb[n] for n in names if n in wb.sheetnames and not n.startswith("Answer ") and n not in ("Changes (Goal Engine)",)]


def _hdr(ws, t):
    return t.header_row if t is not None else 1


def _has_formulas(ws) -> bool:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                return True
    return False


def _safe_sheet_name(wb, title):
    base = re.sub(r"[\[\]:*?/\\]", " ", f"Chart - {title}").strip()[:31] or "Chart"
    name, n = base, 2
    while name in wb.sheetnames:
        suffix = f" ({n})"
        name, n = base[:31 - len(suffix)] + suffix, n + 1
    return name


def _chart_sheet(p, t, wb, source):
    """Write the grouped values first, then chart exactly those cells."""
    target = source if p.get("next_to_data") else wb.create_sheet(_safe_sheet_name(wb, p["title"]))
    start_col = source.max_column + 2 if target is source else 1
    group, value = p["group_col"], p["value_col"]
    totals = {}
    for row in t.rows:
        if group >= len(row) or value >= len(row):
            continue
        g, v = row[group], row[value]
        if g in (None, "") or not isinstance(v, (int, float)) or isinstance(v, bool):
            continue
        if p.get("month"):
            if not isinstance(g, (date, datetime)):
                continue
            g = (g.year, g.month)
        else:
            g = str(g).strip()
        totals[g] = totals.get(g, 0) + v
    items = sorted(totals.items(), key=lambda x: x[0])
    target.cell(1, start_col, "Month" if p.get("month") else t.header[group]).font = Font(bold=True)
    target.cell(1, start_col + 1, f"Total {t.header[value]}").font = Font(bold=True)
    for r, (label, total) in enumerate(items, start=2):
        if isinstance(label, tuple):
            label = date(label[0], label[1], 1).strftime("%b %Y")          # in date order, shown as "Sep 2026"
        target.cell(r, start_col, label)
        target.cell(r, start_col + 1, total)
    chart = PieChart() if p["chart_type"] == "pie" else LineChart() if p["chart_type"] == "line" else BarChart()
    if isinstance(chart, BarChart):
        chart.type = p.get("direction", "col")
    chart.title = p["title"]
    if not isinstance(chart, PieChart):
        chart.x_axis.title = "Month" if p.get("month") else t.header[group]
        chart.y_axis.title = f"Total {t.header[value]}"
    else:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showCatName = True
        chart.dataLabels.showVal = False
        chart.dataLabels.showSerName = False
        chart.dataLabels.showLegendKey = False
        chart.dataLabels.showLeaderLines = True
    chart.width, chart.height = 12, 7.5
    last = len(items) + 1
    chart.add_data(Reference(target, min_col=start_col + 1, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(target, min_col=start_col, min_row=2, max_row=last))
    if isinstance(chart, LineChart):
        for series in chart.series:
            series.smooth = False                      # straight lines between months: a curve would invent values
    target.add_chart(chart, f"{get_column_letter(start_col + 3)}2")
    target.column_dimensions[get_column_letter(start_col)].width = 18
    target.column_dimensions[get_column_letter(start_col + 1)].width = 20
    return target.title


def _add_conditional_format(ws, p, hr, last, t=None):
    first = hr + 1
    if last < first:
        return 0
    end = get_column_letter(max(ws.max_column, 1))
    rows = p.get("rows")
    if p["col"] == "all":
        letter, cell_range = "A", f"A{first}:{end}{last}"
    else:
        letter = get_column_letter(p["col"] + 1)
        cell_range = f"A{first}:{end}{last}" if rows else f"{letter}{first}:{letter}{last}"
    ref = f"${letter}{first}" if rows else f"{letter}{first}"          # whole rows: the column stays fixed
    span = f"${letter}${first}:${letter}${last}"
    rule = p["rule"]
    if rule == "color_scale":
        cf = ColorScaleRule(start_type="min", start_color="FFF8696B", mid_type="percentile", mid_value=50,
                            mid_color="FFFFEB84", end_type="max", end_color="FF63BE7B")
    else:
        fill = PatternFill(fill_type="solid", fgColor=COLOURS[p.get("colour", "yellow")])
        if rule in ("above", "below", "equal", "between") and not rows:
            op = {"above": "greaterThan", "below": "lessThan", "equal": "equal", "between": "between"}[rule]
            formulas = [str(v) for v in (p.get("values") or [p.get("value")])]
            cf = CellIsRule(operator=op, formula=formulas, fill=fill)
        elif rule in ("top10", "bottom") and not rows:
            cf = Rule(type="top10", rank=int(p.get("rank", 10)), bottom=(rule == "bottom"), dxf=DifferentialStyle(fill=fill))
        else:
            v = p.get("value")
            f = {"above": f"AND(ISNUMBER({ref}),{ref}>{v})", "below": f"AND(ISNUMBER({ref}),{ref}<{v})",
                 "equal": f"{ref}={v}",
                 "between": f"AND(ISNUMBER({ref}),{ref}>={(p.get('values') or [0, 0])[0]},{ref}<={(p.get('values') or [0, 0])[1]})",
                 "top10": f"AND(ISNUMBER({ref}),{ref}>=LARGE({span},{int(p.get('rank', 10))}))",
                 "bottom": f"AND(ISNUMBER({ref}),{ref}<=SMALL({span},{int(p.get('rank', 10))}))",
                 "overdue": f"AND(ISNUMBER({ref}),{ref}<TODAY())",
                 "duplicate": f'AND({ref}<>"",COUNTIF({span},{ref})>1)',
                 "blank": f'LEN(TRIM({ref}&""))=0',
                 "text_equal": f'{ref}="{str(v).replace(chr(34), chr(34) * 2)}"'}[rule]
            cf = FormulaRule(formula=[f], fill=fill)
    ws.conditional_formatting.add(cell_range, cf)
    p["matched"] = p.get("matched", 0) + _matches(t, p)
    return 1


def _matches(t, p) -> int:
    """How many cells the highlight lights up now (said in the reply), from the values Excel shows."""
    if t is None:
        return 0
    rule = p["rule"]
    if p["col"] == "all":
        cells = [v for row in t.rows for v in row]
    else:
        cells = list(t.col(p["col"]))
    nums = [v for v in cells if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if rule == "above":
        return sum(1 for v in nums if v > p["value"])
    if rule == "below":
        return sum(1 for v in nums if v < p["value"])
    if rule == "equal":
        return sum(1 for v in nums if v == p["value"])
    if rule == "between":
        lo, hi = p["values"]
        return sum(1 for v in nums if lo <= v <= hi)
    if rule in ("top10", "bottom"):
        return min(int(p.get("rank", 10)), len(nums))
    if rule == "blank":
        return sum(1 for v in cells if v is None or (isinstance(v, str) and not v.strip()))
    if rule == "text_equal":
        return sum(1 for v in cells if isinstance(v, str) and v.strip() == str(p["value"]))
    if rule == "duplicate":
        from collections import Counter
        cnt = Counter(v for v in cells if v not in (None, ""))
        return sum(n for n in cnt.values() if n > 1)
    if rule == "overdue":
        today = date.today()
        return sum(1 for v in cells if isinstance(v, (date, datetime)) and (v.date() if isinstance(v, datetime) else v) < today)
    return 0


def _indian(n) -> str:
    if not isinstance(n, (int, float)) or isinstance(n, bool):
        return str(n)
    n = float(n)
    whole, frac = int(abs(n)), abs(n) - int(abs(n))
    s = str(whole)
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        head = ",".join([head[max(0, i - 2):i] for i in range(len(head), 0, -2)][::-1])
        s = f"{head},{tail}"
    if frac:
        s += f"{frac:.2f}"[1:]
    return ("-" if n < 0 else "") + s


def _cf_say(p, col, where, hi):
    colour = p.get("colour", "yellow")
    hcol = {"red": "laal", "green": "hara", "yellow": "peela", "orange": "narangi", "blue": "neela", "pink": "gulabi", "grey": "grey"}[colour]
    rule, n = p["rule"], p.get("matched", 0)
    what = col or ("the table" if not hi else "table")
    if rule == "color_scale":
        return (f"{where} mein {col} pe colour scale laga diya: chhote laal, beech wale peele, bade hare." if hi else
                f"Added a colour scale to {col} in {where}: low values red, middle yellow, high green.")
    en = {"above": f"{col} above {_indian(p.get('value', 0))}", "below": f"{col} below {_indian(p.get('value', 0))}",
          "equal": f"{col} equal to {_indian(p.get('value', 0))}",
          "between": f"{col} between {_indian((p.get('values') or [0, 0])[0])} and {_indian((p.get('values') or [0, 0])[1])}",
          "top10": f"the top {p.get('rank', 10)} {col}", "bottom": f"the bottom {p.get('rank', 10)} {col}",
          "duplicate": f"repeated {col} values", "blank": f"blank cells in {what}", "overdue": f"overdue dates in {col}",
          "text_equal": f"rows where {col} is {p.get('value')}"}[rule]
    hn = {"above": f"{col} {_indian(p.get('value', 0))} se zyada", "below": f"{col} {_indian(p.get('value', 0))} se kam",
          "equal": f"{col} {_indian(p.get('value', 0))} ke barabar",
          "between": f"{col} {_indian((p.get('values') or [0, 0])[0])} aur {_indian((p.get('values') or [0, 0])[1])} ke beech",
          "top10": f"sabse bade {p.get('rank', 10)} {col}", "bottom": f"sabse chhote {p.get('rank', 10)} {col}",
          "duplicate": f"{col} ki dohrai values", "blank": f"{what} ke khaali cells", "overdue": f"{col} ki nikal chuki dates",
          "text_equal": f"jahan {col} {p.get('value')} hai"}[rule]
    unit = ("row" if not hi else "row") if p.get("rows") else "cell"
    cnt = f" ({n} {unit}{'s' if n != 1 else ''})" if rule not in ("top10", "bottom") else ""
    return (f"{where} mein {hn} {hcol} kar diye{cnt}." if hi else f"Highlighted {en} in {colour} in {where}{cnt}.")


def run_edit(p, tables, wb, hi, n) -> dict:
    if wb is None:
        raise ValueError("needs the workbook")
    op = p["op"]
    t = tables.get(p["table"])
    done, notes, output_sheet = 0, [], p["table"]
    for ws in _sheets(p, wb, tables):
        tt = tables.get(ws.title) or t
        hr = _hdr(ws, tt)
        last = ws.max_row
        col = p.get("col")
        cidx = (col + 1) if isinstance(col, int) else None
        if op == "chart":
            output_sheet = _chart_sheet(p, tt, wb, ws)
            done += 1
        elif op == "conditional_format":
            done += _add_conditional_format(ws, p, hr, last, tt)
        elif op == "bold_header":
            for c in ws[hr]:
                if c.value not in (None, ""):
                    c.font = Font(bold=True)
                    done += 1
        elif op == "freeze":
            cell = p.get("cell", "A2")
            if cell.endswith("2"):
                cell = f"{cell[0]}{hr + 1}"
            ws.freeze_panes = cell
            done += 1
        elif op == "autofit":
            for ci, column in enumerate(ws.iter_cols(min_row=1, max_row=min(last, 2000)), start=1):
                width = max((len(str(c.value)) for c in column if c.value not in (None, "")), default=6)
                ws.column_dimensions[get_column_letter(ci)].width = min(60, max(8, width + 2))
                done += 1
        elif op == "filter":
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(ws.max_column)}{last}"
            done += 1
        elif op == "remove_blank_rows":
            if _has_formulas(ws):
                notes.append(ws.title)
                continue
            blanks = [r for r in range(last, hr, -1) if all(c.value in (None, "") or (isinstance(c.value, str) and not c.value.strip()) for c in ws[r])]
            for r in blanks:
                ws.delete_rows(r)
            done += len(blanks)
        elif op == "delete_column":
            if _has_formulas(ws):
                notes.append(ws.title)
                continue
            ws.delete_cols(cidx)
            done += 1
        elif op == "rename":
            ws.cell(hr, cidx).value = p["new"]
            done += 1
        elif op == "serial":
            if _has_formulas(ws):
                ci = ws.max_column + 1                      # inserting would shift formulas: add it at the end instead
                notes.append(ws.title)
            else:
                ws.insert_cols(1)
                ci = 1
            ws.cell(hr, ci).value = "S.No"
            ws.cell(hr, ci).font = Font(bold=True)
            k = 0
            for r in range(hr + 1, ws.max_row + 1):
                if any(c.value not in (None, "") for c in ws[r] if c.column != ci):
                    k += 1
                    ws.cell(r, ci).value = k
            done += k
        elif op == "split":
            sep = p.get("sep", " ")
            parts_n = 2 if p.get("names") else max((len(str(ws.cell(r, cidx).value).split(sep)) for r in range(hr + 1, last + 1)
                                                    if ws.cell(r, cidx).value not in (None, "")), default=2)
            parts_n = min(parts_n, 8)
            base = ws.max_column + 1
            heads = ["First Name", "Last Name"] if p.get("names") else [f"{ws.cell(hr, cidx).value} {k + 1}" for k in range(parts_n)]
            for k, h in enumerate(heads):
                ws.cell(hr, base + k).value = h
                ws.cell(hr, base + k).font = Font(bold=True)
            for r in range(hr + 1, last + 1):
                v = ws.cell(r, cidx).value
                if v in (None, ""):
                    continue
                s = str(v).strip()
                bits = s.split(sep, 1) if p.get("names") else s.split(sep)
                bits = [b.strip() for b in bits]
                for k in range(len(heads)):
                    ws.cell(r, base + k).value = bits[k] if k < len(bits) else None
                done += 1
        elif op == "dropdown":
            vals = [str(v).replace(",", " ") for v in p.get("values", [])]
            formula = '"' + ",".join(vals) + '"'
            if len(formula) > 255:                          # Excel's limit for a typed list: put the list on its own sheet
                lst = wb.create_sheet(f"List {ws.cell(hr, cidx).value}"[:31])
                for k, v in enumerate(vals, start=1):
                    lst.cell(k, 1).value = v
                formula = f"='{lst.title}'!$A$1:$A${len(vals)}"
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
            ws.add_data_validation(dv)
            letter = get_column_letter(cidx)
            dv.add(f"{letter}{hr + 1}:{letter}{max(last, hr) + 500}")
            done += 1
        elif op == "replace":
            f, rp = p["find"], p["repl"]
            for row in ws.iter_rows(min_row=hr + 1):
                for c in row:
                    if cidx is not None and c.column != cidx:
                        continue
                    v = c.value
                    if isinstance(v, str) and not v.startswith("=") and f.lower() in v.lower() and rp.lower() not in v.lower():
                        c.value = re.sub(re.escape(f), rp, v, flags=re.I)
                        done += 1
                    elif isinstance(v, (int, float)) and str(v) == f.strip():
                        c.value = rp
                        done += 1
        elif op == "date_format":
            fmt = p.get("fmt", "dd-mm-yyyy")
            cols = [cidx] if cidx else [c for c in range(1, ws.max_column + 1)]
            for r in range(hr + 1, last + 1):
                for ci in cols:
                    c = ws.cell(r, ci)
                    if isinstance(c.value, (datetime, date)):
                        c.number_format = fmt
                        done += 1
        elif op == "money_format":
            for r in range(hr + 1, last + 1):
                c = ws.cell(r, cidx)
                if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                    c.number_format = INDIAN.replace("₹", "₹" if p.get("rupee", True) else "")
                    done += 1
        elif op == "round":
            places = p.get("places", 2)
            for r in range(hr + 1, last + 1):
                c = ws.cell(r, cidx)
                if isinstance(c.value, str) and c.value.startswith("=") and not c.value.upper().startswith("=ROUND("):
                    c.value = f"=ROUND({c.value[1:]},{places})"
                    c.number_format = "0" if places == 0 else "0." + "0" * places
                    done += 1
                elif isinstance(c.value, float):
                    c.value = round(c.value, places)
                    c.number_format = "0" if places == 0 else "0." + "0" * places
                    done += 1
        elif op == "case":
            how = p["how"]
            for r in range(hr + 1, last + 1):
                c = ws.cell(r, cidx)
                if isinstance(c.value, str) and not c.value.startswith("="):
                    new = c.value.upper() if how == "upper" else c.value.lower() if how == "lower" else c.value.title()
                    if new != c.value:
                        c.value = new
                        done += 1
        elif op == "split_sheets":
            groups = {}
            for r in range(hr + 1, last + 1):
                v = ws.cell(r, cidx).value
                if v in (None, ""):
                    continue
                groups.setdefault(str(v).strip(), []).append(r)
            if len(groups) > 60:
                raise ValueError(f"{len(groups)} different values: too many sheets. Pick a column with fewer groups.")
            header = [c.value for c in ws[hr]]
            for k, (val, rows) in enumerate(sorted(groups.items())):
                name = re.sub(r"[\[\]:*?/\\]", " ", val)[:31] or f"Group {k + 1}"
                if name in wb.sheetnames:
                    name = f"{name[:27]} ({k + 1})"
                new = wb.create_sheet(name)
                new.append(header)
                for c in new[1]:
                    c.font = Font(bold=True)
                for r in rows:
                    new.append([c.value for c in ws[r]])
                done += 1
    return {"say": _say(p, done, notes, tables, hi), "sheet": output_sheet, "changes": done}


def _say(p, done, notes, tables, hi):
    op, t = p["op"], tables.get(p["table"])
    where = ("every sheet" if not hi else "har sheet") if p.get("all") and not (op == "replace" and p.get("col") is not None) else p["table"]
    col = t.header[p["col"]] if (t is not None and isinstance(p.get("col"), int) and p["col"] < len(t.header)) else ""
    en = {
        "chart": f"Made a {p.get('chart_type', 'bar')} chart of {p.get('title', '')} in {where}",
        "conditional_format": "",
        "bold_header": f"Made the heading row bold in {where}",
        "freeze": f"Froze the {'heading row and first column' if p.get('cell', '').startswith('B2') else 'first column' if p.get('cell') == 'B1' else 'heading row'} in {where}, so it stays visible while scrolling",
        "autofit": f"Fitted the column widths in {where} to their contents",
        "filter": f"Added filter buttons to the headings in {where}",
        "remove_blank_rows": f"Removed {done} blank rows from {where}",
        "delete_column": f"Deleted the column {col} from {where}",
        "rename": f"Renamed the column {col} to {p.get('new', '')} in {where}",
        "serial": f"Added a S.No column numbering {done} rows in {where}",
        "split": f"Split {col} into {'First Name and Last Name' if p.get('names') else 'separate columns'} for {done} rows in {where} (new columns at the end; {col} kept as it was)",
        "dropdown": f"Added a dropdown to {col} in {where} with: {', '.join(p.get('values', [])[:8])}{' …' if len(p.get('values', [])) > 8 else ''}",
        "replace": f"Replaced \"{p.get('find', '')}\" with \"{p.get('repl', '')}\" in {done} cells{' of ' + col if col else ''} in {where}",
        "date_format": f"Changed {done} dates{' in ' + col if col else ''} to {p.get('fmt', 'dd-mm-yyyy')} in {where}",
        "money_format": f"Showed {done} values of {col} in rupee format (₹1,23,456.00) in {where}",
        "round": f"Rounded {done} values of {col} to {p.get('places', 2)} decimal places in {where}",
        "case": f"Changed {done} values of {col} to {p.get('how', 'proper').replace('proper', 'Proper Case').replace('upper', 'UPPER CASE').replace('lower', 'lower case')} in {where}",
        "split_sheets": f"Made {done} new sheets from {where}, one for each {col}",
    }[op]
    if op == "conditional_format":
        en = _cf_say(p, col, where, False).rstrip(".")
    if hi:
        hn = {
            "chart": f"{where} ke liye {p.get('title', '')} ka {p.get('chart_type', 'bar')} chart bana diya",
            "conditional_format": "",
            "bold_header": f"{where} ki heading row bold kar di",
            "freeze": f"{where} mein {'heading row aur pehla column' if p.get('cell', '').startswith('B2') else 'pehla column' if p.get('cell') == 'B1' else 'heading row'} freeze kar diya; scroll karne pe bhi dikhega",
            "autofit": f"{where} ke columns ki width unke data jitni kar di",
            "filter": f"{where} ki headings pe filter buttons laga diye",
            "remove_blank_rows": f"{where} se {done} khaali rows hata di",
            "delete_column": f"{where} se {col} column hata diya",
            "rename": f"{where} mein {col} column ka naam {p.get('new', '')} kar diya",
            "serial": f"{where} mein S.No column jod diya ({done} rows)",
            "split": f"{where} mein {col} ko {'First Name aur Last Name' if p.get('names') else 'alag columns'} mein tod diya ({done} rows; naye columns aakhir mein, {col} waisa hi hai)",
            "dropdown": f"{where} ke {col} mein dropdown laga diya: {', '.join(p.get('values', [])[:8])}{' …' if len(p.get('values', [])) > 8 else ''}",
            "replace": f"{where} mein {done} jagah \"{p.get('find', '')}\" ko \"{p.get('repl', '')}\" kar diya",
            "date_format": f"{where} mein {done} dates {p.get('fmt', 'dd-mm-yyyy')} format mein kar di",
            "money_format": f"{where} mein {col} ke {done} values rupee format (₹1,23,456.00) mein kar diye",
            "round": f"{where} mein {col} ke {done} values {p.get('places', 2)} decimal tak round kar diye",
            "case": f"{where} mein {col} ke {done} values {p.get('how', 'proper').replace('proper', 'Proper Case').replace('upper', 'UPPER CASE').replace('lower', 'lower case')} mein kar diye",
            "split_sheets": f"{where} se {done} nayi sheets bana di, har {col} ki ek",
        }[op]
        if op == "conditional_format":
            hn = _cf_say(p, col, where, True).rstrip(".")
        if notes:
            hn += f". {', '.join(notes)} mein formulas hain, isliye wahan rows/columns nahi khiskaye" if op in ("remove_blank_rows", "delete_column") \
                else f". {', '.join(notes)} mein S.No aakhir mein lagaya taaki formulas sahi rahein"
        if done == 0 and op in ("remove_blank_rows", "replace", "case", "round", "date_format", "money_format"):
            hn = f"{where} mein badalne layak kuch nahi mila; kuch nahi badla"
        return hn
    if notes:
        en += (f". On {', '.join(notes)} rows and columns weren't moved because those sheets have formulas that would break"
               if op in ("remove_blank_rows", "delete_column") else f". On {', '.join(notes)} the S.No column went at the end, so the formulas stay correct")
    if done == 0 and op in ("remove_blank_rows", "replace", "case", "round", "date_format", "money_format"):
        any_ = "any sheet" if where == "every sheet" else where
        en = {"remove_blank_rows": f"No blank rows in {any_}; nothing changed",
              "replace": f"\"{p.get('find', '')}\" isn't in {any_}; nothing changed"}.get(op, f"Nothing to change in {where}")
    return en
