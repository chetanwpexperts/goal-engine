"""Builds the Word document the tests use (a small company report with headings, text, two tables and messy spacing)."""

from docx import Document


def build(path):
    d = Document()
    d.add_heading("Quarterly Sales Report", 0)
    d.add_paragraph("Prepared by the finance team on 05-10-2026 for the board.  Contact: finance@acme.in or 98765 43210.")
    d.add_heading("Summary", 1)
    d.add_paragraph("Sales grew strongly this quarter, led by the South region. Total revenue reached ₹34,14,200 across all regions. "
                    "The North region remained the weakest and needs a new distributor before the deadline of 31 Dec 2026.")
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_heading("Regional performance", 1)
    d.add_paragraph("Region figures:")
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Region", "Sales", "Manager"
    for r, s, m in [("South", "12,33,200", "Priya"), ("East", "8,93,900", "Rahul"), ("West", "8,52,600", "Anita"), ("North", "4,34,500", "Vikram")]:
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = r, s, m
    d.add_heading("Hiring plan", 1)
    d.add_paragraph("We plan to hire  twelve new sales staff in Delhi and Mumbai. Training will start in January with a budget of Rs 5 lakh.")
    d.add_paragraph("Delhi office will move to a larger floor. The Delhi team supports the North region.   ")
    d.add_paragraph("Staff list:")
    t2 = d.add_table(rows=1, cols=3)
    t2.rows[0].cells[0].text, t2.rows[0].cells[1].text, t2.rows[0].cells[2].text = "Name", "City", "Salary"
    for n, c_, s in [("Aman Gupta", "Delhi", "50000"), ("Neha Singh", "Mumbai", "62000"), ("Ravi Kumar", "Delhi", "48000")]:
        c = t2.add_row().cells
        c[0].text, c[1].text, c[2].text = n, c_, s
    d.add_heading("Risks", 1)
    d.add_paragraph("The main risk is late payments from distributors. More details are at https://acme.in/risks.")
    d.save(path)


if __name__ == "__main__":
    build("tests/data/report.docx")
