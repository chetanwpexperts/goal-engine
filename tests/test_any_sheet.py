"""The Excel agent on sheets it has never seen: school marks, clinic bills (amounts as "₹ 1,200" text), GST invoices +
customers, a godown stock list, two monthly sheets. Data is random; the right answer is worked out separately here."""

import datetime
import random
import re
from collections import defaultdict

import pytest
from openpyxl import Workbook, load_workbook

from app.office import ask
from app.office.sheetdata import load_tables


@pytest.fixture(scope="module")
def books(tmp_path_factory):
    DIR = tmp_path_factory.mktemp("sheets")
    random.seed(7)
    def save(name, sheets):
        wb=Workbook(); wb.remove(wb.active)
        for title, header, rows in sheets:
            ws=wb.create_sheet(title); ws.append(header)
            for r in rows: ws.append(r)
        wb.save(f"{DIR}/{name}.xlsx")
    d0 = datetime.date(2026,1,1)
    # 1. school marks
    names=["Aarav","Diya","Kabir","Meera","Rohan","Sana","Tara","Vivek","Zoya","Ishan"]
    rows=[[f"S{i:03}", random.choice(names)+" "+random.choice(["Sharma","Gill","Khan"]), random.choice(["6A","6B","7A","7B"]), random.choice(["Maths","Science","English"]), random.randint(30,100)] for i in range(1,121)]
    save("school",[("Marks",["Roll No","Student Name","Class","Subject","Marks"],rows)])
    # 2. clinic billing with rupee text amounts and dates
    rows=[[f"B-{i}", d0+datetime.timedelta(days=random.randint(0,280)), random.choice(["Dr Mehta","Dr Rao","Dr Singh"]), random.choice(["Consultation","X-Ray","Blood Test","ECG"]), f"₹ {random.choice([500,800,1200,1500,2500]):,}", random.choice(["Paid","Unpaid"])] for i in range(1,201)]
    save("clinic",[("Bills",["Bill No","Visit Date","Doctor","Service","Fee","Payment Status"],rows)])
    # 3. shop GST invoices + customers in 2 sheets
    cust=[[f"C{i}", random.choice(["Ram Traders","Shiv Stores","Gupta & Sons","Jain Mart","Kumar Agencies"])+f" {i}", random.choice(["Ludhiana","Jalandhar","Amritsar"])] for i in range(1,21)]
    inv=[[f"INV/{i}", random.choice(cust)[0], d0+datetime.timedelta(days=random.randint(0,300)), round(random.uniform(1000,90000),2), random.choice([5,12,18])] for i in range(1,301)]
    save("gst",[("Invoices",["Invoice No","Customer Code","Invoice Date","Taxable Value","GST Rate"],inv),("Customers",["Customer Code","Customer Name","City"],cust)])
    # 4. hindi-ish headers, warehouse
    rows=[[f"P{i}", random.choice(["Chawal","Atta","Dal","Cheeni","Tel"]), random.choice(["Godown 1","Godown 2"]), random.randint(0,500), random.randint(50,200)] for i in range(1,61)]
    save("godown",[("Stock",["Item ID","Item","Godown","Current Stock","Min Stock"],rows)])
    # 5. two monthly sheets same columns (combine)
    mk=lambda: [[f"E{i}", random.choice(["Asha","Bina","Chetan","Deep"]), random.randint(10,40)*1000] for i in range(1,26)]
    save("monthly",[("Jan",["Emp Code","Name","Commission"],mk()),("Feb",["Emp Code","Name","Commission"],mk())])
    return DIR


def _rows(DIR, f, s):
    ws = load_workbook(f"{DIR}/{f}.xlsx")[s]
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def _cases(DIR):
    rows = lambda f, s: _rows(DIR, f, s)  # noqa: E731
    money = lambda v: float(re.sub(r"[^\d.]", "", v)) if isinstance(v, str) else v  # noqa: E731
    fmt = ask._fmt
    M = rows("school", "Marks"); B = rows("clinic", "Bills"); I = rows("gst", "Invoices"); S = rows("godown", "Stock")  # noqa: E702

    def grp(rs, k, v, fn=sum):
        d = defaultdict(list)
        for r in rs:
            d[r[k]].append(v(r))
        return {a: fn(b) for a, b in d.items()}
    avg = lambda x: sum(x) / len(x)  # noqa: E731
    exp = {
     ("school","class wise average marks"): [f"{k}: {fmt(v)}" for k,v in grp(M,2,lambda r:r[4],avg).items()][:2],
     ("school","maths mein average marks kitne hain"): [fmt(avg([r[4] for r in M if r[3]=="Maths"]))],
     ("school","sabse zyada marks kisne liye"): [fmt(max(r[4] for r in M))],
     ("school","6A mein kitne students hain"): [str(sum(1 for r in M if r[2]=="6A"))],
     ("school","40 se kam marks wale highlight karo"): [str(sum(1 for r in M if r[4]<40))+" rows"],
     ("clinic","total fee kitni aayi"): [fmt(sum(money(r[4]) for r in B))],
     ("clinic","doctor wise total fee"): [f"{k}: {fmt(v)}" for k,v in sorted(grp(B,2,lambda r:money(r[4])).items(), key=lambda x:-x[1])][:1],
     ("clinic","unpaid bills dikhao"): [str(sum(1 for r in B if r[5]=="Unpaid"))+" rows"],
     ("clinic","month wise fee"): ["2026-01"],
     ("clinic","unpaid bills ka total"): [fmt(sum(money(r[4]) for r in B if r[5]=="Unpaid"))],
     ("gst","total taxable value"): [fmt(sum(r[3] for r in I))],
     ("gst","Customers se customer name Invoices mein laao"): ["Customer Name (from Customers)"],
     ("gst","invoice aging"): ["90+"],
     ("gst","kis customer ka taxable value sabse zyada hai"): [sorted(grp(I,1,lambda r:r[3]).items(), key=lambda x:-x[1])[0][0]],
     ("gst","18% gst wale invoices kitne hain"): [str(sum(1 for r in I if r[4]==18))],
     ("godown","godown wise current stock"): [f"{k}: {fmt(v)}" for k,v in grp(S,2,lambda r:r[3]).items()][:1],
     ("godown","jahan current stock min stock se kam hai highlight karo"): [str(sum(1 for r in S if r[3]<r[4]))+" rows"],
     ("godown","atta ka total stock"): [fmt(sum(r[3] for r in S if r[1]=="Atta"))],
     ("monthly","dono sheets combine karo"): ["50 rows"],
     ("monthly","jan mein total commission"): [fmt(sum(r[2] for r in rows("monthly","Jan")))],
    }
    return exp


def test_any_sheet(books):
    misses = []
    for (f, q), want in _cases(books).items():
        path = f"{books}/{f}.xlsx"
        try:
            plan = ask.understand(q, load_tables(path))
            got = ask.run(plan, path, f"{books}/out.xlsx", True, 1)["say"]
        except ask.NeedMore as e:
            got = "ASKED: " + e.hi
        if not all(w in got for w in want):
            misses.append(f"[{f}] {q!r} -> {got[:160]!r}, wanted {want}")
    assert not misses, "\n".join(misses)


def test_messy_hindi_and_csv_sheets(tmp_path):
    random.seed(3)
    wb = Workbook()
    ws = wb.active
    ws.title = "Expense Report"
    ws.append(["ABC Pvt Ltd - Expense Report FY 2026"])         # a title row above the real header
    ws.append([])
    ws.append(["  Date ", "Employee", None, " Category ", "Amount (Rs) ", "Approved?"])
    rows = []
    for _ in range(80):
        d = datetime.date(2026, 1, 1) + datetime.timedelta(days=random.randint(0, 250))
        r = [d.strftime("%d/%m/%Y"), random.choice(["Asha", "Ravi", "Neha"]), None, random.choice(["Travel", "Food", "Hotel"]),
             f"{random.randint(5, 90) * 100:,}", random.choice(["Yes", "No"])]
        ws.append(r)
        rows.append(r)
    wb.save(tmp_path / "messy.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "बिक्री"
    ws.append(["तारीख", "ग्राहक", "शहर", "राशि"])
    H = []
    for i in range(50):
        r = [datetime.date(2026, 1, 1) + datetime.timedelta(days=i * 3), random.choice(["राम", "श्याम", "गीता"]),
             random.choice(["दिल्ली", "पटना"]), random.randint(1, 50) * 100]
        ws.append(r)
        H.append(r)
    wb.save(tmp_path / "hindi.xlsx")
    lines = ["order_id,city,amount"] + [f"{i},{random.choice(['Pune', 'Delhi'])},{random.randint(100, 999)}" for i in range(100)]
    (tmp_path / "orders.csv").write_text("\n".join(lines) + "\n")
    C = [l.split(",") for l in lines[1:]]
    amt = lambda v: int(v.replace(",", ""))  # noqa: E731
    fmt = ask._fmt
    cases = [("messy.xlsx", "travel ka total amount", fmt(sum(amt(r[4]) for r in rows if r[3] == "Travel"))),
             ("messy.xlsx", "approved expenses kitne hain", str(sum(1 for r in rows if r[5] == "Yes"))),
             ("messy.xlsx", "not approved expenses ka total", fmt(sum(amt(r[4]) for r in rows if r[5] == "No"))),
             ("hindi.xlsx", "शहर wise राशि", fmt(sum(r[3] for r in H if r[2] == "दिल्ली"))),
             ("hindi.xlsx", "कुल राशि", fmt(sum(r[3] for r in H))),
             ("hindi.xlsx", "shahar wise rashi", "दिल्ली"),
             ("orders.csv", "city wise total amount", "Pune"),
             ("orders.csv", "delhi ka total amount", fmt(sum(int(r[2]) for r in C if r[1] == "Delhi")))]
    misses = []
    for f, q, want in cases:
        path = str(tmp_path / f)
        try:
            got = ask.run(ask.understand(q, load_tables(path)), path, str(tmp_path / "o.xlsx"), True, 1)["say"]
        except ask.NeedMore as e:
            got = "ASKED: " + e.hi
        if want not in got:
            misses.append(f"{f} {q!r} -> {got[:150]!r}, wanted {want}")
    assert not misses, "\n".join(misses)


def test_department_jobs(tmp_path):
    """Accounts / finance / sales / HR jobs from Chetan's list, on a fresh workbook."""
    random.seed(11)
    wb = Workbook()
    wb.remove(wb.active)
    led = wb.create_sheet("Ledger")
    led.append(["Date", "Particulars", "Amount"])
    bank = wb.create_sheet("Bank")
    bank.append(["Txn Date", "Narration", "Withdrawal"])
    missing_from_bank = 0
    for i in range(40):
        d = datetime.date(2026, 9, 1) + datetime.timedelta(days=random.randint(0, 29))
        a = random.randint(5, 500) * 100
        led.append([d, f"Pay {i}", a])
        if i % 8 != 0:
            bank.append([d + datetime.timedelta(days=random.randint(0, 3)), f"NEFT {i}", a])
        else:
            missing_from_bank += 1
    bank.append([datetime.date(2026, 9, 30), "Bank charges", 590])
    bud = wb.create_sheet("Budget")
    bud.append(["Department", "Budget", "Actual"])
    B = []
    for dpt in ["HR", "IT", "Sales", "Admin"]:
        r = [dpt, random.randint(5, 20) * 10000, random.randint(5, 20) * 10000]
        bud.append(r)
        B.append(r)
    st = wb.create_sheet("Staff")
    st.append(["First Name", "Last Name", "Region", "Product", "Sales", "Phone"])
    S = []
    for i in range(30):
        r = [random.choice(["Asha", "Ravi", "Neha"]), random.choice(["Gill", "Rao"]), random.choice(["North", "South"]),
             random.choice(["A", "B", "C"]), random.randint(1, 90) * 1000, None if i % 5 == 0 else "98xxxx"]
        st.append(r)
        S.append(r)
    path = str(tmp_path / "office.xlsx")
    wb.save(path)
    fmt = ask._fmt
    cases = [
        ("bank aur ledger reconcile karo", [f"Sirf Ledger mein: {missing_from_bank}", "Sirf Bank mein: 1 (590)"]),
        ("budget vs actual variance nikalo", [f"Kul farak {fmt(sum(r[2] - r[1] for r in B))}"]),
        ("full name banao", ["Full Name", '=TRIM(A2&" "&B2)']),
        ("sales ka 5% commission nikalo", ["Commission (5%)", fmt(sum(r[4] for r in S) * 5 / 100)]),
        ("region aur product wise sales", ["Region × Product"]),
        ("phone khaali kahan hai", [f"{sum(1 for r in S if r[5] is None)} rows mein Phone khaali"]),
        ("sales ke hisaab se sort karo bade se chhote", ["SORT(", ",5,-1)"]),
        ("ledger amount ka running total", ["Running total of Amount"]),
        ("file review karo", ["Ledger: 40 rows", "Phone: 6 khaali"]),
    ]
    misses = []
    for i, (q, want) in enumerate(cases, 1):
        try:
            r = ask.run(ask.understand(q, load_tables(path)), path, str(tmp_path / f"o{i}.xlsx"), True, i)
            got = r["say"] + " " + r["formula"]
        except ask.NeedMore as e:
            got = "ASKED: " + e.hi
        if not all(w in got for w in want):
            misses.append(f"{q!r} -> {got[:180]!r}, wanted {want}")
    assert not misses, "\n".join(misses)
    # the file the person downloads opens with the new columns and formulas in place
    out = load_workbook(str(tmp_path / "o4.xlsx"))
    assert out["Staff"]["G1"].value == "Commission (5%)" and out["Staff"]["G2"].value == "=E2*5%"
