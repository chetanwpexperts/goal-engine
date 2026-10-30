"""General data quality for any sheet: work out what each column IS from its values, then apply that kind's rules.

Nothing here knows a particular file. A column's kind (date, email, phone, number, rate, category, id, formula, text)
comes from what is in it, helped by its heading. Each kind has general rules:
  fix what is safe and certain (logged, with Undo), FLAG what only a person can decide (logged as "flag_*", never guessed).
"""

import difflib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl.utils import get_column_letter

EMAIL_OK = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,24}$")
MONEY_JUNK = re.compile(r"(?i)^\s*(rs\.?|inr|₹|\$|€|£)\s*|\s*(rs\.?|inr|/-|₹|\$)\s*$")
DATE_HINT = re.compile(r"(?i)\b(date|dt|dob|birth|joining|joined|due|day|delivery|dispatch|expiry|paid on|on)\b")
RATE_HINT = re.compile(r"(?i)(%|\b(discount|disc|tax|gst|vat|rate|percent|percentage|commission|interest|margin)\b)")
ID_HINT = re.compile(r"(?i)\b(id|no|num|number|code|invoice|bill|order|receipt|voucher|roll|ref|reference|sku)\b\.?")
EMAIL_HINT = re.compile(r"(?i)\b(e-?mail|mail)\b")
PHONE_HINT = re.compile(r"(?i)\b(phone|mobile|contact|whatsapp|cell|tel)\b")
TOTAL_ROW = re.compile(r"(?i)^\s*(grand\s+)?(total|sum|totals|kul|jod)\b")
NUM_DATE = re.compile(r"^\s*(\d{1,4})[/.-](\d{1,2})[/.-](\d{2,4})\s*$")
DAYFIRST_FMTS = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%y"]
MONTHFIRST_FMTS = ["%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y", "%m-%d-%y"]
OTHER_FMTS = ["%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d-%b-%y", "%d-%b-%Y", "%d %b %Y", "%d %b %y", "%d %B %Y", "%d-%B-%Y",
              "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%d %b, %Y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M"]
EXCEL_EPOCH = datetime(1899, 12, 30)


def _vocab() -> set[str]:
    words = set(json.loads((Path(__file__).resolve().parents[1] / "knowledge" / "data" / "heading_words.json").read_text()))
    try:
        from app.mind import packs

        def walk(x):
            if isinstance(x, str):
                words.update(w.lower() for w in re.findall(r"[A-Za-z]{3,}", x))
            elif isinstance(x, dict):
                for k, v in x.items():
                    walk(k)
                    walk(v)
            elif isinstance(x, list):
                for i in x:
                    walk(i)

        pack = packs.get("office")
        for key in ("column_synonyms", "concepts", "category_words"):
            walk(pack.get(key, {}))
    except Exception:  # noqa: BLE001 - the list on its own still works
        pass
    return words


# ------------------------------------------------------------------ reading values
def as_number(v):
    """A number, or None. '₹1,499', 'Rs 200/-', ' 1,200 ', '12.5' -> number. Never a date or a code like 'INV-12'."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if not isinstance(v, str):
        return None
    s = MONEY_JUNK.sub("", v.strip()).replace(",", "").strip()
    if not re.fullmatch(r"[-+]?\d+(\.\d+)?", s):
        return None
    n = float(s)
    return int(n) if n.is_integer() and "." not in s else n


def as_rate(v):
    """A rate as a fraction or None: '10%' -> 0.1, 0.1 -> 0.1, '0.10' -> 0.1; plain 5 is returned as 5 (scale decided by the column)."""
    if isinstance(v, str) and v.strip().endswith("%"):
        n = as_number(v.strip()[:-1])
        return None if n is None else n / 100
    return as_number(v)


def parse_date(v, dayfirst=True):
    """(date, how) with how in ok/fixed, or (None, 'impossible'/'unknown')."""
    if isinstance(v, datetime):
        return v, "ok"
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day), "ok"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 <= v <= 80000:                         # an Excel date serial shown as a number
            return EXCEL_EPOCH + timedelta(days=float(v)), "fixed"
        return None, "unknown"
    if not isinstance(v, str) or not v.strip():
        return None, "unknown"
    s = re.sub(r"\s+", " ", v.strip().replace("st ", " ").replace("nd ", " ").replace("rd ", " ").replace("th ", " "))
    fmts = (DAYFIRST_FMTS + MONTHFIRST_FMTS if dayfirst else MONTHFIRST_FMTS + DAYFIRST_FMTS) + OTHER_FMTS
    for f in fmts:
        try:
            d = datetime.strptime(s, f)
        except ValueError:
            continue
        if f.endswith("%y") and d.year > datetime.now().year + 20:
            d = d.replace(year=d.year - 100)
        return d, "fixed"
    m = NUM_DATE.match(s)
    if m:                                               # looks like a date, but no calendar has it (31/02/2026)
        return None, "impossible"
    return None, "unknown"


def _dayfirst(values) -> bool:
    """Does this column write the day first? Any '25/03/...' says yes, any '03/25/...' says no; India's default is day first."""
    first_big = second_big = 0
    for v in values:
        if isinstance(v, str):
            m = NUM_DATE.match(v)
            if m and len(m.group(1)) <= 2:
                a, b = int(m.group(1)), int(m.group(2))
                first_big += a > 12
                second_big += b > 12
    return not (second_big > first_big)


# ------------------------------------------------------------------ the shape of a sheet
def _header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = [c.value for c in ws[r]]
        filled = [v for v in vals if v not in (None, "")]
        texts = [v for v in filled if isinstance(v, str) and not v.startswith("=") and as_number(v) is None]
        if len(filled) >= 2 and len(texts) >= 0.6 * len(filled) and len(filled) >= 0.5 * max(1, sum(1 for v in vals if v is not None)):
            nxt = [c.value for c in ws[r + 1]] if r < ws.max_row else []
            if sum(1 for v in nxt if v not in (None, "")) >= 2:
                return r
    return None


def _data_rows(ws, hr: int) -> list[int]:
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        vals = [c.value for c in ws[r]]
        if not any(v not in (None, "") for v in vals):
            continue
        first_text = next((v for v in vals if isinstance(v, str) and v.strip() and not v.startswith("=")), "")
        if TOTAL_ROW.match(str(first_text)):
            continue                                    # a totals row is not data
        rows.append(r)
    return rows


def _kind(header: str, vals: list, pct_format: float = 0) -> str:
    filled = [v for v in vals if v not in (None, "")]
    if not filled:
        return "empty"
    n = len(filled)
    h = header or ""
    formulas = sum(1 for v in filled if isinstance(v, str) and v.startswith("="))
    if formulas >= 0.5 * n:
        return "formula"
    plain = [v for v in filled if not (isinstance(v, str) and v.startswith("="))]
    df = _dayfirst(plain)
    dates = 0
    for v in plain:
        if isinstance(v, (int, float)) and not isinstance(v, bool) and not DATE_HINT.search(h):
            continue                                    # a plain number is a date only in a column whose heading says so
        d, how = parse_date(v, df)
        dates += d is not None or how == "impossible"
    if dates >= 0.6 * n or (DATE_HINT.search(h) and dates >= 0.3 * n):
        return "date"
    if EMAIL_HINT.search(h) or sum(1 for v in plain if isinstance(v, str) and "@" in v) >= 0.4 * n:
        return "email"
    if PHONE_HINT.search(h):
        return "phone"
    codes = sum(1 for v in plain if re.fullmatch(r"\+?\d{8,}|0\d+", str(v).strip()))
    if codes >= 0.6 * n:
        return "phone"                                  # long digit strings / leading zeros: phones, accounts, codes - never maths
    nums = [as_rate(v) for v in plain]
    numeric = sum(1 for x in nums if x is not None)
    pct_text = sum(1 for v in plain if isinstance(v, str) and v.strip().endswith("%"))
    if numeric >= 0.6 * n and (RATE_HINT.search(h) or pct_text >= 0.3 * n or pct_format >= 0.5):
        return "rate"
    if numeric >= 0.6 * n:
        texts = [v for v in plain if isinstance(v, str) and as_number(v) is None]
        if ID_HINT.search(h) and len({str(v) for v in plain}) >= 0.9 * n and not texts:
            return "id"
        return "number"
    sv = [str(v).strip() for v in plain]
    distinct = len({_norm(v) for v in sv})
    if ID_HINT.search(h) and distinct >= 0.8 * n:
        return "id"
    if n >= 4 and distinct <= max(2, int(0.5 * n)) and max(len(v) for v in sv) <= 40:
        return "category"
    return "text"


def _one_slip(a: str, b: str) -> bool:
    """One typing slip apart (a letter swapped with the next, missing, extra or wrong) in words long enough that it
    can't be a different name. Never across digits or one-letter parts ('Class A' / 'Class B', 'Unit 1' / 'Unit 2')."""
    if min(len(a), len(b)) < 5 or abs(len(a) - len(b)) > 1 or a == b:
        return False
    i = 0
    while i < min(len(a), len(b)) and a[i] == b[i]:
        i += 1
    ja, jb = len(a), len(b)
    while ja > i and jb > i and a[ja - 1] == b[jb - 1]:
        ja, jb = ja - 1, jb - 1
    da, db = a[i:ja], b[i:jb]
    if any(ch.isdigit() for ch in da + db):
        return False
    one = (len(da) <= 1 and len(db) <= 1) or (len(da) == len(db) == 2 and da == db[::-1])
    if not one:
        return False
    ta, tb = a.split(), b.split()
    return len(ta) != len(tb) or all(len(x) >= 3 for x, y in zip(ta, tb) if x != y)


def _norm(v: str) -> str:
    s = re.sub(r"\s+", " ", str(v).strip().lower())
    if len(s) > 4 and s.endswith("ies"):
        s = s[:-3] + "y"
    elif len(s) > 4 and s.endswith(("ches", "shes", "sses", "xes")):
        s = s[:-2]
    elif len(s) > 3 and s.endswith("s") and not s.endswith("ss"):
        s = s[:-1]
    return s


# ------------------------------------------------------------------ the rules
def check_sheet(ws, log, fix=True) -> dict:
    """Apply the general rules to one sheet. Returns what each column was read as."""
    from app.office import best_form

    hr = _header_row(ws)
    if hr is None:
        return {}
    rows = _data_rows(ws, hr)
    if len(rows) < 3:
        return {}
    sheet = ws.title
    width = ws.max_column
    headers = {c: (str(ws.cell(hr, c).value).strip() if ws.cell(hr, c).value not in (None, "") else "") for c in range(1, width + 1)}
    if fix:
        _headings(ws, hr, headers, log)               # first, so 'Moble' is read as the phone column it is
        headers = {c: (str(ws.cell(hr, c).value).strip() if ws.cell(hr, c).value not in (None, "") else "") for c in headers}
    kinds = {}
    for c in range(1, width + 1):
        if not headers[c]:
            continue
        vals = [ws.cell(r, c).value for r in rows]
        pct = sum(1 for r in rows if "%" in (ws.cell(r, c).number_format or "")) / max(1, len(rows))
        kinds[c] = _kind(headers[c], vals, pct)

    for c, kind in kinds.items():
        cells = [ws.cell(r, c) for r in rows]
        filled = [x for x in cells if x.value not in (None, "")]
        mostly_filled = len(filled) >= 0.7 * len(cells)
        name = headers[c]
        if kind == "date":
            df = _dayfirst([x.value for x in filled])
            fmt = next((x.number_format for x in filled if isinstance(x.value, (datetime, date)) and x.number_format != "General"), "DD-MM-YYYY")
            for x in cells:
                if x.value in (None, ""):
                    if mostly_filled:
                        log.add("flag_missing", sheet, x.coordinate, "", name)
                    continue
                d, how = parse_date(x.value, df)
                if d is None:
                    log.add("flag_bad_date", sheet, x.coordinate, x.value, name)
                elif how == "fixed" and fix:
                    log.add("fix_dates", sheet, x.coordinate, x.value, d.strftime("%d-%m-%Y"))
                    x.value, x.number_format = d, fmt
        elif kind == "email":
            for x in cells:
                v = x.value
                if v in (None, ""):
                    if mostly_filled:
                        log.add("flag_missing", sheet, x.coordinate, "", name)
                elif not EMAIL_OK.match(str(v).strip()):
                    log.add("flag_bad_email", sheet, x.coordinate, v, name)
        elif kind in ("number", "rate"):
            parsed = {x.coordinate: (as_rate(x.value) if kind == "rate" else as_number(x.value)) for x in filled}
            good = [p for p in parsed.values() if p is not None]
            fraction_scale = kind == "rate" and sum(1 for p in good if 0 <= p <= 1) >= 0.6 * max(1, len(good))
            positive = sum(1 for p in good if p >= 0) >= 0.8 * max(1, len(good))
            for x in cells:
                v = x.value
                if v in (None, ""):
                    if mostly_filled:
                        log.add("flag_missing", sheet, x.coordinate, "", name)
                    continue
                p = parsed.get(x.coordinate)
                if p is None:
                    log.add("flag_not_number", sheet, x.coordinate, v, name)
                    continue
                new = p
                if kind == "rate" and fraction_scale and 1 < p <= 100:
                    new = p / 100                         # '5' in a column of 0.05s and 10%s means 5%
                if new != v and fix:
                    log.add("fix_rates" if kind == "rate" else "text_to_number", sheet, x.coordinate, v,
                            f"{new:.0%}" if kind == "rate" and fraction_scale else new)
                    x.value = new
                    if kind == "rate" and fraction_scale:
                        x.number_format = "0%"
                if positive and isinstance(new, (int, float)) and new < 0:
                    log.add("flag_negative", sheet, x.coordinate, v, name)
        elif kind == "category":
            groups = defaultdict(Counter)
            for x in filled:
                groups[_norm(x.value)][re.sub(r"\s+", " ", str(x.value).strip())] += 1
            main = {k: best_form(cnt) for k, cnt in groups.items()}
            size = {k: sum(cnt.values()) for k, cnt in groups.items()}
            keys = sorted(size, key=lambda k: -size[k])
            merge, maybe = {}, {}
            for k in keys:                                # a rare spelling close to a common one is that one ('Electronic', 'Electornics')
                for big in keys:
                    if big == k or size[big] < size[k] or big in merge or (size[big] == size[k] and big in maybe):
                        continue
                    rk, rb = main[k].lower(), main[big].lower()
                    slip = (_one_slip(k, big) or _one_slip(rk, rb) or
                            (rk.replace(" ", "") == rb.replace(" ", "") and rk != rb))
                    if size[big] > size[k] and len(k) >= 4 and (difflib.SequenceMatcher(None, k, big).ratio() >= 0.86 or
                                                                 (size[k] * 2 <= size[big] and slip)):
                        merge[k] = big
                        break
                    if slip and len(k) >= 4 and k not in maybe:   # as common as each other: can't know which is right
                        maybe[k] = big
            for k, cnt in groups.items():                 # 'Appliance' x2, 'Appliances' x2: either could be right
                tops = cnt.most_common(2)
                if len(tops) == 2 and tops[0][1] == tops[1][1] and tops[0][0].lower() != tops[1][0].lower():
                    maybe.setdefault(k, k)
            for k, big in maybe.items():
                if k in merge or big in merge:
                    continue
                forms = sorted({f for g in (k, big) for f in groups[g]}, key=str.lower)
                for x in filled:
                    if _norm(x.value) in (k, big):
                        others = [f for f in forms if f != re.sub(r"\s+", " ", str(x.value).strip())]
                        log.add("flag_maybe_same", sheet, x.coordinate, x.value, f"{name}: same as '{others[0]}'?" if others else name)
            unsure = {g for pair in maybe.items() for g in pair if pair[0] not in merge and pair[1] not in merge}
            mixed = sum(1 for x in filled if any(ch.isupper() for ch in str(x.value)) and any(ch.islower() for ch in str(x.value)))
            styled = mixed >= 0.8 * len(filled)
            for x in filled:
                k = _norm(x.value)
                if k in unsure:
                    continue                               # asked about, not changed
                target = main[merge.get(k, k)]
                letters = [ch for ch in target if ch.isalpha()]
                if styled and letters and (target == target.lower() or (target == target.upper() and len(letters) >= 5)):
                    target = target.title()                # a lone 'west' / 'CENTRAL' among 'North', 'South'...
                if str(x.value) != target and fix:
                    log.add("fix_categories", sheet, x.coordinate, x.value, target)
                    x.value = target
        elif kind == "id":
            seen = {}
            for x in filled:
                key = str(x.value).strip().lower()
                if key in seen:
                    log.add("flag_repeated_id", sheet, x.coordinate, x.value, f"{name} (same as {seen[key]})")
                else:
                    seen[key] = x.coordinate
        elif kind == "formula":
            _formula_column(ws, c, rows, cells, name, log, fix)

    return kinds


def _formula_column(ws, c, rows, cells, name, log, fix):
    """The column's own formula is the one most of its rows use. Rows with another formula, a typed-in number or nothing
    get that formula back (wrong reference, hard-coded total, missing total)."""
    from app.office import formula_from_shape, formula_shape

    shapes = Counter(formula_shape(x.value, x.row, x.column) for x in cells if isinstance(x.value, str) and x.value.startswith("="))
    if not shapes:
        return
    main, n = shapes.most_common(1)[0]
    if n < 3 or n / sum(shapes.values()) < 0.6:
        return
    for x in cells:
        v = x.value
        if isinstance(v, str) and v.startswith("="):
            if formula_shape(v, x.row, x.column) == main:
                continue
            what = "fix_formulas"
        elif v in (None, ""):
            row_has_data = sum(1 for cc in range(1, ws.max_column + 1) if cc != c and ws.cell(x.row, cc).value not in (None, "")) >= 2
            if not row_has_data:
                continue
            what = "fix_missing_formula"
        else:
            what = "fix_typed_total"
        new = formula_from_shape(main, x.row, x.column)
        if fix:
            log.add(what, ws.title, x.coordinate, v, new)
            x.value = new


def _headings(ws, hr, headers, log):
    vocab = _vocab()
    for c, h in headers.items():
        if not h:
            continue
        parts = re.split(r"(\W+)", h)
        changed = False
        for i, w in enumerate(parts):
            lw = w.lower()
            if len(lw) < 5 or not lw.isalpha() or lw in vocab:
                continue
            m = difflib.get_close_matches(lw, vocab, n=1, cutoff=0.85)
            if m and abs(len(m[0]) - len(lw)) <= 2:
                parts[i] = m[0].capitalize() if w[:1].isupper() else m[0]
                changed = True
        if changed:
            new = "".join(parts)
            log.add("fix_headings", ws.title, ws.cell(hr, c).coordinate, h, new)
            ws.cell(hr, c).value = new


def column_letter(c: int) -> str:
    return get_column_letter(c)
