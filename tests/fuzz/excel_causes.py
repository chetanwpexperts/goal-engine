"""Generator for the Excel reasoning: real-looking workbooks with known problems (alone and in COMBINATIONS no single
playbook describes), with the values Excel would have saved. The agent must find exactly the causes that are there,
prove them, and its fix must make them go away. Run: python -m tests.fuzz.excel_causes [count] [seed]
"""

import random
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook

CAUSES = ["text_in_totals", "short_range", "hidden_in_totals", "typed_over_formula", "manual_calc", "lookup_spaces",
          "lookup_type", "lookup_missing", "approx_lookup", "broken_refs", "div0", "text_dates"]
EXCLUSIVE = [{"approx_lookup", "lookup_spaces", "lookup_type", "lookup_missing"}, {"lookup_spaces", "lookup_type"},
             {"short_range", "hidden_in_totals"}]
SIDE = {"text_in_totals": {"text_numbers"}, "lookup_spaces": {"spaces"}, "lookup_type": {"text_numbers"}}
SAY = {
    "total_wrong": ["my total is wrong", "total galat aa raha hai", "the sum is less than it should be", "grand total not matching",
                    "sum kam aa raha hai", "total doesn't add up"],
    "lookup_fails": ["vlookup is showing #N/A", "lookup not finding the customers", "vlookup kaam nahi kar raha #N/A aa raha",
                     "my vlookup gives n/a for some rows"],
    "wrong_value": ["vlookup returns the wrong value", "lookup galat value la raha hai", "vlookup brings wrong discount"],
    "not_updating": ["formulas are not updating", "values don't change when I edit", "formula update nahi ho raha"],
    "formula_error": ["some formulas show #DIV/0! and #REF!", "formula errors in my sheet", "#ref aa raha hai formula mein"],
    "sort_wrong": ["dates are not sorting properly", "date filter not working right", "sort karne par dates galat order mein"],
}
COMPLAINT = {"text_in_totals": "total_wrong", "short_range": "total_wrong", "hidden_in_totals": "total_wrong",
             "typed_over_formula": "not_updating", "manual_calc": "not_updating", "lookup_spaces": "lookup_fails",
             "lookup_type": "lookup_fails", "lookup_missing": "lookup_fails", "approx_lookup": "wrong_value",
             "broken_refs": "formula_error", "div0": "formula_error", "text_dates": "sort_wrong"}


def pick(rng: random.Random, k: int) -> set:
    while True:
        s = set(rng.sample(CAUSES, k))
        if not any(len(s & e) > 1 for e in EXCLUSIVE):
            return s


def build(path: str, inject: set, rng: random.Random) -> dict:
    """Write the workbook; returns what the agent should find and the right total."""
    n = rng.randint(14, 40)
    look = rng.choice(["vlookup", "xlookup", "indexmatch"] if "approx_lookup" not in inject else ["vlookup", "indexmatch"])
    layout = rng.choice(["below", "summary"])
    numeric = "lookup_type" in inject
    codes = [101 + i for i in range(10)] if numeric else [f"C{i:03d}" for i in range(1, 11)]
    disc = {c: rng.choice([0, 5, 10, 12, 15]) for c in codes}
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    rates = wb.create_sheet("Rates")
    rates.append(["Customer", "Discount"])
    for c in codes:
        rates.append([c, disc[c]])
    ws.append(["Invoice", "Customer", "Region", "Qty", "Price", "Amount", "Discount", "Date", "Tax", "Per unit"])
    cache = {"Sales": {}, "Rates": {}}
    amounts = []
    start = date(2026, 4, 1)
    rows = list(range(2, n + 2))
    text_rows = set(rng.sample(rows, 2)) if "text_in_totals" in inject else set()
    space_rows = set(rng.sample(rows, 2)) if "lookup_spaces" in inject else set()
    type_rows = set(rng.sample(rows, 2)) if "lookup_type" in inject else set()
    miss_rows = set(rng.sample(rows, 2)) if "lookup_missing" in inject else set()
    typed_row = rng.choice(rows) if "typed_over_formula" in inject else None
    ref_rows = set(rng.sample(rows, 2)) if "broken_refs" in inject else set()
    div_rows = set(rng.sample(rows, 2)) if "div0" in inject else set()
    date_rows = set(rng.sample(rows, 2)) if "text_dates" in inject else set()
    for r in rows:
        qty, price = rng.randint(1, 20), rng.choice([120, 250, 499, 1200, 75])
        if r in div_rows:
            qty = 0
        amt = qty * price
        amounts.append(amt)
        code = rng.choice(codes)
        key = code
        if r in space_rows:
            key = code + " "
        elif r in type_rows:
            key = str(code)
        elif r in miss_rows:
            key = 999 if numeric else "C999"
        d = start + timedelta(days=rng.randint(0, 120))
        ws.append([f"INV-{1000 + r}", key, rng.choice(["North", "South", "East", "West"]), qty, price,
                   f"{amt:,}" if r in text_rows else amt,
                   _lookup(look, r, "approx_lookup" in inject),
                   d.strftime("%d/%m/%Y") if r in date_rows else d,
                   round(amt * 0.18, 2) if r == typed_row else ("=#REF!*0.18" if r in ref_rows else f"=F{r}*0.18"),
                   f"=F{r}/D{r}"])
        found = key == code
        cache["Sales"][f"G{r}"] = disc[code] if found or "approx_lookup" in inject else "#N/A"
        if r != typed_row:
            cache["Sales"][f"I{r}"] = "#REF!" if r in ref_rows else round(amt * 0.18, 2)
        cache["Sales"][f"J{r}"] = "#DIV/0!" if qty == 0 else round(amt / qty, 4)
    last = n + 1
    end = last - 3 if "short_range" in inject else last
    shown = sum(a for r, a in zip(rows, amounts) if r <= end and r not in text_rows)
    if layout == "below":
        tot_row = n + 3
        ws.cell(row=tot_row, column=5, value="Total")
        ws.cell(row=tot_row, column=6, value=f"=SUM(F2:F{end})")
        cache["Sales"][f"F{tot_row}"] = shown
        total_cell = ("Sales", f"F{tot_row}")
    else:
        summ = wb.create_sheet("Summary")
        summ.append(["Total sales", f"=SUM(Sales!F2:F{end})"])
        cache["Summary"] = {"B1": shown}
        total_cell = ("Summary", "B1")
    if "hidden_in_totals" in inject:
        for r in rng.sample(rows, 2):
            ws.row_dimensions[r].hidden = True
    if "manual_calc" in inject:
        wb.calculation.calcMode = "manual"
    wb.save(path)
    _with_cache(path, wb.sheetnames, cache)
    return {"right_total": sum(amounts), "shown_total": shown, "total_cell": total_cell}


def _lookup(kind: str, r: int, approx: bool) -> str:
    if kind == "xlookup":
        return f"=XLOOKUP(B{r},Rates!$A$2:$A$11,Rates!$B$2:$B$11)"
    if kind == "indexmatch":
        return f"=INDEX(Rates!$B$2:$B$11,MATCH(B{r},Rates!$A$2:$A$11{'' if approx else ',0'}))"
    return f"=VLOOKUP(B{r},Rates!$A$2:$B$11,2{'' if approx else ',FALSE'})"


def _with_cache(path: str, order: list[str], cache: dict) -> None:
    """Put in the values Excel saves next to each formula (openpyxl writes formulas only)."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", item.filename)
            if m and int(m.group(1)) <= len(order):
                vals = cache.get(order[int(m.group(1)) - 1], {})
                xml = data.decode("utf-8")

                def put(mm):
                    addr, attrs, f = mm.group(1), re.sub(r'\s+t="[^"]*"', "", mm.group(2)), mm.group(3)
                    if addr not in vals:
                        return mm.group(0)
                    v = vals[addr]
                    if isinstance(v, str) and v.startswith("#"):
                        return f'<c r="{addr}"{attrs} t="e">{f}<v>{escape(v)}</v></c>'
                    if isinstance(v, str):
                        return f'<c r="{addr}"{attrs} t="str">{f}<v>{escape(v)}</v></c>'
                    return f'<c r="{addr}"{attrs}>{f}<v>{v}</v></c>'
                xml = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)>(<f>.*?</f>)(?:<v>.*?</v>|<v\s*/>)?</c>', put, xml)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def case(i: int, seed: int, folder: str):
    rng = random.Random(seed * 1000 + i)
    k = rng.choice([0, 1, 1, 2, 2, 3])
    inject = pick(rng, k) if k else set()
    path = str(Path(folder) / f"case{i}.xlsx")
    truth = build(path, inject, rng)
    if inject:
        said = {COMPLAINT[c] for c in inject}
    else:
        said = {rng.choice(list(SAY))}
    text = " and ".join(rng.choice(SAY[c]) for c in sorted(said))
    return path, inject, text, truth


def expected(inject: set, said: dict) -> set:
    from app.mind.reason import FOUND_MIN
    from app.officeknow import diagnose

    k = diagnose.knowledge()["causes"]
    present = set(inject).union(*(SIDE.get(c, set()) for c in inject)) if inject else set()
    return {c for c in present if max((k[c]["explains"].get(s, 0) * w for s, w in said.items()), default=0) >= FOUND_MIN}


def check(i: int, seed: int, folder: str) -> list[str]:
    """[] when the agent got it exactly right; otherwise what went wrong."""
    from app import office
    from app.officeknow import diagnose, evidence

    path, inject, text, truth = case(i, seed, folder)
    report = office.scan(path)
    facts = evidence.collect(path, report)
    res = diagnose.diagnose(text, facts)
    found = {x["cause"] for x in res["found"]}
    want = expected(inject, res["complaints"])
    errs = []
    if found != want:
        errs.append(f"case {i}: '{text}' injected {sorted(inject)}: found {sorted(found)}, expected {sorted(want)}")
        return errs
    msg, ops = diagnose.report(res, False, "case.xlsx")
    if not found and "none of them is there" not in msg:
        errs.append(f"case {i}: nothing found but the reply doesn't say so")
    if ops:
        out = path.replace(".xlsx", "-fixed.xlsx")
        office.fix(path, out, ops)
        facts2 = evidence.collect(out, office.scan(out))
        for c in found:
            fx = diagnose.knowledge()["causes"][c].get("fix")
            if fx and facts2.get(diagnose.knowledge()["causes"][c]["shown_by"], {}).get("seen"):
                errs.append(f"case {i}: {c} still there after the fix {fx}")
        if found & {"text_in_totals", "short_range"} and not ({"hidden_in_totals"} & inject):
            from openpyxl import load_workbook

            wb = load_workbook(out)
            ws = wb["Sales"]
            f = wb[truth["total_cell"][0]][truth["total_cell"][1]].value
            m = re.match(r"=SUM\((?:Sales!)?F(\d+):F(\d+)\)", f or "")
            got = sum(ws.cell(row=r, column=6).value or 0 for r in range(int(m.group(1)), int(m.group(2)) + 1)
                      if isinstance(ws.cell(row=r, column=6).value, (int, float))) if m else None
            if got != truth["right_total"]:
                errs.append(f"case {i}: total after fix {got}, should be {truth['right_total']}")
    return errs


def run(count: int = 200, seed: int = 1) -> tuple[int, list[str]]:
    folder = tempfile.mkdtemp(prefix="excel-causes-")
    errs, ok = [], 0
    try:
        for i in range(count):
            e = check(i, seed, folder)
            errs += e
            ok += not e
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    return ok, errs


if __name__ == "__main__":
    c = int(sys.argv[1]) if len(sys.argv) > 1 else 200
    s = int(sys.argv[2]) if len(sys.argv) > 2 else 1
    ok, errs = run(c, s)
    print(f"{ok}/{c} exactly right")
    for e in errs[:30]:
        print(" ", e)
