"""Word documents: upload a .docx, ask about its text and its tables, change it (new version, undo), in English and Hinglish."""

import io

import pytest
from docx import Document

from app.office import ask, word
from tests.make_docx import build
from tests.test_ml import say
from tests.test_office_agent import _upload, _wait


@pytest.fixture(scope="module")
def doc(tmp_path_factory):
    p = tmp_path_factory.mktemp("w") / "report.docx"
    build(str(p))
    return str(p)


def _ask(q, path, out):
    p = word.plan(q, path)
    if p is None:
        try:
            p = ask.understand(q, ask.load_tables(path))
        except ask.NeedMore:
            p = word.fallback(q, path)
    return ask.run(p, path, out, False, 1)


@pytest.mark.parametrize("q,must", [
    ("how many words", "160 words"), ("kitne tables hain", "2 tables"), ("what are the headings", "Hiring plan"),
    ("summarize this document", "South region"), ("is document ka saar batao", "Key points"),
    ("where is Delhi mentioned", "appears 5 times"), ("Mumbai kahan likha hai", "\"Mumbai\" appears 2 times"),
    ("what is the deadline", "31 Dec 2026"), ("what does it say about training", "January"), ("risk kya hai", "late payments"),
    ("list all dates", "05-10-2026"), ("saare emails batao", "finance@acme.in"), ("phone number kya hai", "98765 43210"),
    ("list all amounts", "₹34,14,200"), ("total sales by region", "South: 12,33,200"), ("average salary", "53,333"),
    ("who is the manager of south", "Priya"),
])
def test_questions_on_a_document(doc, tmp_path, q, must):
    r = _ask(q, doc, str(tmp_path / "o.docx"))
    assert must in r["say"], (q, r["say"])


def test_changes_keep_the_rest(doc, tmp_path):
    out = str(tmp_path / "o.docx")
    r = _ask("replace Delhi with New Delhi", doc, out)
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "5 places" in r["say"] and "New New Delhi" not in text and text.count("New Delhi") == 3
    r = _ask("fix the spacing", doc, out)
    d = Document(out)
    assert "2 extra blank lines" in r["say"] and not any("  " in p.text for p in d.paragraphs)
    r = _ask("make headings uppercase", doc, out)
    assert "HIRING PLAN" in [p.text for p in Document(out).paragraphs]


def test_old_doc_and_fake_docx_are_explained(tmp_path):
    from app import office

    p = tmp_path / "old.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    with pytest.raises(office.OfficeError, match="Save As"):
        office.kind_of(str(p))
    q = tmp_path / "fake.docx"
    q.write_bytes(b"hello")
    with pytest.raises(office.OfficeError, match="real Word file"):
        office.kind_of(str(q))


def test_word_file_in_chat_end_to_end(client, ml_mode, doc):
    r = say(client, "I need help with a word document")
    agent = next(a for a in r["agents"] if a.get("office"))
    f = _upload(client, agent["id"], open(doc, "rb").read(), "report.docx")
    assert f["status"] == "ready" and f["report"]["kind"] == "docx" and f["report"]["summary"]["words"] == 160
    gid = r["goal"]["id"]
    r = say(client, "summarize this document", gid)
    assert "Key points" in r["reply"], r["reply"]
    r = say(client, "total sales by region", gid)
    assert "South: 12,33,200" in r["reply"], r["reply"]
    r = say(client, "replace Delhi with New Delhi", gid)
    assert "5 places" in r["reply"], r["reply"]
    f = _wait(client, f["id"])
    assert len(f["steps"]) == 1 and "Replaced" in f["steps"][0]["say"]
    d = Document(io.BytesIO(client.get(f"/api/office/files/{f['id']}/download").content))
    assert "New Delhi" in "\n".join(p.text for p in d.paragraphs)
    say(client, "undo", gid)
    f = _wait(client, f["id"])
    assert f["steps"] == [] and f["report"]["kind"] == "docx"
    r = say(client, "what is the weather on mars", gid)
    assert "Summarise this document" in r["options"] or "doesn't say" in r["reply"], r
