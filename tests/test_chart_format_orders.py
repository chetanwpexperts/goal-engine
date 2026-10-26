"""Charts and conditional-format orders must change the saved workbook, not just promise a result."""

from collections import defaultdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.office import ask, edits
from app.office.sheetdata import Table, load_tables


PRACTICE = Path(__file__).parent / "data" / "practice.xlsx"


def _do(text, tmp_path):
    tables = load_tables(str(PRACTICE))
    plan = ask.understand(text, tables)
    assert plan["job"] == "edit", (text, plan)
    out = tmp_path / (str(abs(hash(text))) + ".xlsx")
    ask.run(plan, str(PRACTICE), str(out), False, 1)
    return plan, load_workbook(out)


@pytest.mark.parametrize(
    "text,kind",
    [
        ("sales by region chart", "bar"),
        ("month wise sales line chart", "line"),
        ("pie chart of sales by region", "pie"),
        ("region wise bar graph bana do", "bar"),
    ],
)
def test_chart_phrases_make_chart_plans(text, kind):
    plan = ask.understand(text, load_tables(str(PRACTICE)))
    assert plan["op"] == "chart" and plan["chart_type"] == kind


def test_region_chart_has_python_checked_summary_and_real_chart(tmp_path):
    plan, wb = _do("sales by region chart", tmp_path)
    source = next(t for t in load_tables(str(PRACTICE)) if t.name == "Sales Data")
    expected = defaultdict(float)
    for row in source.rows:
        expected[row[plan["group_col"]]] += row[plan["value_col"]]

    ws = wb[next(name for name in wb.sheetnames if name.startswith("Chart - "))]
    actual = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    assert actual == dict(expected)
    assert len(ws._charts) == 1
    assert ws._charts[0].x_axis.title is not None and ws._charts[0].y_axis.title is not None


def test_month_line_chart_and_chart_next_to_data(tmp_path):
    _, wb = _do("month wise sales line chart", tmp_path)
    ws = wb[next(name for name in wb.sheetnames if name.startswith("Chart - "))]
    assert len(ws._charts) == 1 and ws["A1"].value == "Month"

    _, wb = _do("sales by region chart next to the data", tmp_path)
    assert len(wb["Sales Data"]._charts) == 1
    assert not any(name.startswith("Chart - ") for name in wb.sheetnames)


def test_pie_chart_has_data_labels(tmp_path):
    _, wb = _do("pie chart of sales by region", tmp_path)
    chart = wb[next(name for name in wb.sheetnames if name.startswith("Chart - "))]._charts[0]
    assert chart.dataLabels.showPercent is True


def test_expense_chart_uses_named_columns_and_missing_column_asks():
    table = Table("Expenses", ["Expense Date", "Category", "Expense Amount"], [
        [None, "Travel", 100], [None, "Food", 50], [None, "Travel", 25],
    ], 1)
    plan = edits.plan_edit("pie chart of expenses by category", [table])
    assert plan["op"] == "chart" and plan["group_col"] == 1 and plan["value_col"] == 2
    with pytest.raises(ask.NeedMore) as error:
        edits.plan_edit("sales by territory chart", load_tables(str(PRACTICE)))
    assert "Which one" in error.value.en and error.value.options


@pytest.mark.parametrize(
    "text,rule,value,colour",
    [
        ("highlight sales amounts above 100000 in red", "above", 100000, "red"),
        ("1 lakh se upar wale sales green karo", "above", 100000, "green"),
        ("highlight sales below 25k in orange", "below", 25000, "orange"),
        ("highlight sales equal to 1 crore in blue", "equal", 10000000, "blue"),
        ("highlight sales between 50k and 2 lakh in yellow", "between", None, "yellow"),
        ("highlight overdue", "overdue", None, "red"),
        ("highlight duplicates in Invoice", "duplicate", None, "yellow"),
        ("highlight blank cells in Column1", "blank", None, "yellow"),
        ("colour scale on sales", "color_scale", None, "yellow"),
        ("top 10 sales highlight karo", "top10", None, "yellow"),
        ("status Pending wale yellow", "text_equal", "Pending", "yellow"),
    ],
)
def test_conditional_format_phrases_make_right_plan(text, rule, value, colour):
    plan = ask.understand(text, load_tables(str(PRACTICE)))
    assert plan["op"] == "conditional_format" and plan["rule"] == rule and plan["colour"] == colour
    if value is not None:
        assert plan["value"] == value
    if rule == "between":
        assert plan["values"] == [50000, 200000]


@pytest.mark.parametrize(
    "text,sheet,expected_range,rule_type",
    [
        ("highlight sales amounts above 100000 in red", "Sales Data", "J2:J61", "cellIs"),
        ("highlight overdue", "Sales Data", "C2:C61", "expression"),
        ("highlight duplicates in Invoice", "Purchase Data", "B2:B41", "expression"),
        ("highlight blank cells in Column1", "Tasks", "G2:G26", "expression"),
        ("colour scale on sales", "Sales Data", "K2:K61", "colorScale"),
        ("top 10 sales highlight karo", "Sales Data", "K2:K61", "top10"),
        ("status Pending wale yellow", "Purchase Data", "L2:L41", "expression"),
    ],
)
def test_saved_file_has_rule_on_data_cells_only(tmp_path, text, sheet, expected_range, rule_type):
    _, wb = _do(text, tmp_path)
    ws = wb[sheet]
    assert len(ws.conditional_formatting) == 1
    cf = next(iter(ws.conditional_formatting))
    assert str(cf.sqref) == expected_range
    assert ws.conditional_formatting._cf_rules[cf][0].type == rule_type


@pytest.mark.parametrize(
    "text",
    [
        "which amounts are above 1 lakh",
        "which sales are in the top 10",
        "what cells are blank",
        "are any invoices duplicated",
        "kaunse sales 1 lakh se upar hain",
    ],
)
def test_questions_are_not_format_orders(text):
    assert edits.plan_edit(text, load_tables(str(PRACTICE))) is None
