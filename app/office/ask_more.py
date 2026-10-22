"""More jobs for the Excel agent: review the file, reconcile two sheets (bank vs ledger), month-on-month growth, running
totals, empty cells, sorting, budget-vs-actual variance, calculated columns (line total, commission %, full name),
two-way pivots, and charts next to summaries. Each one answers in words, writes the result into the file, and gives the
Excel formula that does the same."""

import re
from collections import defaultdict
from datetime import date, datetime

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.office.ask import (NeedMore, Skip, _answer_sheet, _choose_table, _fmt, _label, _match_filters, _name, _num, _write_table,
                            columns_in, pick_sheet, values_in, _agg_of)  # noqa: F401


def _data_tables(tables):
    return [t for t in tables if len(t.rows) >= 2 and sum(1 for h in t.header if h) >= 2]


def _span(t):
    q = t.ref(0)[:t.ref(0).index("!") + 1]
    return f"{q}A{t.header_row + 1}:{get_column_letter(len(t.header))}{t.header_row + len(t.rows)}"


# ------------------------------------------------------------------ review the whole file
def _plan_review(text, tables):
    t = pick_sheet(text, tables)
    return {"job": "review", "table": t.name if t else None}


def _run_review(p, tables, wb, hi, n):
    from app import office

    ts = [tables[p["table"]]] if p.get("table") else _data_tables(list(tables.values()))
    lines = []
    for t in ts:
        nums = [i for i in range(len(t.header)) if t.kind(i) == "number" and t.header[i]]
        money = [i for i in nums if re.search(r"amount|total|net|sales|salary|value|fee|price", t.header[i], re.I)]
        key = "; ".join(f"total {t.header[i]} {_fmt(sum(v for v in t.col(i) if isinstance(v, (int, float))))}" for i in money[:2])
        empty = [(t.header[i], sum(1 for v in t.col(i) if v in (None, ""))) for i in range(len(t.header)) if t.header[i]]
        empty = [(f"{h}: {c} khaali" if hi else f"{h}: {c} empty") for h, c in empty if c]
        lines.append(f"{t.name}: {len(t.rows)} rows, {sum(1 for h in t.header if h)} columns" + (f"; {key}" if key else "")
                     + (f"; {', '.join(empty[:3])}" if empty else ""))
    issues = ""
    if p.get("_path"):
        s = office.scan(p["_path"])["summary"]
        found = [(s.get("broken_refs"), "toote formula", "broken formulas"), (s.get("errors"), "error wale cells", "error cells"),
                 (s.get("duplicate_rows"), "duplicate rows", "duplicate rows"), (s.get("text_numbers"), "text wale numbers", "numbers saved as text"),
                 (s.get("spaces"), "extra space wale cells", "cells with extra spaces"),
                 (s.get("case_mixed"), "capital/small letters alag", "mixed capitals"), (s.get("odd_formulas"), "alag formula wali rows", "rows with an odd formula")]
        found = [f"{k} {h if hi else e}" for k, h, e in found if k]
        if hi:
            issues = (" Galtiyan: " + ", ".join(found) + ". 'sab theek karo' likho to main theek kar dunga.") if found else " Koi galti nahi mili."
        else:
            issues = (" Problems: " + ", ".join(found) + ". Type 'fix everything' and I'll fix them.") if found else " No problems found."
    return {"say": "File review: " + " | ".join(lines) + "." + issues, "formula": "", "sheet": None}


# ------------------------------------------------------------------ reconcile: bank vs ledger, POS vs deposits
def _two_tables(text, tables):
    named = [t for t in tables if pick_sheet(text, [t])]
    if len(named) >= 2:
        return named[0], named[1]
    data = _data_tables(tables)
    if len(data) == 2:
        return data[0], data[1]
    raise NeedMore("Kaunsi do sheets milani hain? Dono ke naam likho, jaise \"Bank aur Ledger reconcile karo\".",
                   "Which two sheets should I match? Name both, e.g. \"reconcile Bank and Ledger\".", [t.name for t in data][:8])


def _amount_col(t):
    nums = [i for i in range(len(t.header)) if t.kind(i) == "number" and t.header[i]]
    pref = [i for i in nums if re.search(r"amount|amt|value|total|debit|credit|withdraw|deposit|net", t.header[i], re.I)]
    return (pref or nums or [None])[0]


def _plan_reconcile(text, tables):
    a, b = _two_tables(text, tables)
    ai, bi = _amount_col(a), _amount_col(b)
    if ai is None or bi is None:
        raise NeedMore("Dono sheets mein amount ka column chahiye.", "Both sheets need an amount column.")
    ad = next((i for i in range(len(a.header)) if a.kind(i) == "date"), None)
    bd = next((i for i in range(len(b.header)) if b.kind(i) == "date"), None)
    return {"job": "reconcile", "a": a.name, "b": b.name, "a_amt": ai, "b_amt": bi, "a_date": ad, "b_date": bd}


def _run_reconcile(p, tables, wb, hi, n):
    a, b = tables[p["a"]], tables[p["b"]]
    free = defaultdict(list)
    for j, r in enumerate(b.rows):
        v = _num(r[p["b_amt"]]) if p["b_amt"] < len(r) else None
        if v is not None:
            free[round(abs(v), 2)].append(j)
    matched, only_a, used = [], [], set()
    dated = p["a_date"] is not None and p["b_date"] is not None
    for i, r in enumerate(a.rows):
        v = _num(r[p["a_amt"]]) if p["a_amt"] < len(r) else None
        if v is None:
            continue
        cands = [j for j in free.get(round(abs(v), 2), []) if j not in used]
        if dated:
            da = r[p["a_date"]]

            def gap(j, da=da):
                db = b.rows[j][p["b_date"]]
                return abs((db - da).days) if isinstance(da, (datetime, date)) and isinstance(db, (datetime, date)) else 999
            cands = sorted([j for j in cands if gap(j) <= 7], key=gap)
        if cands:
            used.add(cands[0])
            matched.append((i, cands[0]))
        else:
            only_a.append(i)
    only_b = [j for j in range(len(b.rows)) if j not in used and p["b_amt"] < len(b.rows[j]) and _num(b.rows[j][p["b_amt"]]) is not None]
    sa = sum(_num(a.rows[i][p["a_amt"]]) or 0 for i in only_a)
    sb = sum(_num(b.rows[j][p["b_amt"]]) or 0 for j in only_b)
    sname = None
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, "Reconciliation")
        ws.append([f"Matched (same amount{', date within 7 days' if dated else ''})", len(matched)])
        ws.append([f"Only in {a.name}", len(only_a), sa])
        ws.append([f"Only in {b.name}", len(only_b), sb])
        ws.append([])
        ws.append([f"Only in {a.name}"] + a.header)
        for i in only_a:
            ws.append([f"row {a.header_row + 1 + i}"] + list(a.rows[i]))
        ws.append([])
        ws.append([f"Only in {b.name}"] + b.header)
        for j in only_b:
            ws.append([f"row {b.header_row + 1 + j}"] + list(b.rows[j]))
    formula = f"=COUNTIFS({b.ref(p['b_amt'])},{get_column_letter(p['a_amt'] + 1)}{a.header_row + 1})   (0 = not in {b.name})"
    if hi:
        say = (f"{a.name} vs {b.name}: {len(matched)} entries mil gayi (same amount{', 7 din ke andar' if dated else ''}). "
               f"Sirf {a.name} mein: {len(only_a)} ({_fmt(sa)}). Sirf {b.name} mein: {len(only_b)} ({_fmt(sb)}). Dono list nayi sheet mein hain.")
    else:
        say = (f"{a.name} vs {b.name}: {len(matched)} entries match (same amount{', within 7 days' if dated else ''}). "
               f"Only in {a.name}: {len(only_a)} ({_fmt(sa)}). Only in {b.name}: {len(only_b)} ({_fmt(sb)}). Both lists are in a new sheet.")
    return {"say": say, "formula": formula, "sheet": sname}


# ------------------------------------------------------------------ month-on-month growth
def _plan_growth(text, tables):
    t = _choose_table(text, tables, {"number"})
    dates = columns_in(text, t, {"date"}) or [(i, 1) for i in range(len(t.header)) if t.kind(i) == "date"]
    nums = columns_in(text, t, {"number"})
    if not dates:
        raise Skip()
    m = nums[0][0] if nums else _amount_col(t)
    if m is None:
        raise Skip()
    return {"job": "growth", "table": t.name, "date": dates[0][0], "measure": m, "filters": values_in(text, t)}


def _run_growth(p, tables, wb, hi, n):
    t = tables[p["table"]]
    by = defaultdict(float)
    for r in t.rows:
        if not _match_filters(r, p["filters"]):
            continue
        d, v = r[p["date"]], _num(r[p["measure"]])
        if isinstance(d, (datetime, date)) and v is not None:
            by[d.strftime("%Y-%m")] += v
    table, prev = [], None
    for k in sorted(by):
        ch = None if prev in (None, 0) else (by[k] - prev) / prev * 100
        table.append((k, by[k], None if ch is None else round(ch, 1)))
        prev = by[k]
    mname = t.header[p["measure"]]
    formula = "=(B3-B2)/B2   (format as %)"
    sname = None
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, f"{mname} growth")
        _write_table(ws, ["Month", mname, "Change %"], table, formula)
        _chart(ws, len(table), f"{mname} by month", line=True)
    best = max((x for x in table if x[2] is not None), key=lambda x: x[2], default=None)
    worst = min((x for x in table if x[2] is not None), key=lambda x: x[2], default=None)
    lines = "; ".join(f"{k}: {_fmt(v)}" + ("" if c is None else f" ({c:+.1f}%)") for k, v, c in table[-6:])
    if hi:
        say = f"{t.name}, mahine-dar-mahine {mname}: {lines}." + (
            f" Sabse zyada badha: {best[0]} ({best[2]:+.1f}%), sabse zyada gira: {worst[0]} ({worst[2]:+.1f}%)." if best else "")
    else:
        say = f"{t.name}, {mname} month on month: {lines}." + (f" Best: {best[0]} ({best[2]:+.1f}%), worst: {worst[0]} ({worst[2]:+.1f}%)." if best else "")
    return {"say": say, "formula": formula, "sheet": sname}


# ------------------------------------------------------------------ running total
def _plan_running(text, tables):
    t = _choose_table(text, tables, {"number"})
    nums = columns_in(text, t, {"number"})
    m = nums[0][0] if nums else _amount_col(t)
    if m is None:
        raise NeedMore("Kis column ka running total chahiye?", "Running total of which column?")
    return {"job": "running", "table": t.name, "measure": m}


def _run_running(p, tables, wb, hi, n):
    t = tables[p["table"]]
    m = p["measure"]
    c = get_column_letter(m + 1)
    first = t.header_row + 1
    total = sum(_num(r[m]) or 0 for r in t.rows)
    if wb is not None:
        ws = wb[t.name]
        new = len(t.header) + 1
        ws.cell(t.header_row, new, f"Running total of {t.header[m]}").font = Font(bold=True)
        for k in range(len(t.rows)):
            ws.cell(first + k, new, f"=SUM(${c}${first}:{c}{first + k})")
    say = (f"{t.name} mein naya column \"Running total of {t.header[m]}\" jod diya; aakhri row mein {_fmt(total)}." if hi else
           f"Added \"Running total of {t.header[m]}\" to {t.name}; the last row shows {_fmt(total)}.")
    return {"say": say, "formula": f"=SUM(${c}${first}:{c}{first})  then drag down", "sheet": t.name}


# ------------------------------------------------------------------ empty cells
def _plan_missing(text, tables):
    t = pick_sheet(text, tables)
    if t is None:
        cols = [(o, columns_in(text, o)) for o in _data_tables(tables)]
        cols = [(o, c) for o, c in cols if c]
        t = cols[0][0] if cols else None
    if t is None:
        return {"job": "missing", "table": None, "col": None}
    c = columns_in(text, t)
    return {"job": "missing", "table": t.name, "col": c[0][0] if c else None}


def _run_missing(p, tables, wb, hi, n):
    ts = [tables[p["table"]]] if p.get("table") else _data_tables(list(tables.values()))
    if p.get("col") is not None:
        t, i = ts[0], p["col"]
        rows = [r for r in t.rows if (r[i] if i < len(r) else None) in (None, "")]
        sname = None
        if wb is not None:
            ws, sname = _answer_sheet(wb, n, f"Missing {t.header[i]}")
            _write_table(ws, t.header, rows, f'=FILTER({_span(t)},{t.ref(i)}="")')
        label = _label(t)
        names = ", ".join(_name(t, r, label) for r in rows[:10])
        say = (f"{t.name}: {len(rows)} rows mein {t.header[i]} khaali hai. {names}" if hi else
               f"{t.name}: {t.header[i]} is empty in {len(rows)} rows. {names}")
        return {"say": say, "formula": f"=COUNTBLANK({t.ref(i)})", "sheet": sname}
    parts = []
    for t in ts:
        e = [(t.header[i], sum(1 for v in t.col(i) if v in (None, ""))) for i in range(len(t.header)) if t.header[i]]
        e = [f"{h} {c}" for h, c in e if c]
        if e:
            parts.append(f"{t.name}: {', '.join(e)}")
    if hi:
        say = ("Khaali cells: " + " | ".join(parts)) if parts else "Koi khaali cell nahi mila."
    else:
        say = ("Empty cells: " + " | ".join(parts)) if parts else "No empty cells found."
    return {"say": say, "formula": "=COUNTBLANK(range)", "sheet": None}


# ------------------------------------------------------------------ sort
def _plan_sort(text, tables):
    t = _choose_table(text, tables)
    cols = columns_in(text, t)
    if not cols:
        raise NeedMore(f"{t.name} ko kis column se sort karun?", f"Sort {t.name} by which column?", [h for h in t.header if h][:8])
    desc = bool(re.search(r"\b(desc\w*|bade se chhote|zyada se kam|highest|z to a|ulta|reverse|top)\b", text or "", re.I))
    return {"job": "sort", "table": t.name, "col": cols[0][0], "desc": desc}


def _run_sort(p, tables, wb, hi, n):
    t, i = tables[p["table"]], p["col"]

    def key(r):
        v = r[i] if i < len(r) else None
        return (0, v.lower()) if isinstance(v, str) else (0, v) if v is not None else (1, 0)
    filled = [r for r in t.rows if (r[i] if i < len(r) else None) is not None]
    empty = [r for r in t.rows if (r[i] if i < len(r) else None) is None]
    try:
        rows = sorted(filled, key=key, reverse=p["desc"]) + empty
    except TypeError:                                   # mixed text and numbers: sort as text
        rows = sorted(filled, key=lambda r: str(r[i]).lower(), reverse=p["desc"]) + empty
    formula = f"=SORT({_span(t)},{i + 1},{-1 if p['desc'] else 1})"
    sname = None
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, f"Sorted by {t.header[i]}")
        _write_table(ws, t.header, rows, formula)
    order = ("bade se chhote" if p["desc"] else "chhote se bade") if hi else ("high to low" if p["desc"] else "low to high")
    say = (f"{t.name} ko {t.header[i]} ke hisaab se {order} sort karke nayi sheet \"{sname}\" mein rakh diya (original waisi hi hai)." if hi else
           f"Sorted {t.name} by {t.header[i]} ({order}) into a new sheet \"{sname}\" (the original is unchanged).")
    return {"say": say, "formula": formula, "sheet": sname}


# ------------------------------------------------------------------ budget vs actual
def _plan_variance(text, tables):
    t = _choose_table(text, tables, {"number"})
    nums = columns_in(text, t, {"number"})
    if len(nums) < 2:
        bud = [i for i in range(len(t.header)) if t.kind(i) == "number" and re.search(r"budget|target|plan|expected", t.header[i], re.I)]
        act = [i for i in range(len(t.header)) if t.kind(i) == "number" and re.search(r"actual|achieved|spent|real", t.header[i], re.I)]
        if bud and act:
            return {"job": "variance", "table": t.name, "a": act[0], "b": bud[0]}
        raise Skip()
    a, b = nums[0][0], nums[1][0]
    if re.search(r"budget|target|plan|expected", t.header[a], re.I) and not re.search(r"budget|target|plan|expected", t.header[b], re.I):
        a, b = b, a                                    # actual - budget
    return {"job": "variance", "table": t.name, "a": a, "b": b}


def _run_variance(p, tables, wb, hi, n):
    t = tables[p["table"]]
    a, b = p["a"], p["b"]
    ca, cb = get_column_letter(a + 1), get_column_letter(b + 1)
    first = t.header_row + 1
    pairs = [(_num(r[a]), _num(r[b])) for r in t.rows]
    over = sum(1 for x, y in pairs if x is not None and y is not None and x > y)
    diff = sum((x or 0) - (y or 0) for x, y in pairs)
    if wb is not None:
        ws = wb[t.name]
        c1, c2 = len(t.header) + 1, len(t.header) + 2
        ws.cell(t.header_row, c1, f"Difference ({t.header[a]} - {t.header[b]})").font = Font(bold=True)
        ws.cell(t.header_row, c2, "Difference %").font = Font(bold=True)
        for k in range(len(t.rows)):
            r = first + k
            ws.cell(r, c1, f"={ca}{r}-{cb}{r}")
            ws.cell(r, c2, f'=IFERROR(({ca}{r}-{cb}{r})/{cb}{r},"")').number_format = "0.0%"
    say = (f"{t.name} mein 2 naye column jode: Difference ({t.header[a]} - {t.header[b]}) aur Difference %. Kul farak {_fmt(diff)}; "
           f"{over} rows mein {t.header[a]} {t.header[b]} se zyada hai." if hi else
           f"Added 2 columns to {t.name}: Difference ({t.header[a]} - {t.header[b]}) and Difference %. Total difference {_fmt(diff)}; "
           f"{t.header[a]} is above {t.header[b]} in {over} rows.")
    return {"say": say, "formula": f"={ca}{first}-{cb}{first}   and   =({ca}{first}-{cb}{first})/{cb}{first}", "sheet": t.name}


# ------------------------------------------------------------------ calculated columns
def _plan_calc(text, tables):
    t = _choose_table(text, tables)
    low = (text or "").lower()
    txt = columns_in(text, t, {"text"})
    if re.search(r"full name|naam jodo|naam milao", low):
        names = [i for i, h in enumerate(t.header) if re.search(r"first|last|sur", h or "", re.I)]
        if len(names) >= 2:
            return {"job": "calc", "table": t.name, "op": "&", "a": names[0], "b": names[1]}
    if re.search(r"\b(combine|jodo|milao|merge)\b", low) and len(txt) >= 2:
        return {"job": "calc", "table": t.name, "op": "&", "a": txt[0][0], "b": txt[1][0]}
    nums = columns_in(text, t, {"number"})
    pct = re.search(r"(\d+(?:\.\d+)?)\s*%", low)
    if pct and nums and not re.search(r"\b(wale|wali|wala|kitne|kitni|count|ginti|jinka|jinki|where|jahan)\b", low):
        name = re.search(r"\b(commission|gst|tax|discount|bonus|incentive|tds)\b", low)
        return {"job": "calc", "table": t.name, "op": "%", "a": nums[0][0], "rate": float(pct.group(1)),
                "name": (name.group(1).upper() if name and name.group(1) in ("gst", "tds") else (name.group(1).title() if name else "Amount"))
                + f" ({pct.group(1)}%)"}
    if re.search(r"line total", low):
        q = next((i for i, h in enumerate(t.header) if re.search(r"qty|quantity", h or "", re.I)), None)
        pr = next((i for i, h in enumerate(t.header) if re.search(r"price|rate", h or "", re.I)), None)
        if q is not None and pr is not None:
            return {"job": "calc", "table": t.name, "op": "*", "a": q, "b": pr, "name": "Line Total"}
    if len(nums) >= 2:
        op = ("*" if re.search(r"multiply|multiplied|guna|into|times|×|\*|\bx\b", low) else "-" if re.search(r"\bminus\b|\bghata\b| - ", low)
              else "/" if re.search(r"divide|bhag|ratio|/", low) else "+" if re.search(r"\bplus\b|\+", low) else None)
        if op:
            return {"job": "calc", "table": t.name, "op": op, "a": nums[0][0], "b": nums[1][0]}
    raise Skip()


def _run_calc(p, tables, wb, hi, n):
    t = tables[p["table"]]
    a, op = p["a"], p["op"]
    ca = get_column_letter(a + 1)
    cb = get_column_letter(p["b"] + 1) if "b" in p else None
    first = t.header_row + 1
    if op == "&":
        name = "Full Name" if re.search(r"first|last", t.header[a] + t.header[p["b"]], re.I) else f"{t.header[a]} + {t.header[p['b']]}"

        def f(r):
            return f'=TRIM({ca}{r}&" "&{cb}{r})'
    elif op == "%":
        name = p["name"]

        def f(r):
            return f"={ca}{r}*{p['rate']:g}%"
    else:
        sym = {"*": "×", "-": "-", "/": "÷", "+": "+"}[op]
        name = p.get("name") or f"{t.header[a]} {sym} {t.header[p['b']]}"

        def f(r):
            return f'=IFERROR({ca}{r}/{cb}{r},"")' if op == "/" else f"={ca}{r}{op}{cb}{r}"
    total = 0.0
    for r in t.rows:
        x = _num(r[a]) if a < len(r) else None
        y = _num(r[p["b"]]) if "b" in p and p["b"] < len(r) else None
        if op == "%" and x is not None:
            total += x * p["rate"] / 100
        elif op in ("*", "-", "+") and x is not None and y is not None:
            total += {"*": x * y, "-": x - y, "+": x + y}[op]
    if wb is not None:
        ws = wb[t.name]
        c = len(t.header) + 1
        ws.cell(t.header_row, c, name).font = Font(bold=True)
        for k in range(len(t.rows)):
            ws.cell(first + k, c, f(first + k))
    tot = "" if op in ("&", "/") else (f" Kul: {_fmt(total)}." if hi else f" Total: {_fmt(total)}.")
    say = (f"{t.name} mein naya column \"{name}\" jod diya, har row mein formula ke saath.{tot}" if hi else
           f"Added a column \"{name}\" to {t.name}, with a formula in every row.{tot}")
    return {"say": say, "formula": f(first), "sheet": t.name}


# ------------------------------------------------------------------ two-way pivot (region x product)
def _run_pivot(p, tables, wb, hi, n):
    t = tables[p["table"]]
    ri, ci, m = p["rows"], p["cols"], p["measure"]
    cells = defaultdict(list)
    rk, ck = [], []
    for r in t.rows:
        if not _match_filters(r, p["filters"]):
            continue
        x, y = r[ri], r[ci]
        if x in (None, "") or y in (None, ""):
            continue
        x, y = str(x).strip(), str(y).strip()
        if x not in rk:
            rk.append(x)
        if y not in ck:
            ck.append(y)
        cells[(x, y)].append(_num(r[m]) if m is not None else 1)

    def agg(v):
        v = [z for z in v if z is not None]
        if not v:
            return None
        return {"sum": sum(v), "avg": sum(v) / len(v), "count": len(v)}.get(p["agg"], sum(v))
    mname = t.header[m] if m is not None else "count"
    if m is not None and p["agg"] != "count":
        fn = {"sum": "SUMIFS", "avg": "AVERAGEIFS"}.get(p["agg"], "SUMIFS")
        formula = f"={fn}({t.ref(m)},{t.ref(ri)},$A2,{t.ref(ci)},B$1)"
    else:
        formula = f"=COUNTIFS({t.ref(ri)},$A2,{t.ref(ci)},B$1)"
    sname = None
    if wb is not None:
        ws, sname = _answer_sheet(wb, n, f"{t.header[ri]} x {t.header[ci]}")
        ws.append([f"{t.header[ri]} / {t.header[ci]}"] + ck + ["Total"])
        for x in rk:
            vals = [agg(cells.get((x, y), [])) for y in ck]
            ws.append([x] + vals + [sum(v for v in vals if v is not None)])
        ws.append([])
        ws.append(["Excel formula:", formula])
    top = max(((x, y, agg(v)) for (x, y), v in cells.items()), key=lambda z: z[2] or 0, default=None)
    if hi:
        say = (f"{t.name}: {t.header[ri]} × {t.header[ci]} ka {mname} table nayi sheet \"{sname}\" mein ({len(rk)} × {len(ck)})."
               + (f" Sabse zyada: {top[0]} / {top[1]} = {_fmt(top[2])}." if top else ""))
    else:
        say = (f"{t.name}: {mname} by {t.header[ri]} × {t.header[ci]} in a new sheet \"{sname}\" ({len(rk)} × {len(ck)})."
               + (f" Highest: {top[0]} / {top[1]} = {_fmt(top[2])}." if top else ""))
    return {"say": say, "formula": formula, "sheet": sname}


# ------------------------------------------------------------------ charts
def _chart(ws, rows, title, line=False):
    """A chart next to a summary table (labels in column A, values in column B)."""
    from openpyxl.chart import BarChart, LineChart, Reference

    if rows < 2:
        return
    ch = LineChart() if line else BarChart()
    ch.title, ch.height, ch.width, ch.legend = title[:60], 8, 16, None
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=rows + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=rows + 1))
    ws.add_chart(ch, "E2")
