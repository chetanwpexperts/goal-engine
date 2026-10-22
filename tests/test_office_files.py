"""The Microsoft agent's hands work on real files: scan, fix (a copy, with every change listed), compare, big files."""

import os
import resource
import time

from openpyxl import Workbook, load_workbook

from app import office


def _book(tmp_path, name="a.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Invoice No", "Customer", "Qty", "Rate", "Amount"])
    for i in range(2, 12):
        ws.append([f"INV{i:03}", f" Customer  {i} ", i, "1,200" if i == 3 else 100, f"=C{i}*D{i}"])
    ws["E7"] = "=C7*#REF!"           # broken reference
    ws["E9"] = "=C9+D9"             # the odd one out
    ws.append(["INV011", " Customer  11 ", 11, 100, "=C12*D12"])  # duplicate of row 11? (different amount formula row)
    ws["B4"].font = ws["B4"].font.copy(bold=True)
    p = tmp_path / name
    wb.save(p)
    return str(p)


def test_scan_finds_the_real_problems(tmp_path):
    r = office.scan(_book(tmp_path))
    s = r["sheets"][0]
    assert s["broken_refs"] == ["E7"]
    assert "E9" in s["odd_formulas"]
    assert "D3" in s["text_number_cells"]
    assert s["spaces"] >= 10
    assert r["summary"]["rows"] == 11


def test_fix_makes_a_copy_and_lists_changes(tmp_path):
    src = _book(tmp_path)
    before = open(src, "rb").read()
    out = str(tmp_path / "fixed.xlsx")
    res = office.fix(src, out, ["trim_spaces", "text_to_number", "fix_broken_refs"])
    assert open(src, "rb").read() == before                      # original untouched
    assert res["formatting_kept"] and res["changes"]["fix_broken_refs"] == 1
    wb = load_workbook(out)
    ws = wb["Sales"]
    assert ws["E7"].value == "=C7*D7"
    assert ws["D3"].value == 1200
    assert ws["B2"].value == "Customer 2"
    assert ws["B4"].font.bold
    assert wb["Changes by Goal Engine"].max_row > 10


def test_text_number_keeps_ids_and_phones():
    assert office.text_number("1,20,000") == 120000
    assert office.text_number("₹ 450") == 450
    assert office.text_number("15%") == 0.15
    for keep in ("00123", "9876543210123", "INV001", "110 001 A"):
        assert office.text_number(keep) is None


def _sheet(path, rows, header=("Emp ID", "Name", "Salary")):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def test_compare_two_files(tmp_path):
    a, b, out = str(tmp_path / "old.xlsx"), str(tmp_path / "new.xlsx"), str(tmp_path / "report.xlsx")
    _sheet(a, [("E1", "Asha", 1000), ("E2", "Ravi", 2000), ("E3", "Neha", 3000)])
    _sheet(b, [("E1 ", "asha", "1,000"), ("E2", "Ravi", 2500), ("E4", "Amit", 4000)])
    r = office.compare(a, b, out)
    assert r["key"] == "Emp ID"
    assert (r["only_in_a"], r["only_in_b"], r["changed_rows"], r["changed_cells"]) == (1, 1, 1, 1)
    wb = load_workbook(out, read_only=True)
    changed = list(wb["Changed"].iter_rows(values_only=True))
    assert changed[1][:4] == ("E2", "Salary", "2000", "2500")


def test_big_file_streams_with_low_memory(tmp_path):
    n = 150_000
    a, b, out = str(tmp_path / "big_a.xlsx"), str(tmp_path / "big_b.xlsx"), str(tmp_path / "big_report.xlsx")
    _sheet(a, ((f"E{i}", f"Name {i}", i) for i in range(n)))
    _sheet(b, ((f"E{i}", f"Name {i}", i + (1 if i % 1000 == 0 else 0)) for i in range(1, n + 1)))
    start, mem0 = time.time(), resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    r = office.compare(a, b, out)
    took = time.time() - start
    grew_mb = (resource.getrusage(resource.RUSAGE_SELF).ru_maxrss - mem0) / 1024
    assert (r["only_in_a"], r["only_in_b"], r["changed_rows"]) == (1, 1, n // 1000 - 1)
    assert grew_mb < 300, grew_mb
    assert took < 120, took
    scan = office.scan(a)
    assert scan["summary"]["rows"] == n


def test_csv_fix(tmp_path):
    p = tmp_path / "bank.csv"
    p.write_text("Date,Narration,Amount\n01/10/2026, UPI  payment ,\"1,250.50\"\n01/10/2026, UPI  payment ,\"1,250.50\"\n", encoding="utf-8")
    out = str(tmp_path / "bank_fixed.csv")
    r = office.fix(str(p), out, ["trim_spaces", "text_to_number", "remove_duplicates"])
    lines = open(out, encoding="utf-8-sig").read().splitlines()
    assert lines[1] == "01/10/2026,UPI payment,1250.5" and len(lines) == 2
    assert r["changes"]["remove_duplicates"] == 1


def test_same_word_different_capitals_is_found_and_fixed(tmp_path):
    from openpyxl import Workbook as WB

    wb = WB()
    ws = wb.active
    ws.title = "Messy"
    ws.append(["Name", "Department"])
    for n, d in [("A", "Finance"), ("B", "finance"), ("C", "Finance"), ("D", "sales"), ("E", "Sales"), ("F", "Sales"), ("G", "IT")]:
        ws.append([n, d])
    src = str(tmp_path / "messy.xlsx")
    wb.save(src)
    r = office.scan(src)
    assert r["summary"]["case_mixed"] == 2 and set(r["sheets"][0]["case_cells"]) == {"B3", "B5"}
    out = str(tmp_path / "messy_fixed.xlsx")
    res = office.fix(src, out, ["fix_case"])
    assert res["changes"]["fix_case"] == 2
    got = [c.value for c in load_workbook(out)["Messy"]["B"]][1:]
    assert got == ["Finance", "Finance", "Finance", "Sales", "Sales", "Sales", "IT"]
    first = list(load_workbook(out)["Changes by Goal Engine"].iter_rows(values_only=True))[1]
    assert first[0].startswith("Capital/small letters")
