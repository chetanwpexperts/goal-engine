"""Everyday Excel orders done on the file itself (English / Hinglish), and questions that must NOT be read as orders."""

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.office import ask
from app.office.sheetdata import load_tables
from tests.test_follow_orders import open_file  # noqa: F401 - fixture
from tests.test_ml import say
from tests.test_office_agent import _wait

PRACTICE = Path(__file__).parent / "data" / "practice.xlsx"


def _do(text, tmp_path):
    p = ask.understand(text, load_tables(str(PRACTICE)))
    assert p["job"] == "edit", (text, p)
    out = tmp_path / "out.xlsx"
    r = ask.run(p, str(PRACTICE), str(out), False, 1)
    return p, r, load_workbook(out)


ORDERS = [
    ("make the header bold", "bold_header"), ("header bold karo", "bold_header"), ("freeze the top row", "freeze"),
    ("fit column widths", "autofit"), ("add filter buttons", "filter"), ("filter laga do", "filter"),
    ("remove blank rows", "remove_blank_rows"), ("khaali rows hatao", "remove_blank_rows"), ("add serial numbers", "serial"),
    ("split sales data by region into separate sheets", "split_sheets"), ("region wise alag sheets bana do", "split_sheets"),
    ("add a dropdown to Status with Paid, Pending, Overdue", "dropdown"), ("rename Net Sales to Net Amount", "rename"),
    ("change date format to dd/mm/yyyy", "date_format"), ("show net sales in rupee format", "money_format"),
    ("round net sales to 2 decimals", "round"), ("make region uppercase", "case"), ("replace Delhi with New Delhi", "replace"),
]


@pytest.mark.parametrize("text,op", ORDERS, ids=[o[0][:40] for o in ORDERS])
def test_orders_are_recognised(text, op):
    p = ask.understand(text, load_tables(str(PRACTICE)))
    assert p["job"] == "edit" and p["op"] == op, (text, p)


@pytest.mark.parametrize("text", ["how many names are in capital letters", "filter rows where region is north", "total sales kitni hai",
                                  "month wise sales", "sales above 1 lakh show karo", "change in sales from jan to feb",
                                  "what is the date format", "sales by region"])
def test_questions_are_not_orders(text):
    try:
        p = ask.understand(text, load_tables(str(PRACTICE)))
    except ask.NeedMore:
        return                                          # asked back: still not a change to the file
    assert p["job"] != "edit", (text, p)


def test_freeze_bold_filter_on_every_sheet(tmp_path):
    _, r, wb = _do("freeze the top row", tmp_path)
    assert wb["Sales Data"].freeze_panes == "A2" and wb["Messy Data"].freeze_panes == "A2"
    _, r, wb = _do("make the header bold", tmp_path)
    assert wb["Sales Data"]["A1"].font.bold and "every sheet" in r["say"]
    _, r, wb = _do("add filter buttons in sales data", tmp_path)
    assert wb["Sales Data"].auto_filter.ref.startswith("A1:")


def test_dropdown_goes_on_the_named_column(tmp_path):
    p, r, wb = _do("add a dropdown to Status with Paid, Pending, Overdue", tmp_path)
    dv = wb["Tasks"].data_validations.dataValidation[0]
    assert dv.formula1 == '"Paid,Pending,Overdue"' and str(dv.sqref).startswith("F2:")


def test_one_sheet_per_value(tmp_path):
    _, r, wb = _do("region wise alag sheets bana do", tmp_path)
    assert {"East", "North", "South", "West"} <= set(wb.sheetnames)
    assert sum(wb[n].max_row - 1 for n in ("East", "North", "South", "West")) == 60


def test_rename_case_replace_round(tmp_path):
    _, r, wb = _do("rename Net Sales to Net Amount", tmp_path)
    assert "Net Amount" in [c.value for c in wb["Sales Data"][1]]
    _, r, wb = _do("make region uppercase", tmp_path)
    assert all(str(wb["Sales Data"].cell(i, 4).value).isupper() for i in range(2, 62))
    _, r, wb = _do("replace Delhi with New Delhi", tmp_path)
    vals = [c.value for row in wb["Employee Data"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "Delhi" not in vals and "New New Delhi" not in vals and "New Delhi" in vals
    _, r, wb = _do("round net sales to 2 decimals", tmp_path)
    assert str(wb["Sales Data"]["K2"].value).startswith("=ROUND(")


def test_missing_column_asks_with_the_real_columns():
    with pytest.raises(ask.NeedMore) as e:
        ask.understand("delete the Remarks column", load_tables(str(PRACTICE)))
    assert "Order ID" in e.value.en


def test_macro_and_pivot_are_answered_honestly(tmp_path):
    with pytest.raises(ask.NeedMore) as e:
        ask.understand("write a macro to copy rows", load_tables(str(PRACTICE)))
    assert "macros" in e.value.en
    p = ask.understand("make a pivot table of sales by region", load_tables(str(PRACTICE)))
    r = ask.run(p, str(PRACTICE), str(tmp_path / "o.xlsx"), False, 1)
    assert "South: 12,33,200" in r["say"] and "PivotTable itself can't be built" in r["say"]


def test_order_in_chat_changes_the_file_and_undo_puts_it_back(client, open_file):  # noqa: F811
    r = say(client, "make region uppercase", open_file)
    assert "UPPER CASE" in r["reply"] and "undo" in r["reply"], r["reply"]
    f = _wait(client, next(a for a in r["agents"] if a.get("office"))["files"][0]["id"])
    assert len(f["steps"]) == 1
    wb = load_workbook(io.BytesIO(client.get(f"/api/office/files/{f['id']}/download").content))
    assert wb["Sales Data"]["D2"].value.isupper()
    r = say(client, "undo", open_file)
    f = _wait(client, f["id"])
    assert f["steps"] == []
    r = say(client, "header bold karo", open_file)
    assert "bold" in r["reply"] and "undo" in r["reply"], r["reply"]
